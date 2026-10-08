"""The model reads the layout of question and table documents; code reads the text.

Three readings, each answered with positions only and verified in code:
  - a Pop Quiz's lesson sections: which paragraphs start a lesson, and its number
  - a Data Bank's field names: which label means Question, Options, SLO, feedback ...
  - a Document of Specifications' tables: which chapter each table is for, and which column holds the lesson
    number, the lesson name, the SLO and the Bloom's level
The exact text is always taken from the file by code. Readings are cached per file, model and prompt version.
When there is no model, or a reading fails verification, the rule-based readers are used as before.
"""
import re

from .fallback import file_sha1
from .llm import LLMError
from .textutil import normalize, to_western_digits

PROMPT_VERSION = "l1"


def _cached(ctx, doc_path, tag, ask):
    """ask() -> reading (a dict) or None. Cached per (file, model, tag, prompt version); never raises."""
    model = ctx.model
    if model is None:
        return None
    key = f"layout|{PROMPT_VERSION}|{tag}|{model.client.name}|{file_sha1(doc_path)}"
    hit = model.cache.get(key) if model.cache is not None else None
    if hit is not None:
        return hit.get("reading")
    try:
        reading = ask()
    except LLMError as e:
        model.stats.append({"doc": str(doc_path), "structure": True, "usable": False, "note": f"{tag} not read: {e}"})
        return None
    if model.cache is not None:
        model.cache.put(key, {"reading": reading})
    model.stats.append({"doc": str(doc_path), "structure": True, "usable": reading is not None, "note": tag})
    return reading


def _int(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


# ------------------------------------------------------------------ Pop Quiz lesson sections
LESSON_SYSTEM = """You read a school quiz document (Urdu or English) that has a section for each lesson.
You get it as numbered paragraphs: "[12] text". Find the paragraphs that START a lesson's section, such as
"Lesson 3: States of Matter", "سبق 2: صبر و تحمل" or "Chapter 2 - Lesson 1". Questions, answer options, feedback
and instructions are not section starts.

Answer with JSON only: {"lessons": [{"para": 4, "lesson": 1, "part": ""}]}
"lesson" is the lesson's number as written; "part" is a part letter for a split lesson (e.g. "A" for 2A), else "".
If the document has no lesson sections, answer {"lessons": []}."""


def read_lesson_heads(ctx, doc, info):
    """[{num, variant, start, title}] lesson section starts read by the model, or None."""
    from .fallback import chunks_of, listing, render

    def ask():
        items = [(i, t) for i, t in listing(info) if not info.paras[i].in_table]
        out = []
        for chunk in chunks_of(items):
            reply = ctx.model.client.chat_json(LESSON_SYSTEM, render(chunk))
            rows = reply.get("lessons") if isinstance(reply, dict) else None
            if not isinstance(rows, list):
                raise LLMError("the model's answer had no 'lessons' list")
            out += [r for r in rows if isinstance(r, dict)]
        return {"lessons": out}

    reading = _cached(ctx, doc.abs, "lesson-heads", ask)
    if not reading:
        return None
    heads, seen = [], set()
    for r in reading.get("lessons", []):
        i, n = _int(r.get("para")), _int(r.get("lesson"))
        if i is None or n is None or not (0 <= i < len(info.paras)) or i in seen:
            continue
        p = info.paras[i]
        t = to_western_digits(normalize(p.text))
        # verified: a short line outside a table that carries the lesson's number (or no number at all)
        digits = re.findall(r"\d+", t)
        if p.in_table or not t or len(t) > 160 or (digits and str(n) not in digits):
            continue
        part = str(r.get("part") or "").strip().upper()[:1]
        heads.append({"num": n, "variant": part if part.isalpha() else "", "start": i, "title": p.text.strip()})
        seen.add(i)
    heads.sort(key=lambda h: h["start"])
    return heads or None


def group_lessons(heads, n_paras):
    """Lesson heads -> chapter groups (lesson numbers restart in each chapter), with each head's end set."""
    for i, h in enumerate(heads):
        h["end"] = heads[i + 1]["start"] if i + 1 < len(heads) else n_paras
    groups, prev = [], None
    for h in heads:
        if prev is None or h["num"] < prev["num"] or (h["num"] == prev["num"] and not h["variant"] and not prev["variant"]):
            groups.append([])
        groups[-1].append(h)
        prev = h
    return groups


# ------------------------------------------------------------------ Data Bank fields
FIELDS = ["subject", "chapter_lesson", "chapter", "lesson", "slo", "question_type", "level", "question", "options", "answer",
          "feedback_correct", "feedback_incorrect", "item"]
BANK_SYSTEM = """You read the layout of a school Data Bank (a bank of quiz items), in Urdu or English. Its items are in
tables, either one table per item with the field names in the first column, or one table with the field names in
its header row. You get the field names the document uses, one per line.

Map each field name to one of these fields, or to null if it is none of them:
subject, chapter_lesson (one field holding both chapter and lesson), chapter, lesson, slo (learning outcome),
question_type, level (order level or Bloom's level), question, options (the answer choices), answer (the correct
answer), feedback_correct (feedback for a correct answer), feedback_incorrect (feedback for a wrong answer),
item (the item number).

Answer with JSON only: {"fields": {"<field name exactly as given>": "question"}}"""


def bank_labels(path):
    """(layout, distinct field names): 'rows' when each table holds one item as label | value rows, 'header' when a
    table's first row names its columns and each further row is an item."""
    from docx import Document
    d = Document(path)
    row_labels, head_labels = [], []
    for t in d.tables:
        rows = t.rows
        if not rows:
            continue
        if len(rows[0].cells) >= 4 and len(rows) >= 2:
            head_labels += [c.text.strip() for c in rows[0].cells]
        for r in rows:
            if len(r.cells) >= 2 and r.cells[0].text.strip():
                row_labels.append(r.cells[0].text.strip())
    def distinct(xs):
        return list(dict.fromkeys(x for x in xs if 0 < len(x) <= 60))
    # a two-column table per item repeats its labels in every table; a wide table names fields once in its header
    rows_d, head_d = distinct(row_labels), distinct(head_labels)
    if len(row_labels) >= 2 * max(1, len(rows_d)) and len(rows_d) <= 30:
        return "rows", rows_d
    return ("header", head_d) if head_d else ("rows", rows_d)


def read_bank_fields(ctx, path):
    """(layout, {field name as written: canonical field}) read by the model, or None."""
    layout, labels = bank_labels(path)
    if not labels:
        return None

    def ask():
        reply = ctx.model.client.chat_json(BANK_SYSTEM, "\n".join(labels[:60]))
        fields = reply.get("fields") if isinstance(reply, dict) else None
        if not isinstance(fields, dict):
            raise LLMError("the model's answer had no 'fields' object")
        return {"layout": layout, "fields": fields}

    reading = _cached(ctx, path, "bank-fields", ask)
    if not reading:
        return None
    given = {normalize(l).lower(): l for l in labels}
    out = {}
    for k, v in (reading.get("fields") or {}).items():
        if v in FIELDS and normalize(str(k)).lower() in given:
            out[normalize(str(k)).lower()] = v
    # verified: a bank needs at least a question field the model found among the given names
    if "question" not in out.values():
        return None
    return reading.get("layout", layout), out


# ------------------------------------------------------------------ Document of Specifications
SPEC_SYSTEM = """You read the layout of a school "Document of Specifications" (Urdu or English). It has tables that list,
for each chapter, its lessons and their SLOs (learning outcomes). You get every table as its number, the lines just
before it, and its first rows, with the columns numbered from 0.

For every table that lists SLOs, say which chapter it belongs to (a number), and which column holds the lesson
number (lesson_no), the lesson name (lesson_name), the SLO text (slo) and the Bloom's level (bloom); use null for a
column the table does not have. Leave out tables that do not list SLOs.

Answer with JSON only: {"tables": [{"table": 2, "chapter": 1, "lesson_no": 0, "lesson_name": 1, "slo": 2, "bloom": 3}]}"""


def spec_tables(path, rows=3, cell_chars=70):
    """[(table number, lines before it, [[cell text]] of its first rows, column count)] in document order."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    d = Document(path)
    out, before, n = [], [], 0
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            t = "".join(x.text or "" for x in child.iter(qn("w:t"))).strip()
            if t:
                before = (before + [t[:100]])[-3:]
        elif tag == "tbl":
            table = Table(child, d)
            first = [[c.text.strip().replace("\n", " ")[:cell_chars] for c in r.cells] for r in table.rows[:rows]]
            out.append((n, list(before), first, max((len(r) for r in first), default=0)))
            n += 1
            before = []
    return out


def read_spec_tables(ctx, path):
    """{table number: {chapter, lesson_no, lesson_name, slo, bloom}} read by the model, or None."""
    tables = spec_tables(path)
    if not tables:
        return None

    def ask():
        out, block, budget = [], [], 9000
        def flush():
            reply = ctx.model.client.chat_json(SPEC_SYSTEM, "\n\n".join(block))
            rows = reply.get("tables") if isinstance(reply, dict) else None
            if not isinstance(rows, list):
                raise LLMError("the model's answer had no 'tables' list")
            out.extend(r for r in rows if isinstance(r, dict))
        for n, before, first, _ in tables:
            text = f"Table {n}\nLines before it: " + " / ".join(before) + "\n" + \
                   "\n".join("Row: " + " | ".join(f"[{i}] {c}" for i, c in enumerate(r)) for r in first)
            if block and sum(len(b) for b in block) + len(text) > budget:
                flush()
                block = []
            block.append(text)
        if block:
            flush()
        return {"tables": out}

    reading = _cached(ctx, path, "spec-tables", ask)
    if not reading:
        return None
    ncols = {n: c for n, _, _, c in tables}
    plan = {}
    for r in reading.get("tables", []):
        n, ch = _int(r.get("table")), _int(r.get("chapter"))
        if n not in ncols or ch is None or not (1 <= ch <= 60):
            continue
        cols = {k: _int(r.get(k)) for k in ("lesson_no", "lesson_name", "slo", "bloom")}
        if cols["slo"] is None or cols["lesson_no"] is None:
            continue
        if any(v is not None and not (0 <= v < ncols[n]) for v in cols.values()):
            continue
        plan[n] = {"chapter": ch, **cols}
    return plan or None
