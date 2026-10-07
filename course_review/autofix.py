"""Apply the mechanical Writing & Editing rules to a copy of each document.

Formatting rules have one right answer, so the tool applies them instead of commenting:
page setup, spacing, fonts, sizes by grade, Western digits, heading styles and decimal
numbering, table captions, bullets instead of numbered lists, no colour, a footer, a Table
of Contents for long Lesson Plans, the yellow highlight on correct answers, and the file
name. What needs a person (content, SLOs, the story, feedback quality) is left to comments.

Each fixer returns short lines for the "Fixed by the tool" comment. A Fail is left out of the
comments only when the same check, run again on the fixed copy, passes (see remaining).
"""
import re
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt
from lxml import etree

from .checks.formatting import own_number, run_in_split
from .models import FAIL, NA, PASS, REVIEW
from .textutil import EASTERN_DIGITS, VERSE_MARK, is_arabic_scripture, normalize, script_counts, to_western_digits

# Codes the fixer works on. A Fail on one of them is re-checked on the fixed copy (see remaining).
DOCX_FIXED = {"WE4", "WE6", "WE7", "WE8", "WE10", "WE11", "WE12", "WE13", "LP4", "WE15", "WE16", "WE23", "WE25", "WEG1",
              "WEG2", "LPG2", "LP5", "WE1", "WE2", "WE3"}
PPTX_FIXED = {"WE5", "WE6", "WE7", "WE10", "WE23", "WE1", "WE2", "WE3", "WE25"}


def recheck(ctx, doc, path, rel):
    """The per-document checks run again on the fixed copy at `path`, as if it were named `rel`.
    Returns code -> Finding, or {} when the copy cannot be read (then nothing counts as fixed)."""
    from dataclasses import replace
    from .checks import formatting, guidelines, lessonplan
    from .docx_model import parse_docx
    from .pptx_model import parse_pptx
    try:
        info = (parse_docx if doc.ext == "docx" else parse_pptx)(path)
    except Exception:
        return {}
    proxy = replace(doc, rel=rel, abs=str(path))
    from . import structure
    orig = ctx._cache.get(doc.rel)
    if orig is not None and structure.of(orig) is not None:
        info.structure = structure.transfer(structure.of(orig), orig, info)   # the model read the original, not this copy
    had, old = rel in ctx._cache, ctx._cache.get(rel)
    ctx._cache[rel] = info                       # a check that loads this document sees the fixed copy
    try:
        table = formatting.DOCX_CHECKS if doc.ext == "docx" else formatting.PPTX_CHECKS
        out = [fn(ctx, proxy, info) for fn in table.values()]
        out += guidelines.doc_checks(ctx, proxy, info)
        if doc.ext == "docx" and doc.doc_type == "lesson_plan":
            out += [fn(ctx, proxy, info) for fn in lessonplan.DOC_CHECKS.values()]
    except Exception:
        return {}
    finally:
        if had:
            ctx._cache[rel] = old
        else:
            ctx._cache.pop(rel, None)
    return {f.code: f for f in out}


def remaining(findings, ext, after, extra):
    """The findings to write into the fixed copy.

    A Fail or needs-reviewer result on a rule the fixer works on is replaced by what the same
    check says about the fixed copy: left out if the copy passes (or the rule no longer applies),
    otherwise the copy's own result, so the comment describes the file the reader opens. LP10 and
    FG8 are worked out again from the copy's results with the engine's own rule. Everything else
    (passes, rules the fixer does not touch, a copy that could not be re-checked) is kept as it was.
    """
    from .engine import DERIVED, derived_finding
    handled = DOCX_FIXED if ext == "docx" else PPTX_FIXED
    out = []
    for f in findings:
        g = None
        if f.status not in (FAIL, REVIEW) or not after:
            g = f
        elif ext == "docx" and f.code in ("PQ4", "DB4"):     # checked across the package, so not re-run here
            if not (f.status == FAIL and extra.get("ticked", 0) > 0 and "check-mark" in f.message):
                g = f
        elif f.code in DERIVED:
            g = derived_finding(f.code, [after[c] for c in DERIVED[f.code] if c in after]) or f
        elif f.code in handled and f.code in after:
            g = after[f.code]
        else:
            g = f
        if g is None or (g is not f and g.status in (PASS, NA)):
            continue
        if g is not f:
            g.doc, g.lesson = f.doc, f.lesson
        out.append(g)
    return out


TYPE_NAMES = {"lesson_plan": "Lesson-Plan", "chapter_exam": "Assessment", "facilitator_guide": "Facilitator-Guide",
              "video_storyboard": "Storyboard", "worksheet": "Chapter-Exam", "pop_quiz": "Pop-Quiz", "data_bank": "Data-Bank",
              "specification": "Document-of-Specifications"}
TICKS = re.compile(r"^\s*[✅✔☑✓]️?\s*")


# A version is 'v' straight before the number (v2, v1.0; not 'Class V 2') or the word Version / ورژن before it.
# The major number has at most two digits, so a year (2024) is never read as a version.
_VNUM = r"(\d{1,2}(?:\.\d{1,3}){0,3})(?![\d.]*\d)"
VERSION_MARKED = re.compile(r"(?:(?<![A-Za-z])v|\bversion\s*[:#]?\s*|\bver\.\s*|ورژن\s*[:#]?\s*)" + _VNUM, re.I)
# In the text only a line that is a version label counts ('Version 1.2', 'ورژن: 2'), never prose that
# mentions a version ('read the Urdu version 2 of the story').
VERSION_IN_TEXT = re.compile(r"^\s*(?:version|ver\.|ورژن)\s*[:#]?\s*" + _VNUM + r"\s*$", re.I)
# The footer the previous tool wrote ('Name | v0.1 | 2026-10-06 | Page'); its v0.1 was made up, not the document's.
OLD_TOOL_FOOTER = re.compile(r"\|\s*v0\.1\s*\|\s*\d{4}-\d{2}-\d{2}\s*\|\s*(?:Page|$)", re.I)


def _last(rx, text):
    found = rx.findall(to_western_digits(text or ""))
    return found[-1] if found else None


def find_version(doc, info):
    """The document's own version and where it was found, or (None, None).

    Looked for anywhere in the file name (Lesson-Plan-v2-Final, LessonPlan v1.0 (1)), then the
    footer, then the file's properties, then an explicit 'Version 1.2' line near the top. A
    version is never made up: one the tool cannot find is left for the writer to add (WE3 stays
    a Fail with its comment)."""
    stem = doc.rel.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    v = _last(VERSION_MARKED, stem)
    if v:
        return v, "the file name"
    if info is None:
        return None, None
    footer = info.footer_text if doc.ext == "docx" else " ".join(t for _, _, t in info.footers)
    if OLD_TOOL_FOOTER.search(footer):
        footer = OLD_TOOL_FOOTER.sub("|", footer)
    v = _last(VERSION_MARKED, footer)
    if v:
        return v, "the footer"
    m = re.fullmatch(r"\s*v?\s*" + _VNUM + r"\s*", to_western_digits(getattr(info, "core_version", "") or ""), re.I)
    if m:
        return m.group(1), "the file properties"
    if doc.ext == "docx":
        for p in [p for p in info.paras if p.text.strip() and not p.in_table][:10]:
            v = _last(VERSION_IN_TEXT, p.text.strip()[:40])
            if v:
                return v, "the document text"
    return None, None


def new_name(doc, grade_subject, version):
    """File name in the guideline pattern [Type]-[Identifier]-[Chapter/Topic]-v[Version].
    With no version (None) the name has no -v part, so WE1 and WE3 still ask for one."""
    t = TYPE_NAMES.get(doc.doc_type)
    if t is None:
        return None
    if doc.doc_type in ("pop_quiz", "data_bank", "specification"):
        ident = grade_subject or "Subject"
    elif doc.scope == "chapter" or doc.lesson is None:
        ident = f"Chapter-{doc.chapter}" if doc.chapter is not None else "Chapter"
    else:
        ident = f"Lesson-{doc.lesson}{doc.variant}-Chapter-{doc.chapter}"
    return f"{t}-{ident}" + (f"-v{version}" if version else "") + f".{doc.ext}"


# ------------------------------------------------------------------------- Word
def _para_elements(body):
    """w:p elements in the same order as docx_model's paragraphs, so ParaInfo.idx finds its element."""
    out = []

    def walk(container):
        for child in container:
            tag = etree.QName(child).localname
            if tag == "p":
                out.append(child)
            elif tag == "tbl":
                for tr in child.findall(qn("w:tr")):
                    for tc in tr.findall(qn("w:tc")):
                        out.extend(tc.iter(qn("w:p")))
            elif tag == "sdt":
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    walk(content)
    walk(body)
    return out


def _rpr(r):
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        r.insert(0, rpr)
    return rpr


def _set(rpr, tag, attrs):
    el = rpr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        rpr.append(el)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


def _drop(parent, tag):
    for el in parent.findall(qn(tag)):
        parent.remove(el)


def _run_text(r):
    return "".join(t.text or "" for t in r.iter(qn("w:t")))


def _set_size(rpr, pt):
    _set(rpr, "w:sz", {"w:val": str(int(pt * 2))})
    _set(rpr, "w:szCs", {"w:val": str(int(pt * 2))})


def _set_bold(rpr, on):
    for tag in ("w:b", "w:bCs"):
        _drop(rpr, tag)
        if on:
            _set(rpr, tag, {})


def _ensure_heading_style(doc, level):
    name = f"Heading {level}"
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
    ppr = st.element.get_or_add_pPr()
    _set(ppr, "w:outlineLvl", {"w:val": str(level - 1)})
    rpr = st.element.get_or_add_rPr()
    _drop(rpr, "w:color")                 # Word's built-in headings are blue; the no-colour rule wants black
    return st


def _split_paragraph(el, pos):
    """Move everything after character `pos` of paragraph `el` into a new paragraph right after it (same
    paragraph properties, without the style), dropping the ':' and spaces at the cut. Returns the new paragraph."""
    import copy
    new = OxmlElement("w:p")
    ppr = el.find(qn("w:pPr"))
    if ppr is not None:
        nppr = copy.deepcopy(ppr)
        for tag in ("w:pStyle", "w:outlineLvl", "w:numPr"):
            _drop(nppr, tag)
        new.append(nppr)
    seen, moving = 0, False
    for r in list(el):
        if r.tag != qn("w:r"):
            if moving and r.tag != qn("w:pPr"):
                new.append(r)
            continue
        ts = r.findall(qn("w:t"))
        txt = "".join(t.text or "" for t in ts)
        if moving:
            new.append(r)
        elif seen + len(txt) > pos:
            cut = pos - seen
            tail = copy.deepcopy(r)
            for t in ts[1:]:
                r.remove(t)
            for t in tail.findall(qn("w:t"))[1:]:
                tail.remove(t)
            if ts:
                ts[0].text = txt[:cut]
                tt = tail.find(qn("w:t"))
                tt.text = txt[cut:]
                tt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            new.append(tail)
            moving = True
        seen += len(txt)
    # the colon and the spaces around the cut go: the heading ends at its label, the text starts at its first word
    t = next((t for t in new.iter(qn("w:t")) if (t.text or "").strip()), None)
    if t is not None:
        t.text = re.sub(r"^\s*[:：]\s*", "", t.text)
    el.addnext(new)
    return new


REVIEW_AUTHOR = "Course Review"       # the author annotate.py writes its comments as
REVIEW_HIGHLIGHTS = {"red", "cyan", "green", "magenta"}   # red, turquoise, bright green, pink in annotate.py


def drop_review_comments(d):
    """Remove the comments an earlier Course Review run left in a re-uploaded copy (the creator's and reviewers'
    own comments stay). Returns how many were removed."""
    try:
        part = d.part._comments_part
    except Exception:
        return 0
    if part is None:
        return 0
    root = part.element
    ids = set()
    for c in list(root.findall(qn("w:comment"))):
        if c.get(qn("w:author")) == REVIEW_AUTHOR:
            ids.add(c.get(qn("w:id")))
            root.remove(c)
    if not ids:
        return 0
    body = d.element.body
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(body.iter(qn(tag))):
            if el.get(qn("w:id")) in ids:
                el.getparent().remove(el)
    for ref in list(body.iter(qn("w:commentReference"))):
        if ref.get(qn("w:id")) in ids:
            run = ref.getparent()
            run.remove(ref)
            if not [x for x in run if x.tag != qn("w:rPr")]:
                run.getparent().remove(run)
    return len(ids)


BULLET_WHY = "numbering is only for the major headings; lists inside a section use bullets (Writing & Editing Guidelines)"


def fix_docx(src, dst, ctx, doc, info):
    """Write a fixed copy of one .docx; returns (lines for the summary comment, what was done, e.g. ticked)."""
    from .checks.formatting import appendix_heading, doc_headings, grade_sizes, leading_number, misstyled_body
    from .checks.guidelines import lp_ranges
    d = Document(src)
    body = d.element.body
    els = _para_elements(body)
    if len(els) != len(info.paras):        # structure the parser read differently: do not risk a wrong mapping
        return [], {}
    lines, hr = [], ctx.profile.get("house_rules", {})
    old_notes = drop_review_comments(d)
    if old_notes:
        lines.append(f"{old_notes} comment(s) from an earlier Course Review of this file removed; this review replaces them")
    urdu_font = (hr.get("urdu_fonts") or ["Noto Nastaliq Urdu"])[0]
    eng_font = (hr.get("english_fonts") or ["Poppins"])[0]
    grade, sz, _ = grade_sizes(ctx)
    keep_yellow = doc.doc_type in ("pop_quiz", "data_bank")

    # page setup
    for s in d.sections:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    lines.append("Page set to A4 portrait with 1-inch margins")

    # changes to particular lines, each highlighted in the copy with what it was, why, and what it is now
    changes = []                          # (element, before, why, now or None = the line as it now reads)

    # paragraph roles
    first = next((p for p in info.paras if p.text.strip() and not p.in_table), None)
    heads = doc_headings(ctx, info)
    from .structure import of
    st = structure_of = of(info)
    if st is not None and st.kind == "lesson_plan":
        # a heading in dispute is left as it is, with a comment; a section label with its text run on is split
        heads = [p for p in info.paras if p.idx in st.agreed_headings or p.idx in st.run_in]
    head_ids = {p.idx for p in heads}
    levels, numbers = {}, {}
    secs = lp_ranges(ctx, info) if doc.doc_type == "lesson_plan" else {}
    number_heads = doc.doc_type != "lesson_plan" or len(secs) >= 4
    if doc.doc_type == "lesson_plan" and number_heads:
        starts = sorted((r[0], k) for k, r in secs.items())
        top = {i for i, _ in starts}
        n1 = n2 = 0
        for p in heads:
            if appendix_heading(ctx, p.text):
                levels[p.idx] = 1                 # a top-level heading the template leaves unnumbered
            elif p.idx in top:
                n1, n2 = n1 + 1, 0
                levels[p.idx], numbers[p.idx] = 1, f"{n1}"
            elif n1:
                n2 += 1
                levels[p.idx], numbers[p.idx] = 2, f"{n1}.{n2}"
            else:
                n1 += 1
                levels[p.idx], numbers[p.idx] = 1, f"{n1}"
    else:
        n = 0
        for p in heads:
            levels[p.idx] = 1
            if not appendix_heading(ctx, p.text):
                n += 1
                numbers[p.idx] = f"{n}"

    digits = recoloured = flipped = quoted = 0
    for p, el in zip(info.paras, els):
        ppr = el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            el.insert(0, ppr)
        # spacing: 1.15 lines, 8 pt after
        _set(ppr, "w:spacing", {"w:line": "276", "w:lineRule": "auto", "w:after": "160"})
        # direction follows the paragraph's script: RTL for Urdu, LTR for English
        ar = sum(len(r.text) for r in p.runs if r.script == "arabic")
        la = sum(len(r.text) for r in p.runs if r.script == "latin")
        if ar or la:
            if (ar >= la) != p.bidi:
                flipped += 1
            _drop(ppr, "w:bidi")
            if ar >= la:
                b = OxmlElement("w:bidi")
                ps = ppr.find(qn("w:pStyle"))
                (ps.addnext(b) if ps is not None else ppr.insert(0, b))
        if p.idx in head_ids:
            role = "head"
        elif p is first or p.is_title:
            role = "title"
        elif p.in_table:
            role = "theader" if p.row_idx == 0 else "table"
        else:
            role = "body"
        size = {"title": sz["title"], "head": sz["section_heading"], "theader": sz["table_header"],
                "table": sz["table_body"], "body": sz["body"]}[role]
        scripture = is_arabic_scripture(p.text)
        for r in el.iter(qn("w:r")):
            text = _run_text(r)
            if not text:
                continue
            if scripture or is_arabic_scripture(text):
                quoted += 1
                # Quran / hadith / dua text is quoted exactly: no font, size, digit or colour change. Only a review
                # highlight an earlier Course Review run put on it goes, which leaves every letter as it was.
                hl = r.find(qn("w:rPr") + "/" + qn("w:highlight"))
                if hl is not None and hl.get(qn("w:val")) in REVIEW_HIGHLIGHTS:
                    hl.getparent().remove(hl)
                continue
            rpr = _rpr(r)
            ar, la = script_counts(text)
            if ar or la:
                fonts = _set(rpr, "w:rFonts", {})
                if ar >= la:
                    for k in ("w:cs", "w:ascii", "w:hAnsi"):
                        fonts.set(qn(k), urdu_font)
                else:
                    for k in ("w:ascii", "w:hAnsi"):
                        fonts.set(qn(k), eng_font)
                for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
                    fonts.attrib.pop(qn(k), None)
            _set_size(rpr, size)
            if role in ("title", "head", "theader"):
                _set_bold(rpr, True)
            if rpr.find(qn("w:color")) is not None:
                recoloured += 1
            _drop(rpr, "w:color")
            hl = rpr.find(qn("w:highlight"))
            if hl is not None and not (keep_yellow and hl.get(qn("w:val")) == "yellow"):
                rpr.remove(hl)
            shd = rpr.find(qn("w:shd"))
            if shd is not None and not (keep_yellow and (shd.get(qn("w:fill")) or "").upper() in ("FFFF00", "FFFF")):
                rpr.remove(shd)
            for t in r.iter(qn("w:t")):
                if t.text and EASTERN_DIGITS.search(t.text):
                    # verse numbers in ﴿ ﴾ stay as written; other Urdu digits become Western
                    parts = VERSE_MARK.split(t.text)
                    marks_ = VERSE_MARK.findall(t.text)
                    new = "".join(to_western_digits(x) + (marks_[i] if i < len(marks_) else "") for i, x in enumerate(parts))
                    digits += sum(1 for a, b in zip(t.text, new) if a != b)
                    t.text = new
        # paragraph shading
        _drop(ppr, "w:shd")
    if digits:
        lines.append(f"{digits} Urdu digits changed to Western (1, 2, 3)")
    if quoted:
        lines.append("Quran, hadith and dua text (fully vowelled Arabic) was left exactly as written")
    if flipped:
        lines.append(f"Text direction corrected on {flipped} paragraph(s) (right-to-left for Urdu, left-to-right for English)")
    lines.append(f"Line spacing 1.15 and 8 pt after every paragraph; fonts {urdu_font} / {eng_font}; Grade {grade} sizes "
                 f"(title {sz['title']}, headings {sz['section_heading']}, body {sz['body']}, tables {sz['table_body']})")
    if recoloured:
        lines.append(f"Colour removed from {recoloured} text run(s)")

    # table cell shading
    shaded = 0
    for tc in body.iter(qn("w:tc")):
        tcpr = tc.find(qn("w:tcPr"))
        if tcpr is not None and tcpr.find(qn("w:shd")) is not None:
            fill = (tcpr.find(qn("w:shd")).get(qn("w:fill")) or "").upper()
            if fill not in ("", "AUTO", "FFFFFF") and not (keep_yellow and fill in ("FFFF00",)):
                _drop(tcpr, "w:shd")
                shaded += 1
    if shaded:
        lines.append(f"Colour removed from {shaded} table cell(s)")

    # running text that was set in a Heading style goes back to Normal
    restyled = 0
    for p in info.paras:
        if misstyled_body(p, ctx) and not is_arabic_scripture(p.text) and p.idx not in head_ids:
            ppr = els[p.idx].find(qn("w:pPr"))
            if ppr is not None:
                _drop(ppr, "w:pStyle")
                _drop(ppr, "w:outlineLvl")
                restyled += 1
                changes.append((els[p.idx], f"running text set in the '{p.style}' style",
                                "Heading styles are only for headings; text in a Heading style fills the navigation pane and the "
                                "Table of Contents (Writing & Editing Guidelines)", "the Normal style, as body text"))
    if restyled:
        lines.append(f"{restyled} paragraph(s) of running text moved from a Heading style to Normal")

    # headings: built-in Heading styles and decimal numbers
    if heads:
        for p in heads:
            el = els[p.idx]
            before_style, before_text = p.style, p.text.strip()
            cut = run_in_split(ctx, p.text) if not is_arabic_scripture(p.text) else None
            if cut is not None:
                body_p = _split_paragraph(el, len(p.text) - len(p.text.lstrip()) + cut)
                changes.append((body_p, f"this text ran on after the heading in the same paragraph ('{before_text[:50]}…')",
                                "a heading stands on its own line and the text under it is body text (Writing & Editing Guidelines)",
                                "its own paragraph under the heading, word for word"))
            st = _ensure_heading_style(d, levels[p.idx])
            ppr = el.find(qn("w:pPr"))
            _set(ppr, "w:pStyle", {"w:val": st.style_id})
            if p.list_kind:                       # a list item turned heading keeps no list number of its own
                _drop(ppr, "w:numPr")
                off = _set(ppr, "w:numPr", {})
                _set(off, "w:ilvl", {"w:val": "0"})
                _set(off, "w:numId", {"w:val": "0"})
            ts = list(el.iter(qn("w:t")))
            if ts and number_heads and p.idx in numbers and not own_number(ctx, p.text):
                old = leading_number(p.text)
                first_t = next((t for t in ts if (t.text or "").strip()), ts[0])
                text = first_t.text or ""
                if old is not None:
                    text = re.sub(r"^\s*[\d٠-٩۰-۹]+(?:[.][\d٠-٩۰-۹]+)*[.)]?\s*", "", text)
                first_t.text = f"{numbers[p.idx]} {text.lstrip()}"
                first_t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            numbered_now = number_heads and p.idx in numbers and not own_number(ctx, p.text)
            if not p.heading_level or numbered_now and leading_number(before_text) != tuple(int(x) for x in numbers[p.idx].split(".")):
                what = [] if p.heading_level else [f"the '{before_style}' style"]
                if numbered_now:
                    what.append("no decimal number" if leading_number(before_text) is None else "a different number")
                changes.append((el, f"'{before_text[:60]}' with " + " and ".join(what),
                                "headings use built-in Heading styles" + (" and decimal numbers 1, 1.1, 1.1.1" if numbered_now else "")
                                + " (Writing & Editing Guidelines)", f"Heading {levels[p.idx]}" + (" with its number" if numbered_now else "")))
        lines.append(f"{len(heads)} heading(s) given Heading styles" + (" and decimal numbers (1, 1.1)" if number_heads else
                     "; not numbered, because the document does not follow the five Lesson Plan sections (see the comments)"))

    # numbered list items -> bullets (numbering is only for the major headings)
    bulleted = 0
    disputed = structure_of.disputed_headings if structure_of is not None and structure_of.kind == "lesson_plan" else set()
    for p in info.paras:
        if p.text.strip() and not p.in_table and p.idx not in head_ids and p.idx not in disputed and p.list_kind == "decimal":
            el = els[p.idx]
            ppr = el.find(qn("w:pPr"))
            _drop(ppr, "w:numPr")
            off = _set(ppr, "w:numPr", {})               # numId 0 also switches off numbering that comes from a list style
            _set(off, "w:ilvl", {"w:val": "0"})
            _set(off, "w:numId", {"w:val": "0"})
            t = next((t for t in el.iter(qn("w:t")) if (t.text or "").strip()), None)
            if t is not None:
                t.text = "• " + t.text.lstrip()
            bulleted += 1
            changes.append((el, f"item {p.num_label or ''} in a numbered list".replace("  ", " "), BULLET_WHY, "a bullet"))
    # SLO items written as plain paragraphs -> bullets (LP5)
    if doc.doc_type == "lesson_plan":
        from .checks.lessonplan import lp5
        plain = {m["text"] for m in lp5(ctx, doc, info).marks}
        for p in info.paras:
            if p.text.strip() in plain and p.list_kind != "decimal":
                t = next((t for t in els[p.idx].iter(qn("w:t")) if (t.text or "").strip()), None)
                if t is not None:
                    typed = re.match(r"^\s*\(?[\d٠-٩۰-۹]+[.)]\s*", t.text)
                    t.text = "• " + (t.text[typed.end():] if typed else t.text.lstrip())   # a typed '1.' is the list number, not the text
                    bulleted += 1
                    changes.append((els[p.idx], f"SLO '{p.text.strip()[:50]}' " + ("numbered by hand" if typed else "written as a plain paragraph"),
                                    "SLOs are listed as bullets (Lesson Plan Guidelines)", "a bullet"))
    # Key Takeaways written as plain paragraphs -> bullets (LPG2)
    if doc.doc_type == "lesson_plan":
        from .checks.guidelines import _is_bullet, takeaway_paras
        from .checks.lessonplan import is_bullet_text
        items = [p for p in (takeaway_paras(ctx, info) or []) if p.idx not in head_ids and p.text.strip()]
        if items and not any(_is_bullet(p) for p in items):
            for p in items:
                if p.list_kind == "decimal" or is_bullet_text(p.text):
                    continue                              # handled above, or already a bullet
                t = next((t for t in els[p.idx].iter(qn("w:t")) if (t.text or "").strip()), None)
                if t is None:
                    continue
                typed = re.match(r"^\s*\(?[\d٠-٩۰-۹]+[.)]\s*", t.text)
                t.text = "• " + (t.text[typed.end():] if typed else t.text.lstrip())
                bulleted += 1
                changes.append((els[p.idx], f"Key Takeaway '{p.text.strip()[:50]}' " + ("numbered by hand" if typed else "written as a plain paragraph"),
                                "Key Takeaways recap the lesson in bullets (Lesson Plan Guidelines)", "a bullet"))
    if bulleted:
        lines.append(f"{bulleted} numbered or plain list item(s) turned into bullets")

    # check marks on correct answers -> yellow highlight
    ticked = 0
    if keep_yellow:
        for el in els:
            ts = [t for t in el.iter(qn("w:t")) if t.text]
            if ts and TICKS.match("".join(t.text for t in ts)):
                for t in ts:
                    if TICKS.match(t.text):
                        t.text = TICKS.sub("", t.text, count=1)
                        break
                    if t.text.strip():
                        break
                for r in el.iter(qn("w:r")):
                    if _run_text(r).strip():
                        _set(_rpr(r), "w:highlight", {"w:val": "yellow"})
                ticked += 1
                changes.append((el, "the correct answer was marked with a check mark",
                                "the template marks the correct answer with a yellow highlight", "highlighted yellow, check mark removed"))
        if ticked:
            lines.append(f"{ticked} correct answer(s) marked with a check mark now highlighted yellow instead")

    # table captions above tables
    captions = 0
    word = "جدول" if (info.language == "ur") else "Table"
    cap_rx = re.compile(r"^\s*(?:Table|جدول)\s*\d+\s*[:：.\-–]", re.I)
    for n, t in enumerate(info.tables, 1):
        prev = info.paras[t.prev_para] if t.prev_para is not None else None
        if prev is not None and cap_rx.match(to_western_digits(normalize(prev.text))):
            continue
        tbl = _nth_table(body, t.idx)
        if tbl is None:
            continue
        head = next((info.paras[i].text.strip() for i in t.para_idx if info.paras[i].text.strip()), "")
        i = t.prev_para
        while i is not None and i >= 0 and not info.paras[i].text.strip() and not info.paras[i].in_table:
            i -= 1
        if i is not None and i >= 0 and i in head_ids:
            head = re.sub(r"^\s*[\d٠-٩۰-۹]+(?:\.[\d٠-٩۰-۹]+)*[.)]?\s*", "", info.paras[i].text.strip())   # the heading it sits under names it
        cap = _new_para(f"{word} {n}: {head[:50]}".rstrip(": "), info.language, ctx, sz["caption"], italic=True)
        tbl.addprevious(cap)
        captions += 1
        changes.append((cap, "the table had no caption", "every table has a numbered caption directly above it (Writing & Editing "
                        "Guidelines); the title is taken from the heading above or the first cell, so check it", None))
    if captions:
        lines.append(f"{captions} numbered table caption(s) added above the tables (titles taken from the heading above or the table's first cell; check them)")

    # Table of Contents for a Lesson Plan longer than two pages
    if doc.doc_type == "lesson_plan" and heads and not info.has_toc:
        from .checks.formatting import estimate_pages
        if (info.pages or estimate_pages(info)) > 2 and first is not None:
            toc = OxmlElement("w:p")
            r1 = OxmlElement("w:r")
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            r1.append(fc)
            r2 = OxmlElement("w:r")
            it = OxmlElement("w:instrText")
            it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            it.text = 'TOC \\o "1-3" \\h \\z \\u'
            r2.append(it)
            r3 = OxmlElement("w:r")
            fc3 = OxmlElement("w:fldChar")
            fc3.set(qn("w:fldCharType"), "separate")
            r3.append(fc3)
            r4 = _new_para("فہرستِ مضامین" if info.language == "ur" else "Table of Contents", info.language, ctx, sz["body"])
            r4 = r4.find(qn("w:r"))
            r5 = OxmlElement("w:r")
            fc5 = OxmlElement("w:fldChar")
            fc5.set(qn("w:fldCharType"), "end")
            r5.append(fc5)
            if info.language == "ur":
                _ppr_of(toc).append(OxmlElement("w:bidi"))
            _set(_ppr_of(toc), "w:spacing", {"w:line": "276", "w:lineRule": "auto", "w:after": "160"})
            for x in (r1, r2, r3, r4, r5):
                toc.append(x)
            els[first.idx].addnext(toc)
            changes.append((toc, "no Table of Contents", "a Lesson Plan longer than two pages starts with one (Writing & Editing "
                            "Guidelines)", "a Table of Contents; Word fills it in when the file is opened (choose Yes when asked to update fields)"))
            settings = d.settings.element
            _set(settings, "w:updateFields", {"w:val": "true"})
            lines.append("Table of Contents inserted after the title (Word fills it in when the file is opened)")
    d.save(dst)
    done = [{"text": "".join(t.text or "" for t in el.iter(qn("w:t"))).strip(), "before": before, "why": why,
             "now": now or "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()} for el, before, why, now in changes]
    return lines, {"ticked": ticked, "toc": any("Table of Contents" in l for l in lines), "numbered": number_heads,
                   "changes": [c for c in done if c["text"]]}


def _ppr_of(p):
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    return ppr


def _new_para(text, lang, ctx, size, italic=False):
    """A paragraph already in house style: right font for its script, size, 1.15 spacing, 8 pt after."""
    hr = ctx.profile.get("house_rules", {})
    ar, la = script_counts(text)
    font = (hr.get("urdu_fonts") or ["Noto Nastaliq Urdu"])[0] if ar >= la else (hr.get("english_fonts") or ["Poppins"])[0]
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    if ar >= la and ar:                    # direction follows the text's own script
        ppr.append(OxmlElement("w:bidi"))
    _set(ppr, "w:spacing", {"w:line": "276", "w:lineRule": "auto", "w:after": "160"})
    p.append(ppr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(k), font)
    rpr.append(fonts)
    if italic:
        _set(rpr, "w:i", {})
        _set(rpr, "w:iCs", {})
    _set_size(rpr, size)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    p.append(r)
    return p


def _nth_table(body, idx):
    n = -1

    def walk(container):
        nonlocal n
        for child in container:
            tag = etree.QName(child).localname
            if tag == "tbl":
                n += 1
                if n == idx:
                    return child
            elif tag == "sdt":
                c = child.find(qn("w:sdtContent"))
                if c is not None:
                    hit = walk(c)
                    if hit is not None:
                        return hit
        return None
    return walk(body)


def add_footer(path, name, version):
    """Footer with document name, version (left out when there is none), date and page number on every section."""
    d = Document(path)
    for s in d.sections:
        f = s.footer
        f.is_linked_to_previous = False
        for p in list(f.paragraphs):
            p._p.getparent().remove(p._p)
        p = f.add_paragraph()
        p.add_run(f"{name}" + (f" | v{version}" if version else "") + f" | {date.today().isoformat()} | Page ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
        for run in p.runs:
            run.font.size = Pt(9)
    d.save(path)


# --------------------------------------------------------------------- PowerPoint
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _slide_footers(prs, footer, font):
    """A footer text box along the bottom of every slide (the deck has no footer placeholders to fill)."""
    from pptx.util import Emu, Pt as PPt
    w, h = prs.slide_width, prs.slide_height
    for slide in prs.slides:
        if any(sh.name == "Footer" for sh in slide.shapes):
            continue
        box = slide.shapes.add_textbox(Emu(int(w * 0.03)), Emu(int(h * 0.93)), Emu(int(w * 0.94)), Emu(int(h * 0.05)))
        box.name = "Footer"
        p = box.text_frame.paragraphs[0]
        ln = etree.SubElement(p._p.get_or_add_pPr(), f"{{{A}}}lnSpc")
        etree.SubElement(ln, f"{{{A}}}spcPct", val="150000")
        r = p.add_run()
        r.text = footer + " | "
        r.font.size, r.font.name = PPt(9), font
        fld = etree.SubElement(p._p, f"{{{A}}}fld", id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
        rpr = etree.SubElement(fld, f"{{{A}}}rPr", lang="en-US", sz="900")
        etree.SubElement(rpr, f"{{{A}}}latin", typeface=font)
        etree.SubElement(fld, f"{{{A}}}t").text = "‹#›"
        end = p._p.find(f"{{{A}}}endParaRPr")
        if end is not None:
            p._p.remove(end)
            p._p.append(end)


def fix_pptx(src, dst, ctx, footer=None):
    from pptx import Presentation
    from .checks.common import is_grey
    hr = ctx.profile.get("house_rules", {})
    urdu_font = (hr.get("urdu_fonts") or ["Noto Nastaliq Urdu"])[0]
    eng_font = (hr.get("english_fonts") or ["Poppins"])[0]
    prs = Presentation(src)
    digits = paras = recoloured = 0
    for slide in prs.slides:
        for el in slide.shapes._spTree.iter(f"{{{A}}}p"):
            paras += 1
            ppr = el.find(f"{{{A}}}pPr")
            if ppr is None:
                ppr = etree.SubElement(el, f"{{{A}}}pPr")
                el.remove(ppr)
                el.insert(0, ppr)
            for old in ppr.findall(f"{{{A}}}lnSpc"):
                ppr.remove(old)
            ln = etree.Element(f"{{{A}}}lnSpc")
            etree.SubElement(ln, f"{{{A}}}spcPct", val="150000")
            ppr.insert(0, ln)
            for r in el.findall(f"{{{A}}}r"):
                t = r.find(f"{{{A}}}t")
                if t is None or not t.text:
                    continue
                if EASTERN_DIGITS.search(t.text):
                    digits += len(EASTERN_DIGITS.findall(t.text))
                    t.text = to_western_digits(t.text)
                rpr = r.find(f"{{{A}}}rPr")
                if rpr is None:
                    rpr = etree.Element(f"{{{A}}}rPr")
                    r.insert(0, rpr)
                ar, la = script_counts(t.text)
                fill = rpr.find(f"{{{A}}}solidFill")
                if fill is not None and len(fill):
                    c = fill[0]
                    val = c.get("val") or ""
                    grey = c.tag.endswith("srgbClr") and is_grey(val, ctx.th["colour_grey_tolerance"])
                    neutral = c.tag.endswith("schemeClr") and val in ("tx1", "dk1", "bg1", "lt1")
                    if not grey and not neutral:
                        rpr.remove(fill)
                        recoloured += 1
                tag = "cs" if ar >= la and ar else "latin"
                if ar or la:
                    font = urdu_font if tag == "cs" else eng_font
                    e = rpr.find(f"{{{A}}}{tag}")
                    if e is None:
                        e = etree.SubElement(rpr, f"{{{A}}}{tag}")
                    e.set("typeface", font)
                    # schema order inside rPr: fills/effects, highlight, then latin/ea/cs
                    for name in ("latin", "ea", "cs", "sym", "hlinkClick", "rtl", "extLst"):
                        for x in rpr.findall(f"{{{A}}}{name}"):
                            rpr.remove(x)
                            rpr.append(x)
        for tc in slide.shapes._spTree.iter(f"{{{A}}}tcPr"):
            for sf in tc.findall(f"{{{A}}}solidFill"):
                c = sf[0] if len(sf) else None
                if c is not None and c.tag.endswith("srgbClr") and not is_grey(c.get("val"), ctx.th["colour_grey_tolerance"]):
                    tc.remove(sf)
                    recoloured += 1
    if footer:
        _slide_footers(prs, footer, eng_font)
    prs.save(dst)
    lines = [f"Line spacing 1.5 on all {paras} slide paragraph(s); fonts {urdu_font} / {eng_font}"]
    if footer:
        lines.append("Footer added to every slide: document name, version, date, slide number")
    if digits:
        lines.append(f"{digits} Urdu digits changed to Western (1, 2, 3)")
    if recoloured:
        lines.append(f"Colour removed from {recoloured} text run(s) or table cell(s)")
    return lines
