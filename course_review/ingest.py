"""Safe unzip and classification of package files.

Classification is driven by the profile (patterns), not by hard-coded names, and
file names are only a hint: anything that cannot be placed is returned in
Package.unclassified so the report can show it.
"""
import os
import re
import zipfile
from pathlib import Path

import yaml

from .models import DocRef, LessonKey, Package

DEFAULT_PROFILE = Path(__file__).with_name("profile.default.yaml")


def _merge(base, over):
    """Overlay `over` on `base`: dicts merge key by key, other values (lists included) are replaced,
    except extra_question_patterns, which accumulate."""
    out = dict(base)
    for k, v in over.items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            out[k] = _merge(base[k], v)
        elif k == "extra_question_patterns" and isinstance(base.get(k), list):
            out[k] = list(base[k]) + [x for x in v if x not in base[k]]
        else:
            out[k] = v
    return out


def load_profile(path=None) -> dict:
    """The built-in profile, with the file at `path` (if any) laid over it. A profile file only
    needs the settings it changes, for example a few learned question patterns."""
    with open(DEFAULT_PROFILE, encoding="utf8") as f:
        profile = yaml.safe_load(f)
    if path:
        with open(path, encoding="utf8") as f:
            profile = _merge(profile, yaml.safe_load(f) or {})
    return profile


class UnsafeArchive(Exception):
    pass


def safe_extract(zip_path, dest, max_files=20000, max_total=4 << 30, max_ratio=1000):
    """Extract zip_path into dest, refusing path traversal and zip bombs."""
    dest = Path(dest).resolve()
    dest.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as z:
        infos = z.infolist()
        if len(infos) > max_files:
            raise UnsafeArchive(f"archive has {len(infos)} entries (limit {max_files})")
        total = sum(i.file_size for i in infos)
        if total > max_total:
            raise UnsafeArchive(f"archive expands to {total} bytes (limit {max_total})")
        for i in infos:
            if i.compress_size and i.file_size / max(i.compress_size, 1) > max_ratio and i.file_size > (50 << 20):
                raise UnsafeArchive(f"suspicious compression ratio on {i.filename}")
            target = (dest / i.filename).resolve()
            if dest != target and dest not in target.parents:
                raise UnsafeArchive(f"entry escapes the extraction folder: {i.filename}")
        z.extractall(dest)
    return dest


def find_root(dest) -> Path:
    """If the archive has a single top-level folder, that folder is the package root."""
    entries = [p for p in Path(dest).iterdir() if not p.name.startswith(("__MACOSX", "."))]
    if len(entries) == 1 and entries[0].is_dir():
        return entries[0]
    return Path(dest)


_VERSION = re.compile(r"v(\d+(?:\.\d+)*)", re.I)


def _version_key(name: str):
    m = _VERSION.search(name)
    return tuple(int(x) for x in m.group(1).split(".")) if m else ()


def _letter_for_part(n: int) -> str:
    return chr(ord("A") + n - 1) if 1 <= n <= 26 else ""


def _parse_location(parts, profile):
    """parts: path components below the root, last one is the file name."""
    pp = profile["path_patterns"]
    chapter = lesson = None
    variant = ""
    folder_label = ""
    folders, fname = parts[:-1], parts[-1]
    for p in folders:
        m = re.search(pp["chapter"], p)
        if m and chapter is None:
            chapter = int(m.group(1))
    for p in folders:
        m = re.search(pp["lesson"], p)
        if m and lesson is None:
            lesson = int(m.group(1))
            folder_label = p
            part = re.search(r"(?i)part[-_ ]*(\d+)(?:[-_ ]*([A-Za-z])(?=[-_ .]))?", p)
            if part:
                variant = (part.group(2) or _letter_for_part(int(part.group(1)))).upper()
    stem = os.path.splitext(fname)[0]
    if chapter is None:
        m = re.search(pp["chapter"], stem)
        if m:
            chapter = int(m.group(1))
    if lesson is None:
        m = re.search(pp["lesson"], stem)
        if m:
            lesson = int(m.group(1))
    if lesson is not None and not variant:
        m = re.search(pp["lesson"] + r"[-_ ]*([A-Za-z])(?=[-_ .])", stem)
        if m:
            variant = m.group(2).upper()
    return chapter, lesson, variant, folder_label


# ------------------------------------------------------------- by content
GENERIC = "document"          # a Word or PowerPoint file whose kind cannot be told: it still gets the W&E checks and fixes


def convert_legacy(path):
    """An old Word 97-2003 .doc (or .ppt) converted to .docx (.pptx) next to it with LibreOffice, or None when
    LibreOffice is not installed or the conversion fails."""
    import shutil
    import subprocess
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe is None:
        for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
            if os.path.exists(cand):
                exe = cand
                break
    if exe is None:
        return None
    target = "docx" if path.suffix.lower() == ".doc" else "pptx"
    try:
        subprocess.run([exe, "--headless", "--convert-to", target, "--outdir", str(path.parent), str(path)],
                       capture_output=True, timeout=180, check=False)
    except (OSError, subprocess.SubprocessError):
        return None
    out = path.with_suffix("." + target)
    if not out.exists():
        return None
    try:
        path.unlink()                 # the extracted copy of the old file: the converted one replaces it in the output
    except OSError:
        pass
    return out


def _docx_text(path, limit=400):
    """(paragraph texts, table cell texts) of a .docx, read lightly for classification."""
    import docx
    d = docx.Document(str(path))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()][:limit]
    cells = []
    for t in d.tables[:20]:
        for row in t.rows[:60]:
            for c in row.cells:
                if c.text.strip():
                    cells.append(c.text.strip())
    return paras, cells


def _number_after(words, text):
    m = re.search(r"(?i)\b(?:" + "|".join(words) + r")[-_ :]*(\d{1,2})\b", text)
    return int(m.group(1)) if m else None


def classify_by_content(path, ext, profile):
    """(doc type, chapter, lesson, why) for a file whose name does not say what it is: its title lines, the
    labels it uses (the five Lesson Plan sections, 'Lesson Plan Location', feedback fields, storyboard columns)
    and its questions. Returns the generic type when nothing fits."""
    from .questions import question_regex
    from .textutil import normalize, starts_with_label
    if ext == "pptx":
        return "facilitator_guide", None, None, "a slide deck"
    try:
        paras, cells = _docx_text(path)
    except Exception:
        return None, None, None, "the file could not be opened as a Word document"
    title = " | ".join(paras[:8])
    chapter = _number_after(["chapter", "chpater", "unit", "باب"], title)
    lesson = _number_after(["lesson", "سبق"], title)
    # 1. the document names itself in its first lines ('Lesson Plan', 'Pop Quiz', 'Assessment' ...)
    for r in profile["doc_types"]:
        if "docx" in r.get("ext", ["docx"]) and r["type"] != "textbook" and any(re.search(r["pattern"], l, re.I) for l in paras[:6]):
            return r["type"], chapter, lesson, f"its title says so ('{paras[0][:50]}')"
    low = [normalize(x).lower() for x in paras + cells]
    vocab = profile["vocab"]
    # 2. the labels its template uses
    sections = sum(1 for sec in profile["lesson_plan_sections"] if any(starts_with_label(t, sec["labels"]) for t in paras))
    if sections >= 3:
        return "lesson_plan", chapter, lesson, f"it has {sections} of the five Lesson Plan sections"
    if any(normalize(l).lower() in t for l in vocab.get("lesson_plan_location_labels", []) for t in low):
        return "pop_quiz", chapter, lesson, "its questions state a Lesson Plan Location"
    fb = [normalize(l).lower() for l in vocab.get("data_bank_fields", {}).get("feedback_incorrect", [])]
    if cells and any(any(t.startswith(l) for l in fb) for t in [normalize(c).lower() for c in cells]):
        return "data_bank", chapter, lesson, "its items carry feedback fields"
    heads = " ".join(low[:400])
    if ("narration" in heads or "on-screen" in heads or "بیانیہ" in heads) and ("scene" in heads or "منظر" in heads):
        return "video_storyboard", chapter, lesson, "it has storyboard columns (scene, narration)"
    # 3. mostly questions: the per-lesson Assessment
    qrx = question_regex(profile)
    qs = sum(1 for t in paras if qrx.match(t) and re.search(r"[?؟]|_{3,}", t))
    if qs >= 4:
        return "chapter_exam", chapter, lesson, f"it is made of questions ({qs} found)"
    return GENERIC, chapter, lesson, "its kind could not be told from its name or content"


def classify(root, profile) -> Package:
    root = Path(root)
    docs, unclassified = [], []
    chapter_titles, lesson_labels = {}, {}
    rules = profile["doc_types"]
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name.startswith(("~$", ".")):
            continue
        rel = path.relative_to(root).as_posix()
        parts = rel.split("/")
        ext = path.suffix.lower().lstrip(".")
        stem = path.stem
        chapter, lesson, variant, folder_label = _parse_location(parts, profile)
        matched = None
        for r in rules:
            if ext in r.get("ext", [ext]) and re.search(r["pattern"], stem, re.I):
                matched = r
                break
        why = ""
        if ext in ("doc", "ppt"):
            converted = convert_legacy(path)
            if converted is None:
                unclassified.append(f"{rel} (an old Word/PowerPoint 97-2003 file: open it and Save As .docx/.pptx, then upload it again)")
                continue
            path, ext = converted, converted.suffix.lstrip(".")
            rel = path.relative_to(root).as_posix()
            parts = rel.split("/")
            why = "converted from the old format; "
            if matched is not None and ext not in matched.get("ext", [ext]):
                matched = None
            if matched is None:
                for r in rules:
                    if ext in r.get("ext", [ext]) and re.search(r["pattern"], stem, re.I):
                        matched = r
                        break
        if matched is None and ext in ("docx", "pptx"):
            # the name does not say what the file is: its content does
            dtype, c_ch, c_ls, reason = classify_by_content(path, ext, profile)
            if dtype is None:
                unclassified.append(f"{rel} ({reason})")
                continue
            matched = next((r for r in rules if r["type"] == dtype), {"type": dtype, "scope": "lesson"})
            chapter = chapter if chapter is not None else c_ch
            lesson = lesson if lesson is not None else c_ls
            assumed = ""
            if matched.get("scope", "lesson") == "lesson" and dtype != GENERIC:
                if chapter is None:
                    chapter, assumed = 1, " chapter 1"
                if lesson is None:
                    lesson, assumed = 1, assumed + " lesson 1"
            why += f"Recognised as {dtype.replace('_', ' ')} from its content: {reason}." + \
                (f" Neither the name nor the title gives its number, so{assumed} was assumed." if assumed else "")
        if matched is None:
            unclassified.append(rel)
            continue
        ref = DocRef(rel=rel, abs=str(path), ext=ext, doc_type=matched["type"], chapter=chapter,
                     lesson=lesson, variant=variant, scope=matched.get("scope", "lesson"),
                     folder_label=folder_label)
        if matched.get("note"):
            ref.notes.append(matched["note"])
        if why:
            ref.notes.append(why)
        # subject-scope files have no chapter/lesson of their own
        if ref.scope == "subject" or ref.doc_type == GENERIC:
            ref.chapter = ref.lesson = None
            ref.variant = ""
        if ref.scope == "chapter":
            ref.lesson = None
            ref.variant = ""
        docs.append(ref)
        for p in parts[:-1]:
            m = re.search(profile["path_patterns"]["chapter"], p)
            if m and ref.chapter is not None and int(m.group(1)) == ref.chapter and re.match(r"(?i)(chapter|chpater)", p):
                chapter_titles.setdefault(ref.chapter, p)
        if ref.lesson is not None:
            lesson_labels.setdefault(ref.key, folder_label)

    # newest version of each subject-scope document wins; older ones are superseded
    by_type = {}
    for d in docs:
        if d.doc_type in ("data_bank", "pop_quiz", "specification"):
            by_type.setdefault(d.doc_type, []).append(d)
    for dtype, group in by_type.items():
        if len(group) > 1:
            group.sort(key=lambda d: _version_key(os.path.basename(d.rel)))
            for old in group[:-1]:
                old.superseded = True
                old.notes.append(f"Superseded by {os.path.basename(group[-1].rel)}")

    subject = root.name
    return Package(root=str(root), subject=subject, docs=docs, unclassified=unclassified,
                   chapter_titles=chapter_titles, lesson_labels=lesson_labels)
