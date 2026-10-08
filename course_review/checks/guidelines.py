"""Rules from each output's own guidelines and template, beyond the checklist workbook.

The team reviews every output against three sources: the checklist, that output's guideline
document, and the Writing & Editing Guidelines. The workbook covers the first; this module adds
the rules the guideline documents state that code can check. Codes here are not workbook codes,
so they appear in the marked-up documents and the report, not in the filled workbook:

  LPG  Lesson Plan Guidelines and template      FGG  Facilitator's Guide Guidelines and template
  ASG  Assessments and Worksheets Guidelines     SB   Storyboarding Guidelines and template
  WEG  Writing & Editing Guidelines rules the checklist does not list

FG1 and FG4 (checklist codes) get a code-decided result here where the slides make it clear.
"""
import re
from difflib import SequenceMatcher

from ..docx_model import DocxInfo
from ..models import FAIL, NA, PASS, REVIEW
from ..textutil import normalize, to_western_digits
from .common import mark, result

GUIDE_RULES = {
    "LPG1": "Warm-up Activity: 1 to 3 questions, never more (Lesson Plan Guidelines)",
    "LPG2": "Key Takeaways recap the lesson in bullets (Lesson Plan Guidelines)",
    "LPG3": "A Book and SLO Coverage table maps Book Heading -> Lesson Plan Location -> SLOs (Lesson Plan Guidelines and template)",
    "LPG4": "Difficult terms are listed under a 'Glossary Terms' heading at the end of Concept Building (Glossary guidelines)",
    "FGG1": "Facilitator Notes sit in the speaker-notes pane of every slide (Facilitator's Guide Guidelines)",
    "FGG2": "Every slide's notes carry a time allocation, adding up to the lesson's planned duration (Facilitator's Guide Guidelines)",
    "FGG3": "The Cover slide states the topic and the duration (Facilitator's Guide Guidelines)",
    "ASG1": "No 'all of the above' / 'none of the above' options (Assessments and Worksheets Guidelines)",
    "SB1": "Storyboard table has Scene, Duration, On-Screen Text, Narration and Animation/Visual columns (Storyboard template)",
    "SB2": "Video or simulation runs 2 minutes at most (Storyboarding Guidelines)",
    "SB3": "Narration is about 260-300 words for 2 minutes, never more (Storyboarding Guidelines)",
    "SB4": "The top of the storyboard says 'New' or 'Edited from: <source>' (Storyboarding Guidelines)",
    "SB5": "On-screen text is key terms and short labels, not full sentences (Storyboarding Guidelines)",
    "WEG1": "8 pt space after each paragraph (Writing & Editing Guidelines, page setup)",
    "WEG2": "Lists inside a section use bullets; numbering is only for the major headings (Writing & Editing Guidelines)",
    "STR1": "The model and the formatting rules disagree whether a line is a heading; the tool left it as it is",
}


def _labels(*xs):
    return [normalize(x).lower() for x in xs]


# ------------------------------------------------------------------ helpers
def lp_ranges(ctx, info):
    """{section key: (heading paragraph index, end index)} for the five Lesson Plan sections."""
    from .lessonplan import find_sections
    found = sorted((v, k) for k, v in find_sections(ctx, info).items() if v is not None)
    return {k: (v, found[i + 1][0] if i + 1 < len(found) else len(info.paras)) for i, (v, k) in enumerate(found)}


def _paras(info, rng):
    return [p for p in info.paras[rng[0] + 1:rng[1]] if p.text.strip()]


def _is_bullet(p):
    from .lessonplan import is_bullet_text
    return p.list_kind == "bullet" or is_bullet_text(p.text)


def _is_numbered(p):
    return p.list_kind == "decimal" or bool(re.match(r"^\s*\(?\d+[.)]\s", to_western_digits(p.text)))


# ------------------------------------------------------------- Lesson Plan
def lpg1(ctx, doc, info):
    r = lp_ranges(ctx, info).get("warm_up")
    if not r:
        return result("LPG1", REVIEW, "The Warm-up section was not found")
    qlabel = re.compile(r"^\s*[•\-–—*]?\s*(?:" + "|".join(re.escape(normalize(x)) for x in ctx.profile["vocab"]["question_prefixes"]) + r")\s*\d",
                        re.I)
    qs = [p for p in _paras(info, r) if re.search(r"[?؟]", p.text) or qlabel.match(to_western_digits(normalize(p.text)))]
    from ..structure import of
    st = of(info)
    if st is not None:
        mq = [info.paras[i] for i in st.lists.get("warm_up_questions", [])]
        qs = mq or qs
    if not qs:
        return result("LPG1", REVIEW, "No question was found in the Warm-up (it may be an activity without questions)")
    if len(qs) > 3:
        return result("LPG1", FAIL, f"The Warm-up has {len(qs)} questions; the guideline allows 1 to 3",
                      marks=[mark(p.text, "warm-up question beyond the third") for p in qs[3:]])
    return result("LPG1", PASS, f"The Warm-up has {len(qs)} question(s)")


def takeaway_paras(ctx, info):
    """The Key Takeaways items, or None when the section was not found."""
    r = lp_ranges(ctx, info).get("key_takeaways")
    if not r:
        return None
    from ..structure import of
    st = of(info)
    if st is not None and st.lists.get("takeaway_items"):
        return [info.paras[i] for i in st.lists["takeaway_items"]]
    return [p for p in _paras(info, r) if not p.in_table]     # the coverage table after it is not a takeaway


def lpg2(ctx, doc, info):
    ps = takeaway_paras(ctx, info)
    if ps is None:
        return result("LPG2", REVIEW, "The Key Takeaways section was not found")
    bullets = [p for p in ps if _is_bullet(p)]
    numbered = [p for p in ps if _is_numbered(p)]
    if numbered:
        return result("LPG2", FAIL, f"{len(numbered)} Key Takeaways item(s) are numbered, not bullets",
                      marks=[mark(p.text, "numbered; Key Takeaways are written as bullets") for p in numbered])
    if not bullets:
        return result("LPG2", FAIL, "Key Takeaways are not written as bullets",
                      marks=[mark(ps[0].text, "not a bullet; Key Takeaways are written as bullets")] if ps else [])
    return result("LPG2", PASS, f"Key Takeaways recap in {len(bullets)} bullets")


COVERAGE_HEADERS = _labels("Book Heading", "Lesson Plan Location", "SLO", "Possible Lottie", "کتاب کا عنوان", "سبق کا مقام",
                           "حاصلات تعلیم", "لوٹی")


def lpg3(ctx, doc, info):
    for t in info.tables:
        header = " ".join(normalize(info.paras[i].text).lower() for i in t.para_idx if info.paras[i].row_idx == 0)
        if sum(1 for h in COVERAGE_HEADERS if h in header) >= 2:
            return result("LPG3", PASS, f"Book and SLO coverage table found (table {t.idx + 1})", partial=True,
                          evidence=["Whether every book heading and SLO has a row needs a reviewer."])
    title = next((p for p in info.paras if any(l in normalize(p.text).lower() for l in _labels("Book and SLO Coverage", "کتاب اور SLO"))), None)
    if title:
        return result("LPG3", REVIEW, "A coverage heading exists but no table with the template's columns follows it",
                      marks=[mark(title.text, "coverage table expected here")])
    return result("LPG3", FAIL, "No Book and SLO Coverage table (Book Heading / Lesson Plan Location / SLOs)")


def lpg4(ctx, doc, info):
    labels = _labels("Glossary Terms", "Glossary", "فرہنگ", "اصطلاحات", "مشکل الفاظ")
    if any(any(normalize(p.text).lower().startswith(l) for l in labels) for p in info.paras if len(p.text) < 60):
        return result("LPG4", PASS, "A Glossary Terms heading is present")
    return result("LPG4", REVIEW, "No 'Glossary Terms' heading at the end of Concept Building; needed if the lesson uses difficult terms")


# ---------------------------------------------------- Writing & Editing extras
def weg1(ctx, doc, info):
    want = ctx.profile.get("house_rules", {}).get("paragraph_space_after_pt", 8)
    body = [p for p in info.paras if p.text.strip() and not p.in_table]
    if not body:
        return result("WEG1", NA, "No paragraphs")
    off = [p for p in body if p.space_after is None or abs(p.space_after - want) > 0.5]
    if len(off) / len(body) > ctx.th["text_size_tolerance"]:
        seen = {}
        for p in off:
            k = "not set" if p.space_after is None else f"{p.space_after:g} pt"
            seen[k] = seen.get(k, 0) + 1
        return result("WEG1", FAIL, f"{len(off)} of {len(body)} paragraphs do not have {want} pt after (" +
                      ", ".join(f"{k}: {v}" for k, v in sorted(seen.items(), key=lambda kv: -kv[1])[:3]) + ")")
    return result("WEG1", PASS, f"Paragraphs have {want} pt after")


def weg2(ctx, doc, info):
    from .formatting import doc_headings
    heads = {id(p) for p in doc_headings(ctx, info)}
    bad = [p for p in info.paras if p.text.strip() and not p.in_table and id(p) not in heads and p.list_kind == "decimal"]
    if bad:
        return result("WEG2", FAIL, f"{len(bad)} list item(s) are numbered; lists inside a section use bullets",
                      marks=[mark(p.text, "numbered list item; use a bullet") for p in bad])
    return result("WEG2", PASS, "No numbered lists inside sections")


# ------------------------------------------------------- Facilitator's Guide
NOTE_ITEMS = {
    "how to explain": _labels("how to explain", "analogy", "example", "demonstrat", "سمجھا", "تمثیل", "مثال", "وضاحت"),
    "check for understanding": _labels("check for understanding", "check-for-understanding", "show of hands", "poll", "quick question",
                                       "فہم کی جانچ", "ہاتھ اٹھوا", "رائے شماری"),
    "misconceptions": _labels("misconception", "common mistake", "غلط فہمی"),
    "supporting media cue": _labels("video", "simulation", "game", "media", "ویڈیو", "سیمولیشن", "گیم", "میڈیا"),
    "transition line": _labels("transition", "bridge", "next concept", "ربط", "اگلے تصور", "آگے بڑھ"),
}
CB_LABELS = _labels("Concept Building", "تصوراتی تعمیر")
SLO_LABELS = _labels("SLO", "Learning Outcomes", "Session Overview", "حاصلات تعلیم", "سبق کے مقاصد", "سیشن کا خلاصہ")


def _note_text(ctx, s):
    """Speaker notes plus any on-slide text labelled as a facilitator note."""
    labels = _labels(*ctx.profile["vocab"].get("facilitator_note_labels", []), "سہولت کار کے لیے نوٹ", "تدریسی ہدایت")
    on_slide = [t for t in s.texts if any(l in normalize(t).lower()[:80] for l in labels)]
    return normalize(" ".join([s.notes] + on_slide)).lower()


def fgg1(ctx, doc, info):
    slides = info.slides_text
    missing = [s.index for s in slides if not s.notes.strip()]
    if len(missing) == len(slides):
        return result("FGG1", FAIL, "No slide has speaker notes; Facilitator Notes belong in the notes pane, not on the slide")
    if missing:
        return result("FGG1", FAIL, f"Slide(s) {missing} have no speaker notes",
                      marks=[mark(s.title, "no Facilitator Notes in the speaker-notes pane", slide=s.index) for s in slides if s.index in missing])
    return result("FGG1", PASS, "Every slide has speaker notes")


_MIN = re.compile(r"(\d+)\s*(?:[-–—]\s*(\d+))?\s*(?:min|minutes|منٹ)", re.I)


def _minutes(text):
    m = _MIN.search(to_western_digits(normalize(text)))
    if not m:
        return None
    lo = int(m.group(1))
    return lo, int(m.group(2) or lo)


def fgg2(ctx, doc, info):
    slides = info.slides_text
    timed = {s.index: _minutes(s.notes) for s in slides if _minutes(s.notes)}
    untimed = [s.index for s in slides if s.index not in timed]
    total = None
    for s in slides[:3]:
        t = normalize(" ".join(s.texts)).lower()
        if any(l in t for l in _labels("duration", "دورانیہ", "total", "کل")):
            mm = _minutes(" ".join(s.texts))
            if mm:
                total = mm[1]
                break
    if not timed:
        return result("FGG2", FAIL, "No slide's notes carry a time allocation (e.g. 'Time: 3 min' / 'وقت: 3 منٹ')")
    ev = []
    if untimed:
        ev.append(f"slide(s) {untimed} have no time allocation")
    lo, hi = sum(v[0] for v in timed.values()), sum(v[1] for v in timed.values())
    if total and not (lo * 0.9 <= total <= hi * 1.1):
        ev.append(f"slide times add up to {lo}-{hi} min but the lesson is planned for {total} min")
    if ev:
        return result("FGG2", FAIL, "; ".join(ev), ev)
    return result("FGG2", PASS, f"Every slide is timed ({lo}-{hi} min" + (f", planned {total} min)" if total else ")"))


def fgg3(ctx, doc, info):
    if not info.slides_text:
        return result("FGG3", NA, "No slides")
    cover = " ".join(info.slides_text[0].texts + [info.slides_text[0].notes])
    has_time = _minutes(cover) is not None
    if not has_time:
        return result("FGG3", FAIL, "The Cover slide does not state the lesson's duration",
                      marks=[mark(info.slides_text[0].title, "add the duration (e.g. 40 minutes) to the cover", slide=1)])
    return result("FGG3", PASS, "The Cover slide states the duration", partial=True)


def fg4(ctx, doc, info):
    from ..structure import of
    st = of(info)
    if st is not None:
        cb = [s for s in info.slides_text if st.slide_roles.get(s.index) == "concept_building"]
    else:
        cb = [s for s in info.slides_text if any(l in normalize(s.title).lower() for l in CB_LABELS)]
    if not cb:
        return result("FG4", REVIEW, "No Concept Building slide was recognised by its title")
    no_notes, gaps = [], {}
    for s in cb:
        t = _note_text(ctx, s)
        if not t.strip():
            no_notes.append(s.index)
            continue
        miss = [k for k, words in NOTE_ITEMS.items() if not any(w in t for w in words)]
        if miss:
            gaps[s.index] = miss
    marks = [mark(s.title, "Concept Building slide without Facilitator Notes", slide=s.index) for s in cb if s.index in no_notes]
    if no_notes:
        return result("FG4", FAIL, f"Concept Building slide(s) {no_notes} have no Facilitator Notes, so none of the five items",
                      [f"slide {k}: no sign of {', '.join(v)}" for k, v in gaps.items()], marks=marks)
    if gaps:
        return result("FG4", REVIEW, "Some Concept Building notes seem to lack items (found by wording): " +
                      "; ".join(f"slide {k}: {', '.join(v)}" for k, v in gaps.items()))
    return result("FG4", PASS, "Every Concept Building slide's notes mention all five items", partial=True,
                  evidence=["Found by wording; whether each item is done well needs a reviewer."])


def fg1_verbatim(ctx, doc, info):
    """FG1 is a Fail when the Session Overview pastes the Lesson Plan's SLO wording (the guideline asks for
    plain, learner-facing language). Otherwise None: the model or a reviewer decides."""
    from .judgement import lesson_plan_doc
    lp = lesson_plan_doc(ctx, doc.key)
    lpi = ctx.docx(lp) if lp else None
    if lpi is None:
        return None
    from .judgement import slides_for
    from .lessonplan import slo_items
    items, _ = slo_items(ctx, lpi)               # the lesson's SLO statements (the model's list when it read the plan)
    if not items:
        return None
    slo_lines = [normalize(p.text).strip(" •-–—*") for p in items if len(p.text.strip()) >= 25]
    overview = slides_for(ctx, info, "slos")     # the Session Overview slide(s) (the model's reading of the deck)
    slide = overview[0] if overview else next((s for s in info.slides_text if any(l in normalize(s.title).lower() for l in SLO_LABELS)), None)
    if slide is None or not slo_lines:
        return None
    same = []
    for t in slide.texts:
        for line in t.split("\n"):
            nl = normalize(line).strip(" •-–—*")
            if len(nl) >= 25 and any(SequenceMatcher(None, nl, s).ratio() >= 0.9 for s in slo_lines):
                same.append(line)
    if len(same) >= max(1, len(slo_lines) // 2):
        return result("FG1", FAIL, f"Session Overview pastes {len(same)} SLO(s) word for word from the Lesson Plan instead of "
                                   "restating them in plain, learner-facing language",
                      marks=[mark(x, "SLO pasted from the Lesson Plan; restate it in plain language", slide=slide.index) for x in same])
    return None


# ---------------------------------------------------------------- Assessments
ALL_NONE = _labels("all of the above", "none of the above", "مذکورہ بالا تمام", "مندرجہ بالا تمام", "ان میں سے کوئی نہیں",
                   "مذکورہ بالا میں سے کوئی نہیں", "تمام درست ہیں")


def asg1(ctx, doc, info):
    hits = [p for p in info.paras if any(l in normalize(p.text).lower() for l in ALL_NONE)]
    if hits:
        return result("ASG1", FAIL, f"{len(hits)} option(s) use 'all of the above' / 'none of the above'",
                      marks=[mark(p.text, "avoid 'all/none of the above'") for p in hits])
    return result("ASG1", PASS, "No 'all of the above' / 'none of the above' options")


# ----------------------------------------------------------------- Storyboard
SB_COLS = {
    "scene": _labels("Scene", "Step", "منظر"),
    "duration": _labels("Duration", "دورانیہ"),
    "on_screen": _labels("On-Screen", "On Screen", "اسکرین"),
    "narration": _labels("Narration", "بیانیہ"),
    "visual": _labels("Animation", "Visual", "متحرک", "بصری"),
}


def _storyboard_tables(doc):
    """[(column map, rows)] for tables whose header names storyboard columns."""
    from docx import Document
    out = []
    for t in Document(doc.abs).tables:
        if not t.rows:
            continue
        header = [normalize(c.text).lower() for c in t.rows[0].cells]
        cols = {}
        for key, labels in SB_COLS.items():
            for i, h in enumerate(header):
                if any(l in h for l in labels):
                    cols.setdefault(key, i)
        if len(cols) >= 3:
            rows = [[c.text.strip() for c in r.cells] for r in t.rows[1:]]
            out.append((cols, [r for r in rows if any(x for x in r)]))
    return out


def _seconds(text):
    t = to_western_digits(normalize(text))
    times = re.findall(r"(\d+):(\d{2})", t)
    if times:
        return max(int(m) * 60 + int(s) for m, s in times), True          # a range end such as 1:45
    m = re.search(r"(\d+(?:\.\d+)?)\s*(s|sec|seconds|سیکنڈ|m|min|منٹ)", t, re.I)
    if m:
        v = float(m.group(1))
        return (v * 60 if m.group(2).lower() in ("m", "min", "منٹ") else v), False
    return None, False


def storyboard_checks(ctx, doc, info):
    out = []
    try:
        tables = _storyboard_tables(doc)
    except Exception as e:  # unreadable table structure
        return [result("SB1", REVIEW, f"Storyboard tables could not be read ({type(e).__name__})")]
    if not tables:
        return [result("SB1", FAIL, "No storyboard table with Scene / Duration / On-Screen Text / Narration / Animation columns")]
    cols, rows = max(tables, key=lambda t: len(t[1]))
    missing = [k for k in SB_COLS if k not in cols]
    out.append(result("SB1", FAIL if missing else PASS,
                      ("Storyboard table lacks column(s): " + ", ".join(missing)) if missing else f"Storyboard table with {len(rows)} scenes"))
    if "duration" in cols:
        vals = [_seconds(r[cols["duration"]]) for r in rows if cols["duration"] < len(r)]
        ends = [v for v, is_end in vals if v is not None and is_end]
        durs = [v for v, is_end in vals if v is not None and not is_end]
        total = max(ends) if ends else (sum(durs) if durs else None)
        if total is None:
            out.append(result("SB2", REVIEW, "Scene durations could not be read"))
        elif total > 120:
            out.append(result("SB2", FAIL, f"Runs {int(total // 60)}:{int(total % 60):02d}; the cap is 2:00"))
        else:
            out.append(result("SB2", PASS, f"Runs {int(total // 60)}:{int(total % 60):02d}"))
    if "narration" in cols:
        words = sum(len(r[cols["narration"]].split()) for r in rows if cols["narration"] < len(r))
        if words > 300:
            out.append(result("SB3", FAIL, f"Narration is {words} words; 2 minutes holds about 260-300"))
        else:
            out.append(result("SB3", PASS, f"Narration is {words} words"))
    if "on_screen" in cols:
        long = [r[cols["on_screen"]] for r in rows if cols["on_screen"] < len(r) and len(r[cols["on_screen"]].split()) > 12]
        if long:
            out.append(result("SB5", FAIL, f"{len(long)} scene(s) put full sentences on screen; use key terms and short labels",
                              marks=[mark(x, "on-screen text should be a short label") for x in long]))
        else:
            out.append(result("SB5", PASS, "On-screen text is short"))
    top = " ".join(p.text for p in info.paras[:12]).lower()
    if any(l in normalize(top) for l in _labels("new video", "new.", "edited from", "نیا", "ترمیم شدہ", "ماخوذ", "path a", "path b")) or \
            re.search(r"\bnew\b", top):
        out.append(result("SB4", PASS, "States whether it is new or edited from a source"))
    else:
        out.append(result("SB4", FAIL, "The top does not say 'New' or 'Edited from: <source>'"))
    return out


# --------------------------------------------------------------------- wiring
DOC_RULES = {
    "lesson_plan": [lpg1, lpg2, lpg3, lpg4, weg1, weg2],
    "chapter_exam": [asg1, weg1],
    "worksheet": [asg1, weg1],
    "pop_quiz": [asg1],
    "data_bank": [asg1],
    "video_storyboard": [weg1, weg2],
}
SLIDE_RULES = [fgg1, fgg2, fgg3, fg4]


def doc_checks(ctx, doc, info):
    out = []
    if isinstance(info, DocxInfo):
        for fn in DOC_RULES.get(doc.doc_type, []):
            out.append(fn(ctx, doc, info))
        if doc.doc_type == "video_storyboard":
            out.extend(storyboard_checks(ctx, doc, info))
    elif doc.doc_type == "facilitator_guide":
        for fn in SLIDE_RULES:
            out.append(fn(ctx, doc, info))
        f = fg1_verbatim(ctx, doc, info)
        if f is not None:
            out.append(f)
    return out
