"""Parse a .docx into a normalised model with *effective* formatting.

Checks never touch python-docx objects; they read DocxInfo. The model resolves the
style chain (run -> character style -> paragraph style -> defaults -> theme) so a
font or size is what Word would actually show, not just what is set on the run.
Urdu/Arabic text uses the complex-script slots (w:cs, w:szCs, w:bCs), which
python-docx does not expose, so those are read from the XML directly.
"""
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Optional

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from .textutil import normalize, script_counts, to_western_digits

try:
    from PIL import Image
except Exception:  # Pillow is optional; image checks degrade to needs_review
    Image = None

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"


def _val(el):
    return None if el is None else el.get(qn("w:val"))


def _on(el):
    if el is None:
        return None
    return el.get(qn("w:val")) not in ("0", "false", "off")


@dataclass
class RunInfo:
    text: str
    script: Optional[str]          # 'arabic' | 'latin' | None (digits / punctuation only)
    font: Optional[str]            # effective font for this script
    size: Optional[float]          # points, for this script
    bold: Optional[bool]
    italic: Optional[bool]
    color: Optional[str]           # hex 'RRGGBB' or None (auto)
    theme_color: Optional[str]
    highlight: Optional[str]
    shading: Optional[str]         # run-level shading fill
    rtl: bool


@dataclass
class ImageInfo:
    descr: str
    name: str
    width_emu: int
    height_emu: int
    px: Optional[tuple]            # (w, h) pixels, None for vector
    fmt: str
    colourful: Optional[bool]      # None if undecidable
    max_chroma: float = 0.0
    blob_len: int = 0


@dataclass
class ParaInfo:
    idx: int
    text: str
    style: str                     # style name
    heading_level: Optional[int]
    is_title: bool
    bidi: bool
    runs: list
    in_table: bool
    table_idx: Optional[int]
    row_idx: Optional[int]
    num_label: Optional[str]       # auto-number label e.g. '1.2'
    list_kind: Optional[str]       # 'bullet' | 'decimal' | 'other' | None
    line_spacing: tuple            # ('multiple', 1.15) | ('points', 14.0) | ('default', 1.0)
    images: list
    shading: Optional[str]
    all_bold: bool
    in_toc: bool = False
    hanging: int = 0          # hanging indent in twips (0 = none)
    space_after: Optional[float] = None   # points after the paragraph (None = not set anywhere)


@dataclass
class TableInfo:
    idx: int
    rows: int
    cols: int
    para_idx: list
    cell_shadings: list
    prev_para: Optional[int]       # index in paras of the paragraph just before the table
    next_para: Optional[int]


@dataclass
class DocxInfo:
    path: str
    paras: list = field(default_factory=list)
    tables: list = field(default_factory=list)
    blocks: list = field(default_factory=list)       # ('p', para_idx) | ('t', table_idx)
    page_w_in: float = 0
    page_h_in: float = 0
    orientation: str = ""
    margins_in: tuple = ()
    sections: list = field(default_factory=list)
    footer_text: str = ""
    footer_fields: set = field(default_factory=set)
    header_text: str = ""
    has_toc: bool = False
    pages: Optional[int] = None
    images: list = field(default_factory=list)
    style_colors: dict = field(default_factory=dict)
    language: Optional[str] = None
    char_count: int = 0

    def body_paras(self):
        return [p for p in self.paras if not p.in_table]


class _Resolver:
    """Effective-property lookup through the style chain."""

    def __init__(self, doc):
        self.doc = doc
        styles = doc.styles.element
        self.styles = {s.get(qn("w:styleId")): s for s in styles.findall(qn("w:style"))}
        dd = styles.find(qn("w:docDefaults"))
        self.def_r = dd.find(qn("w:rPrDefault") + "/" + qn("w:rPr")) if dd is not None else None
        self.def_p = dd.find(qn("w:pPrDefault") + "/" + qn("w:pPr")) if dd is not None else None
        self.default_pstyle = next((sid for sid, s in self.styles.items()
                                    if s.get(qn("w:type")) == "paragraph" and s.get(qn("w:default")) == "1"), None)
        self.theme = self._theme_fonts()

    def _theme_fonts(self):
        out = {}
        try:
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            part = self.doc.part.part_related_by(RT.THEME)
            root = etree.fromstring(part.blob)
            for kind in ("major", "minor"):
                f = root.find(f".//{{{A_NS}}}{kind}Font")
                if f is None:
                    continue
                latin = f.find(f"{{{A_NS}}}latin")
                cs = f.find(f"{{{A_NS}}}cs")
                out[f"{kind}_latin"] = latin.get("typeface") if latin is not None else None
                cs_face = cs.get("typeface") if cs is not None else None
                if not cs_face:
                    for fe in f.findall(f"{{{A_NS}}}font"):
                        if fe.get("script") == "Arab":
                            cs_face = fe.get("typeface")
                out[f"{kind}_cs"] = cs_face
        except Exception:
            pass
        return out

    def style_name(self, sid):
        s = self.styles.get(sid)
        if s is None:
            return ""
        n = s.find(qn("w:name"))
        return n.get(qn("w:val")) if n is not None else sid

    def chain(self, sid):
        seen = set()
        while sid and sid in self.styles and sid not in seen:
            seen.add(sid)
            s = self.styles[sid]
            yield s
            b = s.find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None

    def _pstyle_id(self, p_el):
        ppr = p_el.find(qn("w:pPr"))
        if ppr is not None:
            ps = ppr.find(qn("w:pStyle"))
            if ps is not None:
                return ps.get(qn("w:val"))
        return self.default_pstyle

    def rpr_candidates(self, r_el, p_el):
        """rPr elements from most to least specific."""
        out = []
        rpr = r_el.find(qn("w:rPr"))
        if rpr is not None:
            out.append(rpr)
            rs = rpr.find(qn("w:rStyle"))
            if rs is not None:
                for s in self.chain(rs.get(qn("w:val"))):
                    e = s.find(qn("w:rPr"))
                    if e is not None:
                        out.append(e)
        for s in self.chain(self._pstyle_id(p_el)):
            e = s.find(qn("w:rPr"))
            if e is not None:
                out.append(e)
        if self.def_r is not None:
            out.append(self.def_r)
        return out

    def first(self, cands, getter):
        for c in cands:
            v = getter(c)
            if v is not None:
                return v
        return None

    def _theme_name(self, theme, kind):
        if not theme:
            return None
        which = "major" if theme.startswith("major") else "minor"
        return self.theme.get(f"{which}_{kind}")

    def font(self, cands, slot):
        """Effective font for slot 'ascii' (Latin) or 'cs' (complex script, e.g. Urdu)."""
        def g(rpr):
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                return None
            if slot == "cs":
                return rf.get(qn("w:cs")) or self._theme_name(rf.get(qn("w:csTheme")), "cs")
            return (rf.get(qn("w:ascii")) or rf.get(qn("w:hAnsi"))
                    or self._theme_name(rf.get(qn("w:asciiTheme")) or rf.get(qn("w:hAnsiTheme")), "latin"))
        return self.first(cands, g)

    def size(self, cands, cs):
        tag = qn("w:szCs") if cs else qn("w:sz")
        v = self.first(cands, lambda r: _val(r.find(tag)))
        return float(v) / 2 if v is not None else None

    def flag(self, cands, cs, base):
        tag = qn("w:" + base + ("Cs" if cs else ""))
        return self.first(cands, lambda r: _on(r.find(tag)))

    def color(self, cands):
        def g(rpr):
            c = rpr.find(qn("w:color"))
            if c is None:
                return None
            return (c.get(qn("w:val")), c.get(qn("w:themeColor")))
        return self.first(cands, g) or (None, None)

    def simple(self, cands, tag, attr="w:val"):
        return self.first(cands, lambda r: (r.find(qn(tag)).get(qn(attr)) if r.find(qn(tag)) is not None else None))

    def ppr_chain(self, p_el):
        out = []
        ppr = p_el.find(qn("w:pPr"))
        if ppr is not None:
            out.append(ppr)
        for s in self.chain(self._pstyle_id(p_el)):
            e = s.find(qn("w:pPr"))
            if e is not None:
                out.append(e)
        if self.def_p is not None:
            out.append(self.def_p)
        return out


class _Numbering:
    def __init__(self, doc):
        self.abs, self.nums = {}, {}
        self.counters = {}
        try:
            part = doc.part.numbering_part
        except Exception:
            return
        root = part.element
        for a in root.findall(qn("w:abstractNum")):
            lv = {}
            for l in a.findall(qn("w:lvl")):
                fmt = _val(l.find(qn("w:numFmt")))
                txt = _val(l.find(qn("w:lvlText")))
                start = _val(l.find(qn("w:start")))
                lv[int(l.get(qn("w:ilvl")))] = (fmt, txt, int(start) if start else 1)
            self.abs[a.get(qn("w:abstractNumId"))] = lv
        for n in root.findall(qn("w:num")):
            aid = _val(n.find(qn("w:abstractNumId")))
            self.nums[n.get(qn("w:numId"))] = aid

    def label(self, num_id, ilvl):
        aid = self.nums.get(num_id)
        lv = self.abs.get(aid)
        if not lv or ilvl not in lv:
            return None, None
        fmt, txt, start = lv[ilvl]
        if fmt == "bullet" or fmt == "none":
            return "", "bullet" if fmt == "bullet" else None
        counters = self.counters.setdefault(aid, {})
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]
        def render(m):
            n = int(m.group(1)) - 1
            c = counters.get(n)
            if c is None:
                c = lv.get(n, (None, None, 1))[2]
            f = lv.get(n, ("decimal",))[0]
            return str(c) if f in ("decimal", "decimalZero") else _fmt_other(c, f)
        label = re.sub(r"%(\d)", render, txt or "")
        return label, "decimal" if fmt in ("decimal", "decimalZero") else "other"


def _fmt_other(n, fmt):
    if fmt == "lowerLetter":
        return chr(ord("a") + (n - 1) % 26)
    if fmt == "upperLetter":
        return chr(ord("A") + (n - 1) % 26)
    return str(n)


def _image_info(blob, fmt):
    px = None
    colourful = None
    max_chroma = 0.0
    if Image is not None and fmt.lower() not in ("emf", "wmf", "svg"):
        try:
            im = Image.open(io.BytesIO(blob))
            px = im.size
            small = im.convert("RGB")
            small.thumbnail((200, 200))
            pixels = list(small.getdata())
            if pixels:
                spread = sorted(max(p) - min(p) for p in pixels)
                max_chroma = spread[int(len(spread) * 0.98)]  # ignore a few anti-aliased edge pixels
                colourful = max_chroma > 20
        except Exception:
            pass
    return px, colourful, max_chroma


def _para_images(doc, p_el):
    out = []
    for d in p_el.iter(qn("w:drawing")):
        holder = d.find(f"{{{WP_NS}}}inline")
        if holder is None:
            holder = d.find(f"{{{WP_NS}}}anchor")
        if holder is None:
            continue
        docpr = holder.find(f"{{{WP_NS}}}docPr")
        ext = holder.find(f"{{{WP_NS}}}extent")
        blip = holder.find(f".//{{{A_NS}}}blip")
        blob, fmt = b"", ""
        if blip is not None:
            rid = blip.get(f"{{{R_NS}}}embed")
            try:
                part = doc.part.related_parts[rid]
                blob = part.blob
                fmt = part.partname.ext.lstrip(".")
            except Exception:
                pass
        px, colourful, chroma = _image_info(blob, fmt) if blob else (None, None, 0.0)
        out.append(ImageInfo(
            descr=(docpr.get("descr") or "").strip() if docpr is not None else "",
            name=docpr.get("name") if docpr is not None else "",
            width_emu=int(ext.get("cx")) if ext is not None else 0,
            height_emu=int(ext.get("cy")) if ext is not None else 0,
            px=px, fmt=fmt, colourful=colourful, max_chroma=chroma, blob_len=len(blob)))
    for pict in p_el.iter(qn("w:pict")):
        out.append(ImageInfo(descr="", name="(legacy picture)", width_emu=0, height_emu=0, px=None,
                             fmt="vml", colourful=None))
    return out


def parse_docx(path) -> DocxInfo:
    doc = Document(path)
    res = _Resolver(doc)
    numbering = _Numbering(doc)
    info = DocxInfo(path=str(path))

    # page setup
    secs = doc.sections
    info.sections = [(s.page_width.inches if s.page_width else 0, s.page_height.inches if s.page_height else 0,
                      (s.top_margin.inches if s.top_margin is not None else None,
                       s.bottom_margin.inches if s.bottom_margin is not None else None,
                       s.left_margin.inches if s.left_margin is not None else None,
                       s.right_margin.inches if s.right_margin is not None else None),
                      "landscape" if (s.page_width or 0) > (s.page_height or 0) else "portrait") for s in secs]
    s0 = info.sections[0]
    info.page_w_in, info.page_h_in, info.margins_in, info.orientation = s0

    # footer / header text and fields
    def part_text_fields(hf):
        text, fields = [], set()
        el = hf._element
        for p in el.iter(qn("w:p")):
            text.append("".join(t.text or "" for t in p.iter(qn("w:t"))))
        for it in el.iter(qn("w:instrText")):
            m = re.match(r"\s*([A-Z]+)", it.text or "")
            if m:
                fields.add(m.group(1))
        for fs in el.iter(qn("w:fldSimple")):
            m = re.match(r"\s*([A-Z]+)", fs.get(qn("w:instr")) or "")
            if m:
                fields.add(m.group(1))
        return " ".join(t for t in text if t.strip()), fields
    try:
        ft, ff = part_text_fields(secs[0].footer)
        if secs[0].different_first_page_header_footer and not ft.strip():
            ft, ff = part_text_fields(secs[0].first_page_footer)
        info.footer_text, info.footer_fields = ft, ff
        info.header_text, _ = part_text_fields(secs[0].header)
    except Exception:
        pass

    # pages from docProps/app.xml
    try:
        with zipfile.ZipFile(path) as z:
            if "docProps/app.xml" in z.namelist():
                m = re.search(rb"<Pages>(\d+)</Pages>", z.read("docProps/app.xml"))
                if m:
                    info.pages = int(m.group(1))
    except Exception:
        pass

    body = doc.element.body
    table_counter = [0]

    def run_info(r_el, p_el):
        text = "".join(t.text or "" for t in r_el.iter(qn("w:t")))
        if not text:
            return None
        ar, la = script_counts(text)
        script = "arabic" if ar > la else ("latin" if la > 0 else None)
        cs = script == "arabic"
        cands = res.rpr_candidates(r_el, p_el)
        color, theme = res.color(cands)
        shd = res.first(cands, lambda r: (r.find(qn("w:shd")).get(qn("w:fill")) if r.find(qn("w:shd")) is not None else None))
        return RunInfo(
            text=text, script=script,
            font=res.font(cands, "cs" if cs else "ascii"),
            size=res.size(cands, cs),
            bold=res.flag(cands, cs, "b"), italic=res.flag(cands, cs, "i"),
            color=color, theme_color=theme,
            highlight=res.simple(cands, "w:highlight"), shading=shd,
            rtl=res.first(cands, lambda r: True if r.find(qn("w:rtl")) is not None else None) is True,
        )

    def line_spacing(p_el):
        for ppr in res.ppr_chain(p_el):
            sp = ppr.find(qn("w:spacing"))
            if sp is not None and sp.get(qn("w:line")) is not None:
                line = int(sp.get(qn("w:line")))
                rule = sp.get(qn("w:lineRule")) or "auto"
                if rule == "auto":
                    return ("multiple", round(line / 240.0, 3))
                return ("points", line / 20.0)
        return ("default", 1.0)

    def space_after(p_el):
        for ppr in res.ppr_chain(p_el):
            sp = ppr.find(qn("w:spacing"))
            if sp is not None and sp.get(qn("w:after")) is not None:
                return int(sp.get(qn("w:after"))) / 20.0
        return None

    def hanging_indent(p_el):
        for ppr in res.ppr_chain(p_el):
            ind = ppr.find(qn("w:ind"))
            if ind is not None and (ind.get(qn("w:hanging")) is not None or ind.get(qn("w:left")) is not None):
                return int(ind.get(qn("w:hanging")) or 0)
        return 0

    def para_info(p_el, in_table, table_idx, row_idx):
        runs = []
        for r_el in p_el.iter(qn("w:r")):
            ri = run_info(r_el, p_el)
            if ri:
                runs.append(ri)
        text = "".join(t.text or "" for t in p_el.iter(qn("w:t")))
        sid = res._pstyle_id(p_el)
        sname = res.style_name(sid)
        low = sname.lower()
        hl = None
        m = re.match(r"heading\s*(\d)", low)
        if m:
            hl = int(m.group(1))
        else:
            for s in res.chain(sid):
                ol = s.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
                if ol is not None and ol.get(qn("w:val")) is not None:
                    hl = int(ol.get(qn("w:val"))) + 1
                    break
        ppr = p_el.find(qn("w:pPr"))
        bidi = False
        for c in res.ppr_chain(p_el):
            b = c.find(qn("w:bidi"))
            if b is not None:
                bidi = _on(b) is not False
                break
        # numbering
        num_label = list_kind = None
        for c in res.ppr_chain(p_el):
            np_ = c.find(qn("w:numPr"))
            if np_ is not None:
                nid = _val(np_.find(qn("w:numId")))
                il = _val(np_.find(qn("w:ilvl")))
                if nid and nid != "0":
                    num_label, list_kind = numbering.label(nid, int(il) if il else 0)
                break
        shd = None
        if ppr is not None and ppr.find(qn("w:shd")) is not None:
            shd = ppr.find(qn("w:shd")).get(qn("w:fill"))
        letters = [r for r in runs if r.script]
        all_bold = bool(letters) and all(r.bold for r in letters)
        return ParaInfo(
            idx=len(info.paras), text=text, style=sname, heading_level=hl,
            is_title=(low == "title"), bidi=bidi, runs=runs, in_table=in_table, table_idx=table_idx,
            row_idx=row_idx, num_label=num_label, list_kind=list_kind, line_spacing=line_spacing(p_el), space_after=space_after(p_el),
            images=_para_images(doc, p_el), shading=shd, all_bold=all_bold,
            in_toc=low.startswith("toc") or low == "table of contents", hanging=hanging_indent(p_el))

    def walk(container, in_toc_sdt=False):
        for child in container:
            tag = etree.QName(child).localname
            if tag == "p":
                pi = para_info(child, False, None, None)
                pi.in_toc = pi.in_toc or in_toc_sdt
                info.paras.append(pi)
                info.blocks.append(("p", pi.idx))
                for it in child.iter(qn("w:instrText")):
                    if re.match(r"\s*TOC\b", it.text or ""):
                        info.has_toc = True
            elif tag == "tbl":
                ti = table_counter[0]
                table_counter[0] += 1
                prev = info.blocks[-1][1] if info.blocks and info.blocks[-1][0] == "p" else None
                idxs, shadings = [], []
                rows = child.findall(qn("w:tr"))
                ncols = 0
                for ri, tr in enumerate(rows):
                    cells = tr.findall(qn("w:tc"))
                    ncols = max(ncols, len(cells))
                    for tc in cells:
                        tcpr = tc.find(qn("w:tcPr"))
                        if tcpr is not None and tcpr.find(qn("w:shd")) is not None:
                            shadings.append(tcpr.find(qn("w:shd")).get(qn("w:fill")))
                        for p_el in tc.iter(qn("w:p")):
                            pi = para_info(p_el, True, ti, ri)
                            info.paras.append(pi)
                            idxs.append(pi.idx)
                info.tables.append(TableInfo(ti, len(rows), ncols, idxs, shadings, prev, None))
                info.blocks.append(("t", ti))
            elif tag == "sdt":
                gallery = child.find(qn("w:sdtPr") + "/" + qn("w:docPartObj") + "/" + qn("w:docPartGallery"))
                is_toc = gallery is not None and "table of contents" in (_val(gallery) or "").lower()
                if is_toc:
                    info.has_toc = True
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    walk(content, in_toc_sdt or is_toc)
    walk(body)
    for t in info.tables:  # next paragraph after each table
        pos = info.blocks.index(("t", t.idx))
        if pos + 1 < len(info.blocks) and info.blocks[pos + 1][0] == "p":
            t.next_para = info.blocks[pos + 1][1]

    for p in info.paras:
        info.images.extend(p.images)
        info.char_count += len(p.text)
    ar = sum(script_counts(p.text)[0] for p in info.paras)
    la = sum(script_counts(p.text)[1] for p in info.paras)
    info.language = None if ar + la == 0 else ("ur" if ar >= la else "en")

    # colours set directly on styles that are used
    used = {p.style for p in info.paras}
    for sid, s in res.styles.items():
        name = res.style_name(sid)
        if name in used:
            c = s.find(qn("w:rPr") + "/" + qn("w:color"))
            if c is not None:
                info.style_colors[name] = (c.get(qn("w:val")), c.get(qn("w:themeColor")))
    return info
