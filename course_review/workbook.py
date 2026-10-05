"""Read the rules from the checklist workbook and write results back into a copy of it.

The workbook is the authority for rule codes, wording and scope. Nothing about a
rule's text is duplicated in code; checks are bound to codes in checks/__init__.py.
"""
import copy
import hashlib
import re
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import MultiCellRange

from .models import Rule

CODE_RE = re.compile(r"^[A-Z]{2,3}\d+$")
HEADER_ROW = 4
EXAMPLE_ROW = 5
FIRST_DATA_ROW = 6
STATUS_WORD = {"pass": "Pass", "fail": "Fail", "na": "N/A"}


def file_sha1(path) -> str:
    h = hashlib.sha1()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()[:12]


class Layout:
    """Where each code lives in the workbook."""

    def __init__(self):
        self.matrix = {}      # sheet -> {codes, col, scope, fails_col, notes_col, last_row}
        self.subject_sheet = None
        self.subject_rows = {}  # code -> row
        self.subject_cols = {}  # status/notes column indexes


def load_rules(path):
    wb = openpyxl.load_workbook(path)
    rules = {}
    key = wb["Key"]
    for row in key.iter_rows(min_row=5, values_only=True):
        code, otype, scope, text = (row + (None,) * 4)[:4]
        if code and CODE_RE.match(str(code).strip()):
            rules[str(code).strip()] = Rule(str(code).strip(), str(otype or "").strip(), str(scope or "").strip(), str(text or "").strip())

    layout = Layout()
    for ws in wb.worksheets:
        if ws.title in ("Index", "Key"):
            continue
        header = [c.value for c in ws[HEADER_ROW]]
        if header and header[0] == "#":
            codes, col = [], {}
            fails_col = notes_col = None
            for i, v in enumerate(header, start=1):
                if v and CODE_RE.match(str(v)):
                    codes.append(str(v))
                    col[str(v)] = i
                elif v == "Fails":
                    fails_col = i
                elif v == "Notes":
                    notes_col = i
            scopes = {rules[c].scope for c in codes if c in rules}
            scope = "chapter" if scopes == {"Per chapter"} else "lesson"
            layout.matrix[ws.title] = dict(codes=codes, col=col, scope=scope, fails_col=fails_col,
                                           notes_col=notes_col, last_row=ws.max_row)
        elif header[:2] == ["Code", "Section"]:
            layout.subject_sheet = ws.title
            layout.subject_cols = {str(v): i for i, v in enumerate(header, start=1) if v}
            for r in range(HEADER_ROW + 1, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if v and CODE_RE.match(str(v)):
                    layout.subject_rows[str(v)] = r
    return rules, layout


def _copy_row_style(ws, src_row, dst_row, max_col):
    for c in range(1, max_col + 1):
        s, d = ws.cell(src_row, c), ws.cell(dst_row, c)
        d._style = copy.copy(s._style)
    if ws.row_dimensions[src_row].height:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def fill_workbook(template_path, out_path, grid, subject_cells, summary_rows, meta):
    """grid: {sheet: [ {label, cells: {code: 'pass'|'fail'|'na'|None}, note} ]}
    subject_cells: {code: (status|None, note)}
    summary_rows: list of tuples for the 'Review Summary' sheet."""
    rules, layout = load_rules(template_path)
    wb = openpyxl.load_workbook(template_path)

    for sheet, spec in layout.matrix.items():
        ws = wb[sheet]
        rows = grid.get(sheet, [])
        max_col = ws.max_column
        first_code_col = min(spec["col"].values())
        last_code_col = max(spec["col"].values())
        # example row becomes a legend row
        for c in range(1, max_col + 1):
            ws.cell(EXAMPLE_ROW, c).value = None
        ws.cell(EXAMPLE_ROW, 1).value = "Auto"
        ws.cell(EXAMPLE_ROW, 2).value = "Filled by Course Review Tool. A blank cell means the tool did not decide it."
        if spec["fails_col"]:
            ws.cell(EXAMPLE_ROW, spec["fails_col"]).value = None
        if spec["notes_col"]:
            ws.cell(EXAMPLE_ROW, spec["notes_col"]).value = "Pass/Fail were decided by code. Evidence is in Notes. Blank = reviewer to complete."

        needed_last = FIRST_DATA_ROW + max(len(rows), 1) - 1
        last_row = max(spec["last_row"], needed_last)
        for r in range(spec["last_row"] + 1, last_row + 1):
            _copy_row_style(ws, spec["last_row"], r, max_col)

        for i in range(FIRST_DATA_ROW, last_row + 1):
            idx = i - FIRST_DATA_ROW
            fl = get_column_letter(first_code_col)
            ll = get_column_letter(last_code_col)
            if idx < len(rows):
                row = rows[idx]
                ws.cell(i, 1).value = idx + 1
                ws.cell(i, 2).value = row["label"]
                for code, col in spec["col"].items():
                    v = row["cells"].get(code)
                    ws.cell(i, col).value = STATUS_WORD.get(v)
                if spec["notes_col"]:
                    ws.cell(i, spec["notes_col"]).value = row.get("note") or None
            else:
                ws.cell(i, 1).value = None
            if spec["fails_col"]:
                ws.cell(i, spec["fails_col"]).value = f'=COUNTIF({fl}{i}:{ll}{i},"Fail")'

        for dv in ws.data_validations.dataValidation:
            dv.sqref = MultiCellRange(f"{get_column_letter(first_code_col)}{EXAMPLE_ROW}:{get_column_letter(last_code_col)}{last_row}")

    if layout.subject_sheet:
        ws = wb[layout.subject_sheet]
        sc, nc = layout.subject_cols.get("Status"), layout.subject_cols.get("Notes")
        for code, row in layout.subject_rows.items():
            status, note = subject_cells.get(code, (None, None))
            ws.cell(row, sc).value = STATUS_WORD.get(status)
            if nc:
                ws.cell(row, nc).value = note or None

    # summary sheet
    if "Review Summary" in wb.sheetnames:
        del wb["Review Summary"]
    ws = wb.create_sheet("Review Summary", 2)
    ws.append(["Course Review Tool run"])
    ws["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    for k, v in meta.items():
        ws.append([k, str(v)])
    ws.append([])
    ws.append(["Code", "Checked by tool", "How", "Wording"])
    for c in ws[ws.max_row]:
        c.font = openpyxl.styles.Font(bold=True)
    for row in summary_rows:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 90
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    wb.save(out_path)
