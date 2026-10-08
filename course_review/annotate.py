"""Marked-up copy of the package: every checklist result written into the documents themselves.

Red highlight = fails a checklist rule, turquoise = the model's suggestion for a reviewer to
check, bright green = passes. Each highlighted paragraph carries a Word comment naming the
checklist code and the reason. Problems that belong to the whole document (fonts, page size,
file name) go into one comment on the first paragraph. Existing highlights are never replaced,
because yellow on a correct answer is itself something the checklist requires.

PowerPoint has no comments the libraries can write, so slides get the highlight plus a small
review box listing the notes for that slide. Results that belong to no file (a missing
document, subject-wide coverage) go into REVIEW-NOTES.txt at the top of the zip.
"""
import re
import shutil
import tempfile
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from . import plain
from .models import FAIL, PASS, REVIEW
from .textutil import normalize, to_western_digits

AUTHOR, INITIALS = "Course Review", "CR"
RANK = {FAIL: 3, REVIEW: 2, PASS: 1}
WORD_COLOUR = {FAIL: WD_COLOR_INDEX.RED, REVIEW: WD_COLOR_INDEX.TURQUOISE, PASS: WD_COLOR_INDEX.BRIGHT_GREEN}
SLIDE_COLOUR = {FAIL: "FF0000", REVIEW: "00FFFF", PASS: "00FF00"}
LABEL = {FAIL: "FAIL", REVIEW: "CHECK", PASS: "PASS"}
CHANGED = "changed"                   # a line the tool corrected in this copy
WORD_COLOUR[CHANGED] = WD_COLOR_INDEX.PINK
COLOUR_ORDER = {FAIL: 4, REVIEW: 3, CHANGED: 2, PASS: 1}
LEGEND = plain.LEGEND


def _model_lines(f, note, first):
    """Comment text for a model suggestion: what it suggests, why, and what the reviewer does."""
    verdict = "fail" if "FAIL" in (note or "") or "Would be fail" in f.message else \
              "pass" if "PASS" in (note or "") or "Would be pass" in f.message else "unclear"
    body = (note or f.message).replace("[model-assisted] ", "")
    body = re.sub(r"^(?:model suggests (?:PASS|FAIL)|Would be (?:pass|fail))\s*:\s*", "", body)
    if not first:
        body = body.split("; the story")[0]
    return plain.ai_lines(f.code, verdict, body, first, _output_name(f.code))
MIN_MATCH = 6


def _loose(text):
    t = re.sub(r"[^\w\s]", " ", to_western_digits(normalize(text or "")))
    return re.sub(r"\s+", " ", t).strip().lower()


def _pieces(text):
    """A mark may be a quote that skips words with '...' or spans several paragraphs."""
    parts = re.split(r"\.\.\.|…|\[\.\.\.[^\]]*\]|\n", text or "")
    return [p for p in (_loose(x) for x in parts) if len(p) >= MIN_MATCH]


def _matches(pieces, para_loose):
    """Indexes of paragraphs a mark points at: the paragraph contains a piece, or (for a quote
    spanning paragraphs) a non-trivial paragraph is wholly inside a piece."""
    hits = []
    for i, t in enumerate(para_loose):
        if not t:
            continue
        if any(p in t for p in pieces) or (len(t) >= MIN_MATCH * 2 and any(t in p for p in pieces)):
            hits.append(i)
    return hits


def _locate(m, para_loose, used, allowed=None):
    """Paragraph indexes a mark points at. A whole-paragraph mark takes one paragraph: the first
    equal one this finding has not used yet, so repeated headings or table cells are marked once
    each, in order. A quote marks every paragraph it touches."""
    ok = (lambda i: True) if allowed is None else allowed
    if not m.get("exact", True):
        return [i for i in _matches(_pieces(m["text"]), para_loose) if ok(i)]
    key = _unnumbered(_loose(m["text"]))
    if not key:
        return []
    # a heading the fixer numbered ('4.5 ...') still matches its unnumbered text, before any containment match
    cands = [i for i, t in enumerate(para_loose) if (t == key or _unnumbered(t) == key) and ok(i)]
    if not cands and len(key) >= MIN_MATCH:
        cands = [i for i, t in enumerate(para_loose) if key in t and ok(i)]
    if not cands:
        return []
    pick = next((i for i in cands if i not in used), cands[0])
    used.add(pick)
    return [pick]


# The workbook's sheet names are swapped against the team's names for these outputs: its "Chapter
# Exam" sheet (CE) is the per-lesson Assessment, its "Worksheet" sheet (WS) is the per-chapter
# Chapter Exam. Comments use the team's names so a reviewer is not confused.
OUTPUT_NAME = {"CE": "Assessment", "WS": "Chapter Exam"}


def _output_name(code):
    return OUTPUT_NAME.get(re.sub(r"\d+$", "", code or ""))


def _code(code):
    name = _output_name(code)
    return f"{code} ({name})" if name else code


def _unnumbered(loose_text):
    return re.sub(r"^(?:\d+\s+)+", "", loose_text or "")


def _line(f, note=""):
    if not f.code:                        # a _Note already holds a finished line
        return f.message
    if (note or "").startswith("model suggests"):
        return _model_lines(f, note, True)[0]
    words = {FAIL: plain.fail_line, REVIEW: plain.check_line}.get(f.status, plain.pass_line)
    return words(f.code, note or f.message, _output_name(f.code))


def _worst(statuses):
    return max(statuses, key=lambda s: RANK.get(s, 0))


def _relevant(findings):
    return [f for f in findings if f.status in RANK]


# ---------------------------------------------------------------------------- Word

def _all_paragraphs(doc):
    body = doc.element.body
    return [Paragraph(p, doc._body) for p in body.iter(qn("w:p"))]


def _text_runs(par):
    return [r for r in par.runs if r.text.strip()]


def annotate_docx(src, dst, findings, fixed=None, changes=None):
    """Write a marked-up copy of one .docx; returns how many paragraphs were marked. `fixed` lists what
    the tool already corrected in this copy, for the summary comment; `changes` are the lines it corrected
    (autofix.fix_docx), each highlighted pink with what it was, why, and what it is now."""
    doc = Document(src)
    paras = _all_paragraphs(doc)
    loose = [_loose(p.text) for p in paras]
    per_para = defaultdict(list)          # para index -> [(status, line)]
    unplaced = []
    for f in _relevant(findings):
        placed, used = False, set()
        model = f.method == "model"
        marks = f.marks
        if model and ("PASS" in (marks[0].get("note", "") if marks else "") or "Would be pass" in f.message) \
                and len({m.get("note") for m in marks}) <= 1:
            marks = marks[:1]             # a suggested pass needs one place to confirm it, not every passage it read
                                          # (marks with different notes are different claims, and each is shown)
        suggested_pass = model and (marks is not f.marks or len(marks) == 1)
        notes_seen = set()
        for m in marks:
            hits = _locate(m, loose, used)
            if suggested_pass:
                hits = hits[:1]               # a quote can span paragraphs (a table); one place is enough to confirm a pass
            for i in hits:
                new_claim = m.get("note") not in notes_seen      # a different note is a claim of its own, explained in full
                lines = _model_lines(f, m.get("note"), new_claim) if model else [_line(f, m.get("note"))]
                for l in lines:
                    per_para[i].append((f.status, l))
                placed = True
                notes_seen.add(m.get("note"))
        if not placed and f.status != PASS:
            unplaced.append(f)

    used_c = set()
    raw = [normalize(p.text).strip() for p in paras]
    for c in changes or []:
        want = normalize(c["text"]).strip()
        same = [i for i, t in enumerate(raw) if t == want and i not in used_c]     # the very line first, then a looser match
        hits = same[:1] or _locate({"text": c["text"], "exact": True}, loose, used_c)[:1]
        used_c.update(hits)
        for i in hits:
            per_para[i].append((CHANGED, plain.fixed_line(c.get("plain") or "see the summary at the top.")))

    for i, items in sorted(per_para.items()):
        runs = _text_runs(paras[i])
        if not runs:
            unplaced_lines = [l for _, l in items if l]
            per_para[i] = []
            unplaced.extend(_Note(l) for l in unplaced_lines)
            continue
        colour = WORD_COLOUR[max((s for s, _ in items), key=lambda s: COLOUR_ORDER.get(s, 0))]
        for r in runs:
            if r.font.highlight_color is None and r._r.find(qn("w:rPr") + "/" + qn("w:shd")) is None:
                r.font.highlight_color = colour
        lines = list(dict.fromkeys(l for _, l in items if l))
        if lines:
            _comment(doc, runs, lines)

    first = next((p for p in paras if _text_runs(p)), None)
    if first is not None:
        problems = [f for f in unplaced if f.status in (FAIL, REVIEW)]
        codes = {f.code for f in problems}
        problems = [f for f in problems if not (f.code in plain.REPEATS and len(codes) > 1)
                    and not (f.code == "WE1" and "WE3" in codes)]       # both only ask for the version in the name
        lines = ["Course Review: how to read this file", LEGEND]
        if fixed:
            lines.append("What we fixed for you:")
            lines += [f"- {l}" for l in fixed]
        if problems:
            lines.append("Still to do in the whole file:")
            lines += [f"- {_whole_line(f)}" for f in problems]
        marked = sum(1 for v in per_para.values() if any(s != CHANGED for s, _ in v))
        pink = sum(1 for v in per_para.values() if v and all(s == CHANGED for s, _ in v))
        lines.append(f"In the text: {marked} place(s) need you (red or turquoise)." if marked else
                     "In the text: nothing else needs you.")
        if pink:
            lines.append(f"{pink} line(s) are pink: we fixed them. Delete those notes once you have looked.")
        _comment(doc, _text_runs(first)[:1], lines)
    from .ooxml import order_run_properties
    order_run_properties(doc.element)
    for part in (doc.part._comments_part,) if doc.comments else ():
        order_run_properties(part.element)
    doc.save(dst)
    return sum(1 for v in per_para.values() if v)


def _whole_line(f):
    """A result about the whole document, as a to-do line."""
    if not f.code:
        return f.message
    if f.status == FAIL:
        return plain.todo_line(f.code, f.message, _output_name(f.code))
    if f.method == "model":
        return plain.ai_lines(f.code, "fail" if "fail" in f.message.lower()[:40] else "pass" if "pass" in f.message.lower()[:40]
                              else "unclear", re.sub(r"^\[model-assisted\] (?:Would be (?:pass|fail): )?", "", f.message), True,
                              _output_name(f.code))[0]
    if f.code in plain.TODO_WORDS:
        return "Please check: " + plain.TODO_WORDS[f.code](f.message) + plain.tag(f.code, _output_name(f.code))
    return plain.check_line(f.code, f.message, _output_name(f.code))


class _Note:
    """A message that could not be anchored to its paragraph, shown in the top comment."""

    def __init__(self, line):
        self.status = FAIL if line.startswith("To fix") else REVIEW if line.startswith(("Please check", "AI check")) else PASS
        self.code, self.message, self.marks = "", line, []


def _comment(doc, runs, lines):
    c = doc.add_comment(runs, text=lines[0], author=AUTHOR, initials=INITIALS)
    for l in lines[1:]:
        c.add_paragraph(l)


# ---------------------------------------------------------------------- PowerPoint

def annotate_pptx(src, dst, findings, fixed=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    from lxml import etree

    prs = Presentation(src)
    slides = list(prs.slides)
    slide_paras = []                      # (slide number, paragraph, loose text)
    for n, s in enumerate(slides, start=1):
        for shape in _shapes(s.shapes):
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    slide_paras.append((n, p, _loose("".join(r.text for r in p.runs))))
    loose = [t for _, _, t in slide_paras]
    notes = defaultdict(list)             # slide number -> [(status, line)]
    marked_paras = defaultdict(list)
    unplaced = []
    for f in _relevant(findings):
        placed, used = False, set()
        model = f.method == "model"
        marks = f.marks
        if model and ("PASS" in (marks[0].get("note", "") if marks else "") or "Would be pass" in f.message):
            marks = marks[:1]
        for m in marks:
            on_slide = None if m.get("slide") is None else (lambda i, s=m["slide"]: slide_paras[i][0] == s)
            hits = _locate(m, loose, used, on_slide)
            for i in (hits[:1] if model and marks is not f.marks else hits):
                n = slide_paras[i][0]
                marked_paras[i].append(f.status)
                for l in (_model_lines(f, m.get("note"), not placed) if model else [_line(f, m.get("note"))]):
                    notes[n].append((f.status, l))
                placed = True
        if not placed and f.status != PASS:
            unplaced.append(f)

    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for i, statuses in marked_paras.items():
        colour = SLIDE_COLOUR[_worst(statuses)]
        for r in slide_paras[i][1].runs:
            rpr = r._r.get_or_add_rPr()
            if rpr.find(f"{{{a}}}highlight") is not None:
                continue
            hl = etree.SubElement(rpr, f"{{{a}}}highlight")
            etree.SubElement(hl, f"{{{a}}}srgbClr", val=colour)
            # schema order: highlight must come after fill/effect elements and before fonts
            for tag in ("latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"):
                for el in rpr.findall(f"{{{a}}}{tag}"):
                    rpr.remove(el)
                    rpr.append(el)

    if slides:
        problems = [f for f in unplaced if f.status in (FAIL, REVIEW)]
        top = [(FAIL, "Course Review. " + LEGEND)] + [(PASS, plain.fixed_line(l)) for l in (fixed or [])] + \
              [(f.status, _whole_line(f)) for f in problems]
        notes[1] = top + notes.get(1, [])
    for n, items in notes.items():
        lines = list(dict.fromkeys(l for _, l in items))
        box = slides[n - 1].shapes.add_textbox(Emu(int(prs.slide_width * 0.02)), Emu(int(prs.slide_height * 0.02)),
                                              Emu(int(prs.slide_width * 0.96)), Emu(int(prs.slide_height * 0.12)))
        box.name = "Course Review notes"
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xFF, 0xF4, 0xF4)
        box.line.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        tf = box.text_frame
        tf.word_wrap = True
        for k, l in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = l
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(0x60, 0x00, 0x00)
    prs.save(dst)
    return len(marked_paras)


def _shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from _shapes(sh.shapes)
        else:
            yield sh


# ------------------------------------------------------------------------ package

def annotate_package(pkg, findings, out_zip, rules=None, profile=None, fix=True, ctx=None):
    """Zip of the whole package with each .docx/.pptx fixed where the rules are mechanical and
    commented where a person must decide, plus REVIEW-NOTES.txt."""
    from . import autofix
    root = Path(pkg.root)
    by_doc = defaultdict(list)
    loose_ends = []
    for f in findings:
        (by_doc[f.doc] if f.doc else loose_ends).append(f)
    if not fix:
        ctx = None
    elif ctx is None and profile is not None:       # the review's own context carries what the model read
        from .engine import Context
        ctx = Context(pkg, profile, rules or {})
    grade_subject = re.sub(r"[^A-Za-z0-9]+", "-", pkg.subject or "").strip("-")
    work = Path(tempfile.mkdtemp(prefix="course-markup-"))
    # "open": (path inside the package, finding) for every Fail / needs-reviewer result still standing in the
    # copies handed out, under the copy's new name; results tied to no file have path None.
    stats = {"documents": 0, "paragraphs": 0, "errors": {}, "renamed": {}, "open": []}
    written = {}                              # original path -> (path of the copy, findings written into it)
    try:
        tree = work / root.name
        shutil.copytree(root, tree)
        for d in pkg.docs:
            fs = by_doc.get(d.rel, [])
            if d.ext not in ("docx", "pptx") or not fs or d.superseded:
                continue
            dst = tree / d.rel
            try:
                fixed_lines, extra, src, changes = [], {}, d.abs, []
                if ctx is not None:
                    info = ctx.docx(d) if d.ext == "docx" else ctx.pptx(d)
                    version, found_in = autofix.find_version(d, info)
                    name = autofix.new_name(d, grade_subject, version)
                    base = name.rsplit(".", 1)[0] if name else Path(d.rel).stem
                    if version:                   # the footer shows the version once, in its own field
                        base = re.sub(r"[\s_-]*(?:v|version\s*)" + re.escape(version) + r"\b.*$", "", base, flags=re.I).strip(" -_") or base
                    staged = work / ("fixed." + d.ext)
                    if d.ext == "docx":
                        if info is not None:
                            fixed_lines, extra = autofix.fix_docx(d.abs, staged, ctx, d, info)
                            changes = extra.get("changes", [])
                    else:
                        footer = f"{base}" + (f" | v{version}" if version else "") + f" | {date.today().isoformat()}"
                        fixed_lines = autofix.fix_pptx(d.abs, staged, ctx, footer if name else None)
                    if fixed_lines:
                        src = str(staged)
                    if name and name != dst.name:
                        dst.unlink(missing_ok=True)
                        dst = dst.with_name(name)
                        stats["renamed"][d.rel] = name
                        fixed_lines.append(f"File renamed to {name} (the team's naming pattern)")
                    if d.ext == "docx" and src == str(staged):
                        autofix.add_footer(src, base, version)
                        fixed_lines.append("Footer added: document name, " + ("version, " if version else "") + "date, page number")
                    fixed_lines[:0] = [plain.source_note(n) for n in d.notes if n.startswith(("Recognised as", "converted from"))]
                    if name and version and found_in != "the file name":
                        fixed_lines.append(f"Version v{version} taken from {found_in}")
                    rel = (d.rel.rsplit("/", 1)[0] + "/" if "/" in d.rel else "") + dst.name
                    fs = autofix.remaining(fs, d.ext, autofix.recheck(ctx, d, src, rel), extra)
                n = annotate_docx(src, dst, fs, fixed_lines, changes) if d.ext == "docx" else annotate_pptx(src, dst, fs, fixed_lines)
                written[d.rel] = (dst.relative_to(tree).as_posix(), fs)
                stats["documents"] += 1
                stats["paragraphs"] += n
            except Exception as e:   # a file the libraries cannot rewrite stays as it was, and says so
                shutil.copy2(d.abs, tree / d.rel)
                stats["errors"][d.rel] = f"{type(e).__name__}: {e}"
                loose_ends.append(_Note(f"{d.rel}: could not be fixed or marked up ({type(e).__name__}); its results are below"))
                loose_ends.extend(f for f in fs if f.status in (FAIL, REVIEW))
                written[d.rel] = (d.rel, by_doc.get(d.rel, []))
        for rel, fs in by_doc.items():
            path, fs = written.get(rel, (rel, fs))
            stats["open"].extend((path, f) for f in fs if f.status in (FAIL, REVIEW) and getattr(f, "code", ""))
        stats["open"].extend((None, f) for f in findings if not f.doc and f.status in (FAIL, REVIEW) and getattr(f, "code", ""))
        (tree / "REVIEW-NOTES.txt").write_text(_notes_text(loose_ends, pkg, rules, stats["renamed"]), encoding="utf-8")
        out_zip = Path(out_zip)
        out_zip.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as z:
            for p in sorted(tree.rglob("*")):
                if p.is_file() and not p.name.startswith("fixed."):
                    z.write(p, p.relative_to(work).as_posix())
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return stats


def _notes_text(findings, pkg, rules, renamed=None):
    lines = ["Course Review notes", "",
             "Open each Word file: the notes (comments) in it say what we fixed and what is left for you.",
             "The first note at the top of each file is a summary.",
             LEGEND, ""]
    if renamed:
        lines += ["Files renamed to the team's naming pattern:"] + [f"- {old}  ->  {new}" for old, new in sorted(renamed.items())] + [""]
    lines += [f.message + "." for f in findings if getattr(f, "code", "") == "NOTE"]
    lines += ["", "Other results (not about one file):", ""]
    groups = defaultdict(lambda: ([], []))   # (status, message) -> (codes, places)
    for f in findings:
        if f.status not in (FAIL, REVIEW):
            continue
        codes, places = groups[(f.status, f.message)]
        if f.code and _code(f.code) not in codes:
            codes.append(_code(f.code))
        where = f.lesson.short() if getattr(f, "lesson", None) else (f"Ch{f.chapter}" if getattr(f, "chapter", None) else "")
        if where and where not in places:
            places.append(where)
    if not groups:
        lines.append("None.")
    for (status, msg), (codes, places) in groups.items():
        lead = "To fix" if status == FAIL else "Please check"
        lines.append(f"- {lead}: {msg.replace('[model-assisted] ', '')}" + (f" ({', '.join(places)})" if places else "") +
                     (f" [{', '.join(codes)}]" if codes else ""))
    if pkg.unclassified:
        lines += ["", "Files that were not reviewed:"]
        lines += [f"- {u}" for u in pkg.unclassified]
    return "\n".join(lines) + "\n"
