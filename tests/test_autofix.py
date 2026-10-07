"""The fixed copies: reviewing them again, the mechanical Writing & Editing rules pass."""
import json
import zipfile

from docx.enum.text import WD_COLOR_INDEX

from course_review import cli

from helpers import A4, add, checklist, new_doc, zip_dir

URDU = "صبر کا مطلب ہے رکنا اور برداشت کرنا۔"


def package(tmp_path, lp_name="Lesson-2-Chapter-4-Lesson-Plan.docx", footer=None):
    root = tmp_path / "Grade-6-Islamiat"
    lesson = root / "Chapter-4-Akhlaq" / "Lesson-2-Sabr"
    lesson.mkdir(parents=True)
    d = new_doc(page=(8.5, 11), margin=0.75, spacing=1.0, size=11)
    add(d, "سبق: صبر و تحمل", bold=True, rtl=True)
    for s in ("تعارف", "حاصلات تعلیم", "تحریکی سرگرمی", "تصوراتی تعمیر", "حاصل کلام"):
        add(d, s, bold=True, size=14, rtl=True)
        add(d, URDU + " آیت ١٥٣", rtl=True, color="2E74B5")
    d.add_paragraph("پہلا نکتہ", style="List Number")
    d.add_paragraph("دوسرا نکتہ", style="List Number")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 0).text = "عنوان", "متن"
    if footer:
        d.sections[0].footer.paragraphs[0].text = footer
    d.save(str(lesson / lp_name))
    pq = new_doc()
    add(pq, "Pop Quiz")
    add(pq, "سبق 2: صبر و تحمل", bold=True)
    add(pq, "سوال 1: صبر کیا ہے؟")
    add(pq, "✅ الف) رکنا")
    add(pq, "ب) بھاگنا")
    (root / "Pop Quiz").mkdir()
    pq.save(str(root / "Pop Quiz" / "Pop Quiz.docx"))
    return zip_dir(str(root), str(tmp_path / "pkg.zip"))


def statuses(out, name):
    r = json.load(open(out / "review.json", encoding="utf-8"))
    return {f["code"]: f["status"] for f in r["findings"] if (f["doc"] or "").endswith(name)}


def test_mechanical_rules_pass_after_fixing(tmp_path):
    out1, out2 = tmp_path / "o1", tmp_path / "o2"
    cl = str(checklist(tmp_path))
    cli.review(package(tmp_path, "Lesson-2-Chapter-4-Lesson-Plan-v1.2.docx"), cl, str(out1))
    before = statuses(out1, "Lesson-2-Chapter-4-Lesson-Plan-v1.2.docx")
    for code in ("WE4", "WE8", "WE10", "WE13", "WE16", "WE23", "WE25", "WEG1", "WEG2"):
        assert before[code] == "fail", code
    fixed = out1 / cli.MARKED_UP
    names = zipfile.ZipFile(fixed).namelist()
    assert any(n.endswith("Lesson-Plan-Lesson-2-Chapter-4-v1.2.docx") for n in names)     # the file's own version kept
    cli.review(str(fixed), cl, str(out2))
    after = statuses(out2, "Lesson-Plan-Lesson-2-Chapter-4-v1.2.docx")
    for code in ("WE1", "WE4", "WE7", "WE8", "WE10", "WE11", "WE12", "WE13", "WE16", "WE25", "WEG1", "WEG2"):
        assert after[code] == "pass", (code, after[code])
    assert after["WE23"] == "pass" or "highlight" in json.dumps(json.load(open(out2 / "review.json", encoding="utf-8")))


def test_check_mark_becomes_a_yellow_highlight_and_lesson_sections_still_split(tmp_path):
    out1, out2 = tmp_path / "o1", tmp_path / "o2"
    cl = str(checklist(tmp_path))
    cli.review(package(tmp_path), cl, str(out1))
    cli.review(str(out1 / cli.MARKED_UP), cl, str(out2))
    r = json.load(open(out2 / "review.json", encoding="utf-8"))
    pq = {f["code"]: f for f in r["findings"] if f["lesson"] == "Ch4-L2" and f["code"].startswith("PQ")}
    assert pq["PQ1"]["status"] in ("pass", "fail") and "Pop Quiz has" not in pq["PQ1"]["message"]
    assert pq["PQ4"]["status"] == "pass"


def test_only_a_fail_the_fixed_copy_no_longer_shows_is_left_out():
    from course_review import autofix
    from course_review.checks.common import result
    fs = [result("LP5", "needs_review", "The SLO section was not found"),       # a rule the fixer does not settle: kept
          result("WE4", "fail", "Letter size"),                                 # fixed: the copy passes
          result("WE3", "fail", "No version number"),                           # the copy still needs a person
          result("WE8", "fail", "Body text is 11 pt"),                          # the copy still fails
          result("WE21", "fail", "Citation is not (Author, Year)"),             # not a code the fixer works on
          result("LP10", "fail", "WE4 fails")]                                  # derived: WE8 still fails on the copy
    for f in fs:
        f.doc = "Lesson-Plan.docx"
    after = {"LP5": result("LP5", "needs_review", "The SLO section was not found"),
             "WE4": result("WE4", "pass", "A4"), "WE3": result("WE3", "needs_review", "Version v0.1 is present"),
             "WE8": result("WE8", "fail", "Body text is 10 pt"), "WE21": result("WE21", "pass", "fine"),
             **{c: result(c, "pass", "") for c in ("WE1", "WE2", "WE6", "WE7")}}
    kept = autofix.remaining(fs, "docx", after, {})
    assert [(f.code, f.status, f.message) for f in kept] == [
        ("LP5", "needs_review", "The SLO section was not found"),
        ("WE3", "needs_review", "Version v0.1 is present"),
        ("WE8", "fail", "Body text is 10 pt"),
        ("WE21", "fail", "Citation is not (Author, Year)"),
        ("LP10", "fail", "WE8: Body text is 10 pt")]                       # worked out again from the copy
    assert all(f.doc == "Lesson-Plan.docx" for f in kept)
    # a copy that could not be re-checked clears nothing
    assert autofix.remaining(fs, "docx", {}, {}) == fs


def fixed_lesson_plan(tmp_path, **kw):
    from docx import Document
    out = tmp_path / "o"
    cli.review(package(tmp_path, **kw), str(checklist(tmp_path)), str(out))
    z = zipfile.ZipFile(out / cli.MARKED_UP)
    name = next(n for n in z.namelist() if "Lesson-Plan" in n)
    z.extract(name, tmp_path / "x")
    d = Document(str(tmp_path / "x" / name))
    return name, d, "\n".join(c.text for c in d.comments)


def test_no_version_is_made_up(tmp_path):
    name, d, text = fixed_lesson_plan(tmp_path)
    assert name.endswith("/Lesson-Plan-Lesson-2-Chapter-4.docx")          # renamed, but no invented -v0.1
    assert "v0.1" not in d.sections[0].footer.paragraphs[0].text
    assert "WE3 FAIL" in text                  # the writer is still asked to add the version
    assert "WE4 FAIL" not in text and "WE8 FAIL" not in text


def test_version_is_taken_from_the_footer(tmp_path):
    name, d, text = fixed_lesson_plan(tmp_path, footer="Sabr lesson plan | Version 2.1 | Ali")
    assert name.endswith("/Lesson-Plan-Lesson-2-Chapter-4-v2.1.docx")
    assert "v2.1" in d.sections[0].footer.paragraphs[0].text
    assert "Version v2.1 taken from the footer" in text
    assert "WE3 CHECK" in text                 # a person still confirms it matches the document's stage


def test_quran_hadith_and_dua_text_is_never_edited(tmp_path):
    from docx import Document
    root = tmp_path / "Grade-6-Islamiat"
    ch = root / "Chapter-1-Quran"
    ch.mkdir(parents=True)
    d = new_doc(size=11)
    add(d, "باب اول: قرآن مجید", bold=True, rtl=True)
    add(d, "سورۃ العصر", bold=True, rtl=True)
    verse = "وَالْعَصْرِ ﴿١﴾ إِنَّ الْإِنسَانَ لَفِي خُسْرٍ ﴿٢﴾"
    add(d, "بِسْمِ اللَّهِ الرَّحْمَٰنِ الرَّحِيمِ", bold=True, rtl=True, urdu_font="Traditional Arabic")
    add(d, verse, rtl=True, urdu_font="Traditional Arabic")
    add(d, "ترجمہ: زمانے کی قسم ﴿١﴾ آیت ٢ میں", rtl=True)
    d.save(str(ch / "Chapter-1-Lesson-Plan.docx"))
    out = tmp_path / "out"
    cli.review(zip_dir(str(root), str(tmp_path / "p.zip")), str(checklist(tmp_path)), str(out))
    z = zipfile.ZipFile(out / cli.MARKED_UP)
    name = next(n for n in z.namelist() if n.endswith(".docx"))
    (tmp_path / "f.docx").write_bytes(z.read(name))
    texts = [p.text for p in Document(str(tmp_path / "f.docx")).paragraphs]
    assert "بِسْمِ اللَّهِ الرَّحْمَٰنِ الرَّحِيمِ" in texts            # not numbered as a heading
    assert verse in texts                                       # ayah numbers kept as written
    assert "ترجمہ: زمانے کی قسم ﴿١﴾ آیت 2 میں" in texts           # Urdu digit fixed, the verse mark kept
    fonts = {r.font.element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs")
             for p in Document(str(tmp_path / "f.docx")).paragraphs if p.text == verse for r in p.runs}
    assert fonts == {"Traditional Arabic"}


def test_headings_left_unnumbered_keep_their_numbering_comment(tmp_path):
    # A Lesson Plan without the five sections gets Heading styles but no numbers: WE12 / LP4 stay for a person.
    from docx import Document
    root = tmp_path / "Grade-6-Islamiat"
    lesson = root / "Chapter-1-Quran" / "Lesson-1-Asr"
    lesson.mkdir(parents=True)
    d = new_doc(size=11)
    add(d, "سبق: سورۃ العصر", bold=True, rtl=True)
    for h in ("تعارف", "حاصلات تعلیم", "ب - حفظ"):
        add(d, h, bold=True, size=14, rtl=True)
        add(d, URDU, rtl=True)
    d.save(str(lesson / "Lesson-1-Chapter-1-Lesson-Plan.docx"))
    out = tmp_path / "o"
    cli.review(zip_dir(str(root), str(tmp_path / "p.zip")), str(checklist(tmp_path)), str(out))
    z = zipfile.ZipFile(out / cli.MARKED_UP)
    (tmp_path / "f.docx").write_bytes(z.read(next(n for n in z.namelist() if n.endswith(".docx"))))
    text = "\n".join(c.text for c in Document(str(tmp_path / "f.docx")).comments)
    assert "not numbered" in text
    assert "WE12 FAIL" in text and "LP4 FAIL" in text
    assert "WE13 FAIL" not in text                       # the Heading styles were applied


def test_review_items_and_summary_items_follow_the_fixed_copy():
    from course_review import autofix
    from course_review.checks.common import result
    passes = lambda *codes: {c: result(c, "pass", "") for c in codes}
    # a needs-reviewer item the fix settled is left out (no Table of Contents -> one was added)
    assert autofix.remaining([result("WE15", "needs_review", "No page count")], "docx",
                             {"WE15": result("WE15", "pass", "Has a Table of Contents")}, {}) == []
    # FG8 on a Word Facilitator Guide clears even though WE5 (slides only) never runs on Word files
    assert autofix.remaining([result("FG8", "fail", "WE6: Poppins missing")], "docx", passes("WE6", "WE7", "WE23"), {}) == []
    # LP10 with no failing part but one needing a reviewer stays as needs-reviewer, as the engine would say
    after = {**passes("WE1", "WE2", "WE4", "WE6", "WE7", "WE8"), "WE3": result("WE3", "needs_review", "Stage?")}
    kept = autofix.remaining([result("LP10", "fail", "WE4: Letter size")], "docx", after, {})
    assert [(f.code, f.status) for f in kept] == [("LP10", "needs_review")] and "WE3" in kept[0].message


def test_version_patterns_do_not_invent_or_lose_versions():
    from course_review import autofix
    from course_review.models import DocRef

    class Info:
        footer_text, core_version, paras = "", "", []

    def version(name="x.docx", footer="", props=""):
        info = Info()
        info.footer_text, info.core_version = footer, props
        return autofix.find_version(DocRef(rel=name, abs="", ext="docx", doc_type="lesson_plan", chapter=1, lesson=1), info)[0]
    assert version("Lesson-Plan-v2-Final.docx") == "2"                     # not only at the end of the name
    assert version("LessonPlan v1.0 (1).docx") == "1.0"
    assert version(footer="Class V 2024") is None                          # a class and a year, not a version
    assert version(footer="Class V 2") is None
    assert version(footer="Grade 6 v2024") is None
    assert version(footer="Sabr | Version 2.1 | Ali") == "2.1"
    assert version(footer="ورژن ٣") == "3"
    assert version(props="1.4") == "1.4"


def test_versions_come_only_from_real_version_labels():
    from types import SimpleNamespace as NS
    from course_review.autofix import find_version
    doc = NS(rel="p/Lesson-Plan-Lesson-1-Chapter-2.docx", ext="docx")

    def info(footer="", lines=(), core=""):
        paras = [NS(text=t, in_table=False) for t in lines]
        return NS(footer_text=footer, paras=paras, core_version=core)
    assert find_version(doc, info(lines=["Read the Urdu version 2 of the story."])) == (None, None)
    assert find_version(doc, info(lines=["Version 1.2"])) == ("1.2", "the document text")
    assert find_version(doc, info(footer="Lesson-Plan-Lesson-1-Chapter-2 | v0.1 | 2026-10-06 | Page 1")) == (None, None)
    assert find_version(doc, info(footer="Lesson-Plan | v0.3 | 2026-10-06 | Page 1")) == ("0.3", "the footer")
    assert find_version(doc, info(footer="Class V 2024")) == (None, None)


def test_font_size_check_leaves_quoted_scripture_alone(tmp_path):
    from course_review.checks import formatting as F
    from course_review.docx_model import parse_docx
    from helpers import make_ctx, ref, save
    d = new_doc(size=12)
    add(d, "Title", size=20, bold=True)
    add(d, "Body text in the right size.", size=12)
    add(d, "بِسْمِ اللَّهِ الرَّحْمَٰنِ الرَّحِيمِ", size=16, rtl=True)      # quoted as written, in its own size
    assert F.we8(make_ctx(), ref(), parse_docx(save(d, tmp_path))).status == "pass"
