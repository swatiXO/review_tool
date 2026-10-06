"""SLO extraction from the Document of Specifications, tagging, and the coverage rules."""
import os

import docx
import pytest

from course_review import cli
from course_review.models import FAIL, NA, PASS, REVIEW
from course_review.slo import Slo, load_slos, tags_in

from helpers import add, checklist, new_doc, zip_dir

SLO_TEXT = ["Understand the meaning of a square root.", "Find the square root of a perfect square.", "Apply roots to real problems."]
HEADER = ["Lesson #", "Lesson", "SLOs", "Bloom's Taxonomy Level", "Corresponding Content Coverage"]


def spec_doc(path, chapters=((1, [("01", "Roots", SLO_TEXT)]),)):
    d = docx.Document()
    d.add_paragraph("Part A. Lesson Breakdown")
    t = d.add_table(rows=2, cols=4)
    for c, h in zip(t.rows[0].cells, ["Lesson #", "Lesson Name", "Lesson Breakdown", "Time"]):
        c.text = h                              # a table that is not an SLO table
    for ch, lessons in chapters:
        d.add_paragraph(f"Chapter {ch}: Something")
        t = d.add_table(rows=1, cols=5)
        for c, h in zip(t.rows[0].cells, HEADER):
            c.text = h
        for no, name, slos in lessons:
            for s in slos:
                cells = t.add_row().cells
                for c, v in zip(cells, [no, name, s, "Understand", "Covered in lesson plan"]):
                    c.text = v
    d.save(path)


def test_load_slos_reads_tables_per_chapter_and_skips_other_tables(tmp_path):
    p = str(tmp_path / "spec.docx")
    spec_doc(p, chapters=((1, [("01", "Roots", SLO_TEXT)]), (2, [("01", "Cubes", SLO_TEXT[:2]), ("02", "More", SLO_TEXT[:1])])))
    slos = load_slos(p)
    assert sorted(slos) == [(1, 1), (2, 1), (2, 2)]
    assert [s.index for s in slos[(1, 1)]] == [1, 2, 3] and slos[(1, 1)][0].text == SLO_TEXT[0]
    assert len(slos[(2, 1)]) == 2 and slos[(2, 2)][0].lesson_name == "More"


def make_slos():
    return [Slo(1, 1, "Roots", i, t) for i, t in enumerate(SLO_TEXT, 1)]


def test_tags_from_markers_and_quoted_text():
    slos = make_slos()
    assert tags_in("Question 1: what is a root? SLO 2", slos) == {2}
    assert tags_in("Find the square root of a perfect square - then explain it", slos) == {2}      # quoted SLO text
    assert tags_in("SLO-1 and SLO 3", slos) == {1, 3}
    assert tags_in("SLO 9", slos) == set()                                                          # no such SLO
    assert tags_in("an untagged question", slos) == set()


def test_chapter_level_tags_need_the_lesson():
    chapter = make_slos() + [Slo(1, 2, "More", 1, "A different outcome for the second lesson.")]
    assert tags_in("L2 SLO 1: question", None, chapter) == {(2, 1)}
    assert tags_in("A different outcome for the second lesson.", None, chapter) == {(2, 1)}
    assert tags_in("SLO 1", None, chapter) == set()          # ambiguous across lessons, so not a tag


# ------------------------------------------------------------------ coverage
def package(tmp_path, question_tags, slos=SLO_TEXT):
    root = tmp_path / "Grade-6-Test"
    lesson = root / "Chapter-1-Intro" / "Lesson-1-Roots"
    lesson.mkdir(parents=True)
    d = new_doc()
    add(d, "Introduction:")
    d.save(str(lesson / "Lesson-1-Chapter-1-Lesson-Plan.docx"))
    (root / "Spec").mkdir()
    spec_doc(str(root / "Spec" / "Document-of-Specifications.docx"), chapters=((1, [("01", "Roots", slos)]),))
    q = new_doc()
    add(q, "Lesson 1: Roots")
    for i, tag in enumerate(question_tags, 1):
        add(q, f"Question {i}: pick one {tag}")
        for opt in ("A) a", "B) b", "C) c"):
            add(q, opt)
    (root / "Pop Quiz").mkdir()
    q.save(str(root / "Pop Quiz" / "Pop Quiz.docx"))
    return zip_dir(str(root), str(tmp_path / "pkg.zip"))


def findings(tmp_path, tags, slos=SLO_TEXT):
    zp = package(tmp_path, tags, slos)
    _, res, _, _ = cli.review(zp, checklist(tmp_path), str(tmp_path / "out"))
    by = {}
    for f in res.findings:
        by.setdefault(f.code, []).append(f)
    return by, res


def test_all_slos_tagged_gives_pass_and_no_gaps(tmp_path):
    by, res = findings(tmp_path, ["SLO 1", "SLO 2"], slos=SLO_TEXT[:2])
    assert by["PQ5"][0].status == PASS
    assert by["ST1"][0].status == PASS and by["ST2"][0].status == PASS
    assert all(r["pop_quiz"] >= 1 for r in res.slo_map)


def test_an_slo_with_no_question_fails_when_every_question_is_tagged(tmp_path):
    by, res = findings(tmp_path, ["SLO 1", "SLO 2"])          # three SLOs, two questions
    f = by["PQ5"][0]
    assert f.status == FAIL and "SLO 3" in f.message
    assert by["ST2"][0].status == FAIL and "1 of 3" in by["ST2"][0].message


def test_untagged_questions_are_reviewed_not_failed_at_lesson_level(tmp_path):
    by, res = findings(tmp_path, ["", ""])
    assert by["PQ5"][0].status == REVIEW and "cannot be decided by code" in by["PQ5"][0].message
    st2 = by["ST2"][0]
    assert st2.status == FAIL and "No question in any assessment carries an SLO tag" in st2.message   # the literal rule


def test_some_untagged_questions_make_a_missing_slo_a_review(tmp_path):
    by, _ = findings(tmp_path, ["SLO 1", ""])
    assert by["PQ5"][0].status == REVIEW and "untagged" in by["PQ5"][0].message


def test_a_package_without_a_specification_says_so(tmp_path):
    root = tmp_path / "Pkg"
    (root / "Chapter-1-A" / "Lesson-1-B").mkdir(parents=True)
    new_doc().save(str(root / "Chapter-1-A" / "Lesson-1-B" / "Lesson-1-Chapter-1-Lesson-Plan.docx"))
    _, res, _, _ = cli.review(zip_dir(str(root), str(tmp_path / "p.zip")), checklist(tmp_path), str(tmp_path / "o"))
    st1 = next(f for f in res.findings if f.code == "ST1")
    assert st1.status == REVIEW and "Document of Specifications" in st1.message


def test_coverage_sheet_and_json_are_written(tmp_path):
    import json
    import openpyxl
    zp = package(tmp_path, ["SLO 1", "SLO 2"], slos=SLO_TEXT[:2])
    _, res, xlsx, _ = cli.review(zp, checklist(tmp_path), str(tmp_path / "out"))
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["SLO Coverage"]
    assert ws["C2"].value == 1 and ws["F2"].value == 1 and ws["J2"].value == 1
    data = json.load(open(tmp_path / "out" / "review.json", encoding="utf8"))
    assert len(data["slo_coverage"]) == 2
