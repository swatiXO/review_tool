"""Model fallback, tested offline with stand-in models and a fake Ollama server."""
import json
import os

import openpyxl
import pytest

from course_review import cli, fallback, llm
from course_review.checks import assessments as A
from course_review.docx_model import parse_docx
from course_review.fallback import Cache, ModelConfig, extract, recover_questions, verify
from course_review.models import FAIL, PASS, REVIEW
from course_review.questions import numbering_is_regular, parse_questions

from fakes import (BadShapeModel, DownModel, HallucinatingModel, RomanModel, start_fake_server)
from helpers import PROFILE, add, checklist, make_ctx, new_doc, ref, save, zip_dir

ROMAN_LINES = [f"{r}. What is the capital of country {n}?" for n, r in enumerate(("i", "ii", "iii", "iv", "v", "vi"), 1)]


def roman_doc(tmp_path, lines=ROMAN_LINES, extra=()):
    d = new_doc()
    add(d, "Chapter exam")
    for l in list(lines) + list(extra):
        add(d, l)
    return parse_docx(save(d, tmp_path))


def ctx_with(model):
    ctx = make_ctx()
    ctx.model = model
    return ctx


def cfg(client, mode="suggest", cache=None):
    return ModelConfig(client=client, cache=cache, mode=mode)


# ---------------------------------------------------------------- verification
def test_verify_accepts_real_proposals_and_rejects_invented_ones(tmp_path):
    info = roman_doc(tmp_path)
    chunk = fallback.listing(info)
    good = {"para": chunk[1][0], "quote": chunk[1][1][:10]}
    accepted, rejected = verify([good, {"para": 9999, "quote": "x" * 8}, {"para": chunk[2][0], "quote": "not in it at all"},
                                 {"para": "abc", "quote": "q"}, {"quote": "no para"}], chunk)
    assert list(accepted) == [chunk[1][0]] and rejected == 4


def test_verify_never_accepts_an_answer_option_as_a_question(tmp_path):
    d = new_doc()
    add(d, "Question text here")
    add(d, "A) an option")
    info = parse_docx(save(d, tmp_path))
    chunk = fallback.listing(info)
    accepted, rejected = verify([{"para": chunk[1][0], "quote": "A) an opt"}], chunk)
    assert not accepted and rejected == 1


def test_quote_comparison_ignores_arabic_letter_variants(tmp_path):
    d = new_doc()
    add(d, "تعليم مثال")        # Arabic yeh
    info = parse_docx(save(d, tmp_path))
    chunk = fallback.listing(info)
    accepted, _ = verify([{"para": chunk[0][0], "quote": "تعلیم"}], chunk)  # Urdu yeh
    assert accepted


# ------------------------------------------------------------------ extraction
def test_extract_returns_verified_starts(tmp_path):
    info = roman_doc(tmp_path)
    ex = extract(info, RomanModel())
    assert ex.usable and len(ex.starts) == 6 and ex.rejected == 0


def test_extraction_with_too_many_invented_proposals_is_discarded(tmp_path):
    ex = extract(roman_doc(tmp_path), HallucinatingModel())
    assert not ex.usable and "failed verification" in ex.note


def test_server_down_and_wrong_shape_are_reported_not_raised(tmp_path):
    info = roman_doc(tmp_path)
    assert "model unavailable" in extract(info, DownModel()).note
    assert "questions" in extract(info, BadShapeModel()).note


def test_long_documents_are_sent_in_chunks_and_merged(tmp_path):
    lines = [f"{'i' if n == 1 else 'ii'}. Question number {n} " + "word " * 40 for n in range(1, 41)]
    info = roman_doc(tmp_path, lines=lines)
    model = RomanModel()
    ex = extract(info, model)
    assert model.calls > 1 and ex.chunks == model.calls and len(ex.starts) == 40


def test_answers_are_cached_per_file_and_model(tmp_path):
    info = roman_doc(tmp_path)
    cache = Cache(tmp_path / "cache")
    first = RomanModel()
    a = extract(info, first, cache, doc_hash="abc")
    second = RomanModel()
    b = extract(info, second, cache, doc_hash="abc")
    assert first.calls == 1 and second.calls == 0 and b.cached and b.starts == a.starts
    other = RomanModel()
    other.name = "another-model"
    extract(info, other, cache, doc_hash="abc")
    assert other.calls == 1                          # a different model never reuses the answer


def test_recover_questions_builds_the_same_objects_as_the_rules(tmp_path):
    info = roman_doc(tmp_path)
    qs, summary = recover_questions(info, cfg(RomanModel()), ref())
    assert [q.num for q in qs] == [1, 2, 3, 4, 5, 6] and numbering_is_regular(qs)
    assert all(q.source == "model" and q.meta["recognised"] == 6 for q in qs)
    assert qs[0].block.startswith("i. What")
    assert recover_questions(info, None, ref())[0] is None


# ------------------------------------------------------------- policy in checks
def exam(ctx, info):
    return {f.code: f for f in A.exam_checks(ctx, ref(doc_type="chapter_exam"), info, "chapter_exam")}


def test_without_a_model_an_unknown_format_is_needs_review(tmp_path):
    res = exam(ctx_with(None), roman_doc(tmp_path))
    assert res["CE1"].status == REVIEW and "not be supported" in res["CE1"].message


def test_suggest_mode_never_writes_a_pass_or_fail(tmp_path):
    res = exam(ctx_with(cfg(RomanModel(), "suggest")), roman_doc(tmp_path))
    f = res["CE1"]
    assert f.status == REVIEW and f.method == "model" and "[model-assisted]" in f.message and "Would be pass" in f.message
    short = exam(ctx_with(cfg(RomanModel(), "suggest")), roman_doc(tmp_path, lines=ROMAN_LINES[:3]))
    assert short["CE1"].status == REVIEW and "Would be fail" in short["CE1"].message


def test_decide_mode_writes_fails_but_a_pass_stays_partial(tmp_path):
    ctx = ctx_with(cfg(RomanModel(), "decide"))
    ok = exam(ctx, roman_doc(tmp_path))["CE1"]
    assert ok.status == PASS and ok.partial and ok.method == "model"
    bad = exam(ctx, roman_doc(tmp_path, lines=ROMAN_LINES[:3]))["CE1"]
    assert bad.status == FAIL and "[model-assisted]" in bad.message


def test_when_the_model_cannot_help_the_reason_is_in_the_evidence(tmp_path):
    res = exam(ctx_with(cfg(DownModel())), roman_doc(tmp_path))
    f = res["CE1"]
    assert f.status == REVIEW and any("Model fallback did not help" in e for e in f.evidence)


def test_the_model_is_not_asked_when_the_rules_already_understand_the_document(tmp_path):
    d = new_doc()
    for i in range(1, 7):
        add(d, f"Question {i}: text")
    model = RomanModel()
    exam(ctx_with(cfg(model)), parse_docx(save(d, tmp_path)))
    assert model.calls == 0


# ----------------------------------------------------------------- end to end
def package_with_roman_pop_quiz(tmp_path):
    root = tmp_path / "Grade-6-Test"
    lesson = root / "Chapter-1-Intro" / "Lesson-1-Roots"
    lesson.mkdir(parents=True)
    d = new_doc()
    for s in ("Introduction", "SLOs", "Warm-up", "Concept Building", "Key Takeaways"):
        add(d, s + ":", bold=True)
    d.save(str(lesson / "Lesson-1-Chapter-1-Lesson-Plan.docx"))
    q = new_doc()
    add(q, "Lesson 1: Roots")
    for roman in ("i. Pick one", "ii. Pick again"):
        add(q, roman)
        for opt in ("A) a", "B) b", "C) c"):
            add(q, opt)
    (root / "Pop Quiz").mkdir()
    q.save(str(root / "Pop Quiz" / "Pop Quiz.docx"))
    return zip_dir(str(root), str(tmp_path / "pkg.zip"))


def test_end_to_end_with_model_fallback(tmp_path):
    zip_path = package_with_roman_pop_quiz(tmp_path)
    model = cfg(RomanModel(), "suggest", Cache(tmp_path / "cache"))
    pkg, res, xlsx, _ = cli.review(zip_path, checklist(tmp_path), str(tmp_path / "out"), model=model)
    pq1 = next(f for f in res.findings if f.code == "PQ1" and f.lesson is not None)
    assert pq1.status == REVIEW and pq1.method == "model" and "Would be pass" in pq1.message
    wb = openpyxl.load_workbook(xlsx)
    assert wb["Pop Quiz"]["C6"].value is None                       # a suggestion is never a Pass in the workbook
    assert "[model-assisted]" in wb["Pop Quiz"]["I6"].value if wb["Pop Quiz"]["I6"].value else True
    summary = {r[0].value: r[1].value for r in wb["Review Summary"].iter_rows(min_row=2, max_row=12) if r[0].value}
    assert "suggest mode" in summary["Model fallback"]
    data = json.load(open(tmp_path / "out" / "review.json", encoding="utf8"))
    assert data["model_calls"] and data["model_calls"][0]["usable"]


def test_end_to_end_without_model_is_unchanged_and_says_so(tmp_path):
    zip_path = package_with_roman_pop_quiz(tmp_path)
    pkg, res, xlsx, _ = cli.review(zip_path, checklist(tmp_path), str(tmp_path / "out"))
    pq1 = next(f for f in res.findings if f.code == "PQ1" and f.lesson is not None)
    assert pq1.status == REVIEW and pq1.method == "deterministic"
    wb = openpyxl.load_workbook(xlsx)
    summary = {r[0].value: r[1].value for r in wb["Review Summary"].iter_rows(min_row=2, max_row=12) if r[0].value}
    assert summary["Model fallback"].startswith("Off")


# ------------------------------------------------------------- the HTTP client
@pytest.fixture
def server():
    s, url = start_fake_server()
    yield s, url
    s.shutdown()


def test_client_available_checks_the_server_and_the_model(server):
    _, url = server
    ok, msg = llm.OllamaClient(url=url, model="qwen3:14b").available()
    assert ok and "qwen3:14b" in msg
    ok, msg = llm.OllamaClient(url=url, model="llama3.1:8b").available()
    assert not ok and "not installed" in msg


def test_client_reports_an_unreachable_server_and_an_offline_ngrok_page(server):
    ok, msg = llm.OllamaClient(url="http://127.0.0.1:1", model="m").available()
    assert not ok and "could not reach" in msg
    from fakes import FakeOllamaHandler
    _, url = server
    FakeOllamaHandler.html_page = True
    ok, msg = llm.OllamaClient(url=url, model="qwen3:14b").available()
    assert not ok and "ngrok tunnel" in msg


def test_client_sends_the_ngrok_header_and_deterministic_options(server):
    from fakes import FakeOllamaHandler
    _, url = server
    FakeOllamaHandler.chat_reply = "```json\n{\"questions\": [{\"para\": 1, \"quote\": \"abc\"}]}\n```"
    client = llm.OllamaClient(url=url + "/", model="qwen3:14b")
    reply = client.chat_json("system text", "[1] abc def")
    assert reply == {"questions": [{"para": 1, "quote": "abc"}]}        # code fences stripped
    method, path, headers, payload = FakeOllamaHandler.seen[-1]
    assert path == "/api/chat" and {k.lower(): v for k, v in headers.items()}.get("ngrok-skip-browser-warning") == "true"
    assert payload["format"] == "json" and payload["stream"] is False and payload["think"] is False
    assert payload["options"]["temperature"] == 0 and "seed" in payload["options"]
    assert payload["messages"][0]["role"] == "system" and payload["messages"][1]["content"] == "[1] abc def"
    assert client.calls == 1


def test_client_rejects_a_reply_that_is_not_json(server):
    from fakes import FakeOllamaHandler
    _, url = server
    FakeOllamaHandler.chat_reply = "Sure! There are six questions."
    with pytest.raises(llm.LLMError):
        llm.OllamaClient(url=url, model="qwen3:14b").chat_json("s", "u")


def test_client_url_comes_from_the_environment(monkeypatch):
    monkeypatch.delenv("OLLAMA_HOST", raising=False)
    monkeypatch.setenv("OLLAMA_URL", "https://abcd.ngrok-free.dev/")
    monkeypatch.setenv("OLLAMA_MODEL", "qwen2.5:14b")
    c = llm.OllamaClient()
    assert c.url == "https://abcd.ngrok-free.dev" and c.model == "qwen2.5:14b"
    monkeypatch.delenv("OLLAMA_URL")
    monkeypatch.setenv("OLLAMA_HOST", "x.ngrok-free.dev")
    assert llm.OllamaClient().url == "https://x.ngrok-free.dev"


def test_make_model_refuses_to_run_silently_without_a_server():
    with pytest.raises(SystemExit) as e:
        cli.make_model("suggest", url="http://127.0.0.1:1", model="m")
    assert "not usable" in str(e.value)
    assert cli.make_model("off") is None


def test_chunks_overlap_so_a_boundary_question_is_seen_whole():
    items = [(i, "word " * 60) for i in range(30)]
    chunks = fallback.chunks_of(items, max_chars=1500)
    assert len(chunks) > 2
    for prev, nxt in zip(chunks, chunks[1:]):
        assert nxt[:2] == prev[-fallback.CHUNK_OVERLAP:] or nxt[0] in prev       # the next chunk starts with the previous chunk's tail
    assert {i for c in chunks for i, _ in c} == set(range(30))                  # nothing is dropped


def test_a_quote_from_the_line_after_a_label_is_accepted_but_not_from_further_away(tmp_path):
    d = new_doc()
    add(d, "Question 3 (analysis)")                       # label line
    add(d, "Explain why the Muslims won the battle of Badr in detail.")   # the wording
    add(d, "Some unrelated paragraph that is far away from the label.")
    info = parse_docx(save(d, tmp_path))
    chunk = fallback.listing(info)
    label = chunk[0][0]
    ok, bad = verify([{"para": label, "quote": "Explain why the Muslims won"}], chunk)
    assert list(ok) == [label] and bad == 0
    ok, bad = verify([{"para": label, "quote": "unrelated paragraph that is far"}], chunk)
    assert not ok and bad == 1
