"""Model-assisted judgement checks, SLO mapping and the book index, with stand-in models."""
import json
import os
import re

import openpyxl
import pytest
from pptx import Presentation
from pptx.util import Inches

from course_review import cli
from course_review.book import BookIndex, parse_pages, tokens
from course_review.fallback import Cache, ModelConfig
from course_review.judge import Material, judge, suggestion, verify_quotes
from course_review.llm import LLMError
from course_review.models import FAIL, PASS, REVIEW

from fakes import _Base
from helpers import add, checklist, make_ctx, new_doc, zip_dir
from test_coverage import SLO_TEXT, spec_doc


class FakeJudge(_Base):
    """Answers every judgement with a fixed verdict and quotes a real line of the material."""

    def __init__(self, verdict="pass", reason="looks fine", extra=None, quote=True):
        super().__init__()
        self.verdict, self.reason, self.extra, self.quote = verdict, reason, extra or {}, quote

    def chat_json(self, system, user):
        self.calls += 1
        self.prompts.append(user)
        if system.startswith("You match exam questions"):
            return self.map_reply(user)
        body = user.split("### ", 1)[1].split("\n", 1)[1] if "### " in user else user
        line = next((l for l in body.splitlines() if len(l.strip()) >= 12), "")
        ev = [{"quote": line.strip()[:40]}] if self.quote and line else []
        return {"verdict": self.verdict, "reason": self.reason, "evidence": ev, **self.extra}

    def map_reply(self, user):
        qs = re.findall(r"^(\d+)\. (.*)$", user.split("Questions:\n", 1)[1], re.M)
        return {"mapping": [{"question": int(n), "slos": [int(n)], "quote": t.strip()[:20]} for n, t in qs]}


def cfg(client, judge_on=True, codes=None, cache=None, mode="suggest"):
    return ModelConfig(client=client, cache=cache, mode=mode, judge=judge_on, codes=codes)


MATS = [Material("Concept Building", "A square root undoes squaring, so the square root of 49 is 7.")]


# -------------------------------------------------------------------- the judge
def test_quotes_must_appear_in_the_material():
    ok, bad = verify_quotes([{"quote": "the square root of 49 is 7"}, {"quote": "invented words here"}, {"quote": "short"}], MATS)
    assert ok == ["the square root of 49 is 7"] and bad == 2


def test_a_supported_verdict_is_usable():
    j = judge(cfg(FakeJudge("pass")), "LP6", "Warm-up is concise", MATS, "decide")
    assert j.usable and j.verdict == "pass" and j.quotes


def test_a_verdict_without_a_verifiable_quote_is_discarded():
    j = judge(cfg(FakeJudge("fail", quote=False)), "LP6", "rule", MATS, "decide")
    assert not j.usable and j.verdict == "unclear" and "no quote" in j.note


def test_unclear_and_malformed_and_unreachable_are_not_usable():
    assert not judge(cfg(FakeJudge("unclear")), "X", "r", MATS, "d").usable

    class Odd(_Base):
        def chat_json(self, s, u):
            return {"answer": "yes"}

    class Down(_Base):
        def chat_json(self, s, u):
            raise LLMError("tunnel down")
    assert "expected form" in judge(cfg(Odd()), "X", "r", MATS, "d").note
    assert "model unavailable" in judge(cfg(Down()), "X", "r", MATS, "d").note


def test_extra_answers_are_validated_against_the_material():
    ok = judge(cfg(FakeJudge("fail", extra={"uncovered_slos": [2]})), "LP1", "r", MATS, "d", extra_keys=("uncovered_slos",),
               validate_extra=lambda e: all(1 <= n <= 3 for n in e["uncovered_slos"]))
    assert ok.usable and ok.extra["uncovered_slos"] == [2]
    bad = judge(cfg(FakeJudge("fail", extra={"uncovered_slos": [9]})), "LP1", "r", MATS, "d", extra_keys=("uncovered_slos",),
                validate_extra=lambda e: all(1 <= n <= 3 for n in e["uncovered_slos"]))
    assert not bad.usable and "inconsistent" in bad.note


def test_answers_are_cached_and_long_material_is_flagged_as_cut(tmp_path):
    cache = Cache(tmp_path / "c")
    first, second = FakeJudge(), FakeJudge()
    judge(cfg(first, cache=cache), "LP6", "r", MATS, "d")
    again = judge(cfg(second, cache=cache), "LP6", "r", MATS, "d")
    assert first.calls == 1 and second.calls == 0 and again.cached and again.usable
    long = [Material("Concept Building", ("A long paragraph about roots and powers. " * 400))]
    j = judge(cfg(FakeJudge()), "LP6", "r", long, "d")
    assert j.cut and "shortened" in " ".join(suggestion("LP6", j).evidence)


def test_suggestion_is_always_needs_review():
    f = suggestion("LP6", judge(cfg(FakeJudge("pass", "concise and previews")), "LP6", "r", MATS, "d"), lesson=None)
    assert f.status == REVIEW and f.method == "model" and "Would be pass: concise and previews" in f.message and f.evidence[0].startswith("Quote:")
    g = suggestion("LP6", judge(cfg(FakeJudge("unclear")), "LP6", "r", MATS, "d"))
    assert g.status == REVIEW and "could not decide" in g.message


# ------------------------------------------------------------ the checks, end to end
def package(tmp_path, tagged=False, with_guide=True):
    root = tmp_path / "Grade-6-Test"
    lesson = root / "Chapter-1-Intro" / "Lesson-1-Roots"
    lesson.mkdir(parents=True)
    d = new_doc()
    for name, body in (("Introduction", "Hamza looks at the stars with his grandfather and wonders which one never moves."),
                       ("SLOs", None), ("Warm-up", "List three things you trust when you are lost and why you trust them."),
                       ("Concept Building", "A square root undoes squaring. The square root of 49 is 7, like the fixed star that guides travellers."),
                       ("Key Takeaways", "Roots undo powers, and Hamza's star shows what steady guidance means.")):
        add(d, name + ":", bold=True)
        if body:
            add(d, body)
    d.save(str(lesson / "Lesson-1-Chapter-1-Lesson-Plan.docx"))
    if with_guide:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for title, text in (("Cover", "Lesson guide"), ("SLOs", "Say these aloud: understand roots."),
                            ("Warm-up", "Ask the class who they trust. Then debrief."), ("Key Takeaways", "Roots undo powers. Ask: what did you learn?")):
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = title
            box = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(8), Inches(1))
            box.text_frame.text = text
        prs.save(str(lesson / "Lesson-1-Chapter-1-Facilitator-Guide.pptx"))
    (root / "Spec").mkdir()
    spec_doc(str(root / "Spec" / "Document-of-Specifications.docx"))
    q = new_doc()
    add(q, "Lesson 1: Roots")
    for i in (1, 2):
        add(q, f"Question {i}: Which statement about roots number {i} is correct?{' SLO ' + str(i) if tagged else ''}")
        for opt in ("A) a", "B) b", "C) c"):
            add(q, opt)
    (root / "Pop Quiz").mkdir()
    q.save(str(root / "Pop Quiz" / "Pop Quiz.docx"))
    return zip_dir(str(root), str(tmp_path / "pkg.zip"))


def review(tmp_path, model, book=None, extra_lp=(), **kw):
    zp = package(tmp_path, **kw)
    return cli.review(zp, checklist(tmp_path, extra_lp), str(tmp_path / "out"), model=model, book=book)


def codes_of(res, method="model"):
    return sorted({f.code for f in res.findings if f.method == method})


def test_without_model_checks_nothing_model_assisted_appears(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _, res, _, _ = review(tmp_path / "a", None)
    assert codes_of(res) == []
    _, res, _, _ = review(tmp_path / "b", cfg(FakeJudge(), judge_on=False))
    assert codes_of(res) == []


def test_all_lesson_plan_and_guide_checks_run_and_are_suggestions(tmp_path):
    model = cfg(FakeJudge("pass"))
    _, res, xlsx, _ = review(tmp_path, model, extra_lp=("LP1", "LP6"))
    got = codes_of(res)
    for code in ("LP1", "LP2", "LP6", "LP7", "LP9", "FG1", "FG3", "FG6"):
        assert code in got, f"{code} missing from {got}"
    assert "LP8" not in got and "CE2" not in got             # these need the textbook
    assert all(f.status == REVIEW for f in res.findings if f.method == "model")
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Lesson Plan"]
    assert ws["C6"].value == "Pass"                            # LP3 is decided by code
    assert ws["G6"].value is None and ws["H6"].value is None   # LP1, LP6: a model suggestion is never a Pass or Fail
    assert "[model-assisted]" in ws["J6"].value
    assert "model" in " ".join(str(r[1].value) for r in wb["Review Summary"].iter_rows(min_row=2, max_row=12) if r[0].value).lower()


def test_only_the_requested_codes_run(tmp_path):
    _, res, _, _ = review(tmp_path, cfg(FakeJudge(), codes={"LP1", "FG1"}))
    assert codes_of(res) == ["FG1", "LP1"]


def test_lp1_reports_which_slos_the_model_found_uncovered(tmp_path):
    model = cfg(FakeJudge("fail", "SLO 3 is never taught", extra={"uncovered_slos": [3]}), codes={"LP1"})
    _, res, _, _ = review(tmp_path, model)
    f = next(f for f in res.findings if f.code == "LP1")
    assert "Would be fail" in f.message and any("Apply roots to real problems" in e for e in f.evidence)


def test_the_book_unlocks_lp8_ce2_and_is_used_by_lp2(tmp_path):
    book = BookIndex({1: "Hamza looks at the stars with his grandfather and wonders which one never moves", 2: "square roots perfect squares"})
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _, res, _, _ = review(tmp_path / "a", cfg(FakeJudge("pass")), book=book)
    got = codes_of(res)
    assert "LP8" in got
    lp2 = next(f for f in res.findings if f.code == "LP2")
    assert "textbook was not consulted" not in " ".join(lp2.evidence)
    _, res2, _, _ = review(tmp_path / "b", cfg(FakeJudge("pass")), book=None)
    assert "textbook was not consulted" in " ".join(next(f for f in res2.findings if f.code == "LP2").evidence)


# -------------------------------------------------------------- SLO mapping
def test_untagged_questions_get_a_model_suggested_slo_mapping(tmp_path):
    _, res, _, _ = review(tmp_path, cfg(FakeJudge("pass"), codes={"PQ5"}), with_guide=False)
    f = next(f for f in res.findings if f.code == "PQ5")
    assert f.status == REVIEW and f.method == "model"
    assert "Would be fail" in f.message and "SLO 3" in f.message         # questions 1, 2 map to SLO 1, 2; SLO 3 has none
    assert any(e.startswith("Q1 -> SLO [1]") for e in f.evidence)


def test_tagged_questions_are_decided_by_code_not_the_model(tmp_path):
    model = FakeJudge("pass")
    _, res, _, _ = review(tmp_path, cfg(model, codes={"PQ5"}), tagged=True, with_guide=False)
    f = next(f for f in res.findings if f.code == "PQ5")
    assert f.method == "deterministic" and f.status == FAIL and model.calls == 0


# --------------------------------------------------------------- book index
def test_book_search_ranks_the_relevant_page_first():
    book = BookIndex({1: "the fast brown fox jumps over the lazy dog", 2: "square roots and perfect squares explained with examples",
                      3: "unrelated chapter about travel and weather"})
    assert book.search("perfect square roots", k=2)[0][0] == 2
    assert book.passages("perfect squares", k=1).startswith("[book page 2]")
    assert book.search("zzz nothing") == []


def test_book_search_works_on_urdu_text():
    book = BookIndex({1: "غزوہ بدر میں مشرکین قتل ہوئے", 2: "نماز کی فرضیت اور اہمیت کے بارے میں", 3: "زکوٰۃ اور حج کے احکام"})
    assert book.search("غزوہ بدر", k=1)[0][0] == 1
    assert book.search("نماز اہمیت", k=1)[0][0] == 2


def test_page_specs_and_index_loading(tmp_path):
    assert parse_pages("1-3,7", 10) == [1, 2, 3, 7] and parse_pages(None, 3) == [1, 2, 3] and parse_pages("9-20", 10) == [9, 10]
    (tmp_path / "pages").mkdir()
    (tmp_path / "pages" / "0004.txt").write_text("square roots", encoding="utf8")
    (tmp_path / "pages" / "0005.txt").write_text("cube roots", encoding="utf8")
    assert sorted(BookIndex.load(tmp_path).pages) == [4, 5]
    assert tokens("Roots, عبادت!") == ["roots", "عبادت"]


# ------------------------------------------------- long text read in parts
class PartsJudge(_Base):
    """Says each part covers the SLO whose keyword appears in it, quoting that line."""
    num_ctx = 4096          # budget ~5000 characters, so the ~7600-character text needs two parts

    def chat_json(self, system, user):
        self.calls += 1
        part = user.split("### Concept Building", 1)[1].split("\n", 1)[1]
        covered = [n for n, w in ((1, "roots"), (2, "stems"), (3, "leaves")) if w in part]
        line = next(l for l in part.splitlines() if any(w in l for w in ("roots", "stems", "leaves", "filler")))
        return {"verdict": "pass" if covered else "unclear", "reason": "r", "evidence": [{"quote": line[:40]}],
                "covered_slos": covered}


def test_long_concept_building_is_read_in_parts_and_an_slo_counts_if_any_part_covers_it():
    from course_review.checks.judgement import lp1, split_parts
    slos = ["Name the roots", "Name the stems", "Name the leaves"]
    filler = "\n".join(["filler sentence about plants that goes on for a while."] * 70)
    cb = "Plant roots hold the soil.\n" + filler + "\nPlant stems carry water up.\n" + filler
    ctx = make_ctx()
    ctx.model = cfg(PartsJudge())
    assert len(split_parts(cb, 3000)) > 1
    f = lp1(ctx, None, None, {"concept_building": cb}, slos)
    assert ctx.model.client.calls > 1
    assert "Would be fail" in f.message and "3 not covered" in f.message      # leaves: in no part
    assert any("3 (Name the leaves)" in e for e in f.evidence)


def test_split_parts_keeps_all_text():
    from course_review.checks.judgement import split_parts
    text = "\n".join(f"paragraph {i} " + "x" * 300 for i in range(40)) + "\n" + "y" * 9000
    parts = split_parts(text, 2500)
    assert all(len(p) <= 2500 for p in parts)
    assert "".join(parts).replace("\n", "") == text.replace("\n", "")
