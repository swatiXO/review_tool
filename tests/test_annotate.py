"""The marked-up copy: highlights and comments written into the documents themselves."""
import zipfile

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from pptx import Presentation
from pptx.util import Inches

from course_review import annotate, cli
from course_review.checks.common import mark, result
from course_review.models import FAIL, PASS, REVIEW

from helpers import add, checklist, new_doc, save
from test_fallback import package_with_roman_pop_quiz


def colours(path):
    d = Document(path)
    return {p.text: {r.font.highlight_color for r in p.runs if r.text.strip()} for p in d.paragraphs}


def comment_texts(path):
    return [c.text for c in Document(path).comments]


def test_failing_heading_is_red_with_a_comment_naming_the_code(tmp_path):
    d = new_doc()
    add(d, "Introduction", bold=True)
    add(d, "Some body text that is fine.")
    src = save(d, tmp_path)
    dst = tmp_path / "out.docx"
    f = result("WE12", FAIL, "1 heading has no number", marks=[mark("Introduction", "heading has no decimal number")])
    assert annotate.annotate_docx(src, dst, [f]) == 1
    c = colours(dst)
    assert c["Introduction"] == {WD_COLOR_INDEX.RED}
    assert c["Some body text that is fine."] == {None}
    texts = comment_texts(dst)
    assert any(t.startswith("To fix: Number this heading") and t.endswith("[WE12]") for t in texts)
    assert any("Red = please fix" in t for t in texts)     # the colour key in the summary at the top


def test_model_suggestion_is_turquoise_and_pass_is_green(tmp_path):
    d = new_doc()
    add(d, "The warm-up asks students what they already know about roots.")
    add(d, "Key Takeaways")
    src = save(d, tmp_path)
    dst = tmp_path / "out.docx"
    fs = [result("LP6", REVIEW, "model", method="model",
                 marks=[mark("asks students what they already know", "model suggests PASS", exact=False)]),
          result("LP3", PASS, "ok", marks=[mark("Key Takeaways", "section present")])]
    annotate.annotate_docx(src, dst, fs)
    c = colours(dst)
    assert c["The warm-up asks students what they already know about roots."] == {WD_COLOR_INDEX.TURQUOISE}
    assert c["Key Takeaways"] == {WD_COLOR_INDEX.BRIGHT_GREEN}


def test_existing_yellow_highlight_on_a_correct_answer_is_kept(tmp_path):
    d = new_doc()
    add(d, "1. Pick one")
    add(d, "A) the right answer", highlight=WD_COLOR_INDEX.YELLOW)
    src = save(d, tmp_path)
    dst = tmp_path / "out.docx"
    f = result("PQ1", FAIL, "x", marks=[mark("A) the right answer", "test")])
    annotate.annotate_docx(src, dst, [f])
    assert colours(dst)["A) the right answer"] == {WD_COLOR_INDEX.YELLOW}


def test_repeated_text_marks_one_paragraph_per_mark_not_all_of_them(tmp_path):
    d = new_doc()
    for _ in range(4):
        add(d, "Question")
        add(d, "Answer text")
    src = save(d, tmp_path)
    dst = tmp_path / "out.docx"
    f = result("WE16", FAIL, "2 tables", marks=[mark("Question", "table 1"), mark("Question", "table 2")])
    assert annotate.annotate_docx(src, dst, [f]) == 2


def test_whole_document_problems_go_in_the_top_comment(tmp_path):
    d = new_doc()
    add(d, "First paragraph")
    src = save(d, tmp_path)
    dst = tmp_path / "out.docx"
    f = result("WE7", FAIL, "Urdu text is not Jamil Noori Nastaliq")
    annotate.annotate_docx(src, dst, [f])
    top = comment_texts(dst)[0]
    assert "To fix: Urdu text is not Jamil Noori Nastaliq. [WE7]" in top
    assert colours(dst)["First paragraph"] == {None}


def test_slides_get_a_highlight_and_a_review_box(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Debrief: pairs share answers"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Other slide"
    src, dst = tmp_path / "in.pptx", tmp_path / "out.pptx"
    prs.save(src)
    f = result("FG3", REVIEW, "m", marks=[mark("pairs share answers", "model suggests PASS", exact=False)])
    assert annotate.annotate_pptx(src, dst, [f]) == 1
    out = Presentation(dst)
    boxes = {i: sh.text_frame.text for i, sl in enumerate(out.slides, 1) for sh in sl.shapes if sh.name == "Course Review notes"}
    assert "AI check, looks fine" in boxes[1] and "[FG3]" in boxes[1]
    assert 2 not in boxes
    xml = out.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]._r.xml
    assert "highlight" in xml and "00FFFF" in xml


def test_review_writes_a_marked_up_zip_with_notes(tmp_path):
    out = tmp_path / "out"
    cli.review(str(package_with_roman_pop_quiz(tmp_path)), str(checklist(tmp_path)), str(out))
    z = zipfile.ZipFile(out / cli.MARKED_UP)
    names = z.namelist()
    assert any(n.endswith("REVIEW-NOTES.txt") for n in names)
    assert any(n.endswith("Lesson-Plan-Lesson-1-Chapter-1.docx") for n in names)      # renamed to the pattern
    assert not any(n.endswith("Lesson-1-Chapter-1-Lesson-Plan.docx") for n in names)
    lp = next(n for n in names if n.endswith("Lesson-Plan-Lesson-1-Chapter-1.docx"))
    (tmp_path / "lp.docx").write_bytes(z.read(lp))
    top = comment_texts(tmp_path / "lp.docx")[-1]
    assert "What we fixed for you" in top and "A4" in top


def test_every_turquoise_highlight_has_a_comment_that_says_what_to_do(tmp_path):
    d = new_doc()
    for t in ("Intro text about roots.", "Roots hold the plant.", "Roots take in water.", "Stems carry water.", "Leaves make food."):
        add(d, t)
    src, dst = save(d, tmp_path), tmp_path / "out.docx"
    suggested_pass = result("LP9", REVIEW, "[model-assisted] Would be pass: fine", method="model",
                            marks=[mark("Roots hold the plant.\nRoots take in water.", "model suggests PASS: fine", exact=False)])
    suggested_fail = result("LP7", REVIEW, "[model-assisted] Would be fail: story missing", method="model",
                            marks=[mark("Stems carry water.", "model suggests FAIL: no story here"),
                                   mark("Leaves make food.", "model suggests FAIL: no story here")])
    annotate.annotate_docx(src, dst, [suggested_pass, suggested_fail])
    c = colours(dst)
    turquoise = [t for t, cs in c.items() if WD_COLOR_INDEX.TURQUOISE in cs]
    assert turquoise == ["Roots hold the plant.", "Stems carry water.", "Leaves make food."]   # a pass: one place only
    texts = comment_texts(dst)
    assert any(t.startswith("AI check, looks fine") and "delete this note" in t for t in texts)
    assert any(t.startswith("AI check, may need fixing") and "tell the writer" in t for t in texts)
    assert sum(1 for t in texts if "[LP7]" in t) == 2                                          # each FAIL place commented


def test_a_re_uploaded_copy_loses_the_earlier_review_but_keeps_peoples_comments(tmp_path):
    from course_review import autofix
    d = new_doc()
    add(d, "Introduction", bold=True)
    add(d, "Body text of the lesson.")
    src = save(d, tmp_path)
    first = tmp_path / "first.docx"
    annotate.annotate_docx(src, first, [result("WE12", FAIL, "x", marks=[mark("Introduction", "no number")])])
    doc = Document(first)
    doc.add_comment(doc.paragraphs[1].runs, text="Please shorten this.", author="Ayesha", initials="A")
    doc.save(first)
    again = Document(first)
    assert autofix.drop_review_comments(again) == 2                        # the WE12 note and the summary
    again.save(tmp_path / "again.docx")
    left = Document(tmp_path / "again.docx")
    assert [c.author for c in left.comments] == ["Ayesha"]
    assert [p.text for p in left.paragraphs] == ["Introduction", "Body text of the lesson."]
