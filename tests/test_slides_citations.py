"""FG2 / FG7 on generated slide decks, WE21 / WE22 on generated documents."""
import os

import pytest
from docx.shared import Twips
from pptx import Presentation
from pptx.util import Inches

from course_review.checks import formatting as F
from course_review.checks import slides as S
from course_review.docx_model import parse_docx
from course_review.models import DocRef, FAIL, NA, PASS, REVIEW
from course_review.pptx_model import parse_pptx

from helpers import add, make_ctx, new_doc, ref, save


def deck(tmp_path, slides, name="g.pptx"):
    """slides: [(title, [text boxes], notes)]"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for title, texts, notes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        for i, t in enumerate(texts):
            box = s.shapes.add_textbox(Inches(0.5), Inches(1.5 + i), Inches(8), Inches(0.8))
            box.text_frame.text = t
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    path = os.path.join(str(tmp_path), name)
    prs.save(path)
    return parse_pptx(path)


FULL = [("Cover", ["Lesson"], ""), ("Introduction", ["a"], ""), ("SLOs", ["b"], ""), ("Warm-up", ["c"], ""),
        ("Concept Building", ["d"], ""), ("Key Takeaways", ["e"], "")]


def fg7(tmp_path, slides):
    return S.fg7(make_ctx(), ref(rel="p/Lesson-1-Guide.pptx", doc_type="facilitator_guide"), deck(tmp_path, slides))


def test_fg7_five_sections_in_order(tmp_path):
    r = fg7(tmp_path, FULL)
    assert r.status == PASS and r.partial


def test_fg7_missing_or_out_of_order(tmp_path):
    r = fg7(tmp_path, [s for s in FULL if s[0] != "Warm-up"])
    assert r.status == FAIL and "warm_up" in r.message
    swapped = [FULL[0], FULL[3], FULL[1], FULL[2], FULL[4], FULL[5]]      # warm-up before introduction
    assert fg7(tmp_path, swapped).status == FAIL


def test_fg7_allows_the_slo_slide_at_the_session_overview_but_not_after_concept_building(tmp_path):
    overview_first = [FULL[0], FULL[2], FULL[1], FULL[3], FULL[4], FULL[5]]
    assert fg7(tmp_path, overview_first).status == PASS
    late = [FULL[0], FULL[1], FULL[3], FULL[4], FULL[2], FULL[5]]
    r = fg7(tmp_path, late)
    assert r.status == FAIL and "slos" in r.message


def test_fg7_matches_urdu_titles_containing_the_section_name(tmp_path):
    slides = [("رہنمائے مقرر", ["x"], ""),
              ("حمزہ سے تعارف", ["x"], ""),
              ("حاصلاتِ تعليم (SLOs)", ["x"], ""),
              ("تحريکی سرگرمی: x", ["x"], ""),
              ("تصوراتی تعمیر: x", ["x"], ""),
              ("حاصل کلام اور اختتام", ["x"], "")]
    assert fg7(tmp_path, slides).status == PASS


def fg2(tmp_path, slides, plan_paragraphs=()):
    docs = []
    if plan_paragraphs:
        d = new_doc()
        for p in plan_paragraphs:
            add(d, p)
        path = save(d, tmp_path, "plan.docx")
        docs.append(DocRef(rel="p/Lesson-1-Lesson-Plan.docx", abs=path, ext="docx", doc_type="lesson_plan", chapter=1, lesson=1))
    ctx = make_ctx(docs=docs)
    return S.fg2(ctx, ref(rel="p/Lesson-1-Guide.pptx", doc_type="facilitator_guide"), deck(tmp_path, slides))


def test_fg2_facilitator_part_via_label_or_speaker_notes(tmp_path):
    slides = [("Cover", ["x"], ""), ("One", ["Content", "Facilitator Notes: read aloud"], ""), ("Two", ["Content"], "Say this slowly")]
    r = fg2(tmp_path, slides)
    assert r.status == PASS and r.partial


def test_fg2_unlabelled_notes_are_a_reviewer_question_not_a_fail(tmp_path):
    r = fg2(tmp_path, [("Cover", ["x"], ""), ("One", ["Content only"], "")])
    assert r.status == REVIEW and "mark facilitator notes another way" in r.message


def test_fg2_copied_lesson_plan_text_fails(tmp_path):
    para = "The square root of a number is the value that, multiplied by itself, gives the original number, as in 7 times 7 equals 49."
    slides = [("Cover", ["x"], ""), ("Concept", [para, "Facilitator Notes: explain"], "")]
    r = fg2(tmp_path, slides, plan_paragraphs=[para])
    assert r.status == FAIL and "Lesson Plan paragraph(s) almost word for word" in r.message


# ------------------------------------------------------------- citations
def run(code, d, tmp_path):
    info = parse_docx(save(d, tmp_path))
    return F.DOCX_CHECKS[code](make_ctx(), ref(), info)


def test_we21_na_without_citations_and_quran_references_are_not_citations(tmp_path):
    d = new_doc()
    add(d, "Allah says: ولقد يسرنا (سورۃ القمر: 17)")
    assert run("WE21", d, tmp_path).status == NA


@pytest.mark.parametrize("text,status", [
    ("Learning improves with practice (Smith, 2019).", PASS),
    ("See (Smith & Jones, 2020) and (Lee et al., 2018).", PASS),
    ("Learning improves with practice [1].", FAIL),
    ("Learning improves with practice (Smith 2019).", FAIL),
])
def test_we21_apa_in_text(tmp_path, text, status):
    d = new_doc()
    add(d, text)
    assert run("WE21", d, tmp_path).status == status


def reference_doc(entries, hanging=True, heading="References", cite=True):
    d = new_doc()
    if cite:
        add(d, "Practice helps (Smith, 2019).")
    add(d, heading)
    for e in entries:
        p = add(d, e)
        if hanging:
            p.paragraph_format.left_indent = Twips(720)
            p.paragraph_format.first_line_indent = Twips(-720)
    return d


def test_we22_alphabetical_with_hanging_indent_passes(tmp_path):
    d = reference_doc(["Adams, B. (2018). A book.", "Smith, J. (2019). Another book."])
    assert run("WE22", d, tmp_path).status == PASS


def test_we22_fails_unsorted_or_no_hanging_indent_or_missing_list(tmp_path):
    r = run("WE22", reference_doc(["Smith, J. (2019). B.", "Adams, B. (2018). A."]), tmp_path)
    assert r.status == FAIL and "alphabetical" in r.message
    r = run("WE22", reference_doc(["Adams, B. (2018). A.", "Smith, J. (2019). B."], hanging=False), tmp_path)
    assert r.status == FAIL and "hanging indent" in r.message
    d = new_doc()
    add(d, "Practice helps (Smith, 2019).")
    assert run("WE22", d, tmp_path).status == FAIL
    d = new_doc()
    add(d, "No citations here.")
    assert run("WE22", d, tmp_path).status == NA
