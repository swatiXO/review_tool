"""Rules from the output guidelines and templates (LPG, FGG, ASG, SB, WEG) and the code-decided
parts of FG1, FG4, CE1, CE5 and slide footers (WE25)."""
import docx
from docx.shared import Pt

from course_review.checks import guidelines as G
from course_review.checks import formatting as F
from course_review.checks.assessments import _answer_key, _split_check, question_level
from course_review.checks.common import result
from course_review.docx_model import parse_docx
from course_review.models import DocRef, FAIL, NA, PASS, REVIEW
from course_review.questions import Question

from helpers import PROFILE, add, make_ctx, new_doc, ref, save
from test_slides_citations import deck

SECTIONS = ["Introduction", "SLOs", "Warm-up", "Concept Building", "Key Takeaways"]


def plan(tmp_path, body, name="plan.docx"):
    """body: {section: [lines]}; a line starting with '#' is a bold sub-heading."""
    d = new_doc()
    add(d, "Lesson Plan: Roots", bold=True)
    for s in SECTIONS:
        add(d, s, bold=True)
        for line in body.get(s, ["Some text here that is a sentence."]):
            if line.startswith("#"):
                add(d, line[1:], bold=True)
            else:
                add(d, line)
    return d, save(d, tmp_path, name)


def run(fn, tmp_path, body, doc=None):
    d, path = plan(tmp_path, body)
    return fn(make_ctx(), doc or ref(), parse_docx(path))


# --------------------------------------------------------------- Lesson Plan
def test_lpg1_warm_up_has_one_to_three_questions(tmp_path):
    assert run(G.lpg1, tmp_path, {"Warm-up": ["What is a root?", "Where do roots grow?"]}).status == PASS
    r = run(G.lpg1, tmp_path, {"Warm-up": ["Q one?", "Q two?", "Q three?", "Q four?"]})
    assert r.status == FAIL and "4 questions" in r.message and len(r.marks) == 1
    assert run(G.lpg1, tmp_path, {"Warm-up": ["سوال ١: جڑ کیا ہے؟"]}).status == PASS


def test_lpg2_key_takeaways_in_bullets(tmp_path):
    assert run(G.lpg2, tmp_path, {"Key Takeaways": ["• Roots hold the plant.", "• Roots take in water."]}).status == PASS
    assert run(G.lpg2, tmp_path, {"Key Takeaways": ["Roots hold the plant. Roots take in water."]}).status == FAIL
    assert run(G.lpg2, tmp_path, {"Key Takeaways": ["1. Roots hold the plant."]}).status == FAIL


def test_lpg3_coverage_table(tmp_path):
    d, path = plan(tmp_path, {})
    assert G.lpg3(make_ctx(), ref(), parse_docx(path)).status == FAIL
    t = d.add_table(rows=2, cols=3)
    for c, h in zip(t.rows[0].cells, ["Book Heading", "Lesson Plan Location", "SLO(s)"]):
        c.text = h
    path = save(d, tmp_path, "plan2.docx")
    assert G.lpg3(make_ctx(), ref(), parse_docx(path)).status == PASS


def test_weg2_numbered_list_inside_a_section_fails(tmp_path):
    d, path = plan(tmp_path, {})
    d.add_paragraph("first item", style="List Number")
    d.add_paragraph("second item", style="List Number")
    path = save(d, tmp_path, "plan3.docx")
    r = G.weg2(make_ctx(), ref(), parse_docx(path))
    assert r.status == FAIL and len(r.marks) == 2


def test_weg1_space_after_8pt(tmp_path):
    d = new_doc()
    d.styles["Normal"].paragraph_format.space_after = Pt(8)
    add(d, "One paragraph.")
    add(d, "Another paragraph.")
    assert G.weg1(make_ctx(), ref(), parse_docx(save(d, tmp_path))).status == PASS
    d.styles["Normal"].paragraph_format.space_after = Pt(0)
    assert G.weg1(make_ctx(), ref(), parse_docx(save(d, tmp_path, "y.docx"))).status == FAIL


# ------------------------------------------------------- Facilitator's Guide
GUIDE = ref(rel="p/Lesson-1-Guide.pptx", doc_type="facilitator_guide")
FIVE = ("Explain with an analogy. Check for understanding: show of hands. Misconception: roots are not stems. "
        "Video: play 0:15-0:45. Transition to the next concept.")


def test_fgg1_notes_in_the_speaker_pane(tmp_path):
    info = deck(tmp_path, [("Cover", ["Roots, 40 minutes"], "Time: 2 min"), ("SLOs", ["a"], "")])
    r = G.fgg1(make_ctx(), GUIDE, info)
    assert r.status == FAIL and "[2]" in r.message


def test_fgg2_times_add_up_to_the_planned_duration(tmp_path):
    ok = [("Cover", ["Duration: 10 minutes"], "Time: 2 min"), ("Intro", ["a"], "Time: 3 min"), ("CB", ["b"], "وقت: 5 منٹ")]
    assert G.fgg2(make_ctx(), GUIDE, deck(tmp_path, ok)).status == PASS
    off = [("Cover", ["Duration: 40 minutes"], "Time: 2 min"), ("Intro", ["a"], "Time: 3 min")]
    r = G.fgg2(make_ctx(), GUIDE, deck(tmp_path, off, "b.pptx"))
    assert r.status == FAIL and "planned for 40" in r.message
    assert G.fgg2(make_ctx(), GUIDE, deck(tmp_path, [("Cover", ["x"], "")], "c.pptx")).status == FAIL


def test_fgg3_cover_states_duration(tmp_path):
    assert G.fgg3(make_ctx(), GUIDE, deck(tmp_path, [("Roots", ["Duration: 40 minutes"], "")])).status == PASS
    assert G.fgg3(make_ctx(), GUIDE, deck(tmp_path, [("Roots", ["Grade 6"], "")], "b.pptx")).status == FAIL


def test_fg4_five_note_items_on_concept_building_slides(tmp_path):
    good = [("Cover", ["x"], ""), ("Concept Building - Part 1", ["Roots"], FIVE)]
    assert G.fg4(make_ctx(), GUIDE, deck(tmp_path, good)).status == PASS
    bare = [("Cover", ["x"], ""), ("Concept Building - Part 1", ["Roots"], "")]
    r = G.fg4(make_ctx(), GUIDE, deck(tmp_path, bare, "b.pptx"))
    assert r.status == FAIL and r.marks
    some = [("Cover", ["x"], ""), ("تصوراتی تعمیر - حصہ 1", ["Roots"], "Explain with an example.")]
    r = G.fg4(make_ctx(), GUIDE, deck(tmp_path, some, "c.pptx"))
    assert r.status == REVIEW and "misconceptions" in r.message


def test_fg1_fails_when_slos_are_pasted_from_the_lesson_plan(tmp_path):
    slo = "Identify the three main parts of a plant root system"
    _, path = plan(tmp_path, {"SLOs": ["• " + slo, "• Explain how water moves up from the roots"]})
    lp = DocRef(rel="p/Lesson-1-Lesson-Plan.docx", abs=path, ext="docx", doc_type="lesson_plan", chapter=1, lesson=1)
    ctx = make_ctx(docs=[lp])
    pasted = deck(tmp_path, [("Cover", ["x"], ""), ("SLOs", [slo + "\nExplain how water moves up from the roots"], "")])
    r = G.fg1_verbatim(ctx, GUIDE, pasted)
    assert r is not None and r.status == FAIL
    plain = deck(tmp_path, [("Cover", ["x"], ""), ("SLOs", ["Today you will find the parts of a root"], "")], "b.pptx")
    assert G.fg1_verbatim(ctx, GUIDE, plain) is None


def test_we25_slide_footers(tmp_path):
    info = deck(tmp_path, [("Cover", ["x"], "")])
    assert F.we25(make_ctx(), GUIDE, info).status == FAIL
    info.footers = [(1, "sldNum", "1"), (1, "dt", "2026-10-06"), (1, "ftr", "Facilitator-Guide-Lesson-1-v0.1")]
    assert F.we25(make_ctx(), GUIDE, info).status == PASS


# ---------------------------------------------------------------- Assessments
def q(n, head, block=""):
    return Question(num=n, text=head, start=n, end=n + 1, block=head + "\n" + block)


def test_question_level_from_labels():
    v = PROFILE["vocab"]
    assert question_level(q(1, "سوال ١ (کثیر الانتخابی — تنقیدی سوچ)"), v) == "higher"
    assert question_level(q(2, "سوال 1  قسم: استدلالی سوال"), v) == "higher"
    assert question_level(q(3, "سوال 2   قسم: تصوراتی وضاحت"), v) == "lower"
    assert question_level(q(4, "Question 2 (Remember)"), v) == "lower"
    assert question_level(q(5, "Question 3"), v) is None


def test_ce1_split_fails_when_labels_show_mostly_higher_order():
    qs = [q(i, f"Question {i} (Analyze)") for i in range(1, 7)]
    f = _split_check(result("CE1", PASS, "6 questions", partial=True), qs, 0.3, "70/30")
    assert f.status == FAIL and "higher-order" in f.message
    qs = [q(i, f"Question {i} (Remember)") for i in range(1, 5)] + [q(i, f"Question {i} (Evaluate)") for i in range(5, 7)]
    assert _split_check(result("CE1", PASS, "6 questions", partial=True), qs, 0.3, "70/30").status == PASS
    unlabelled = [q(i, f"Question {i}") for i in range(1, 7)]
    assert _split_check(result("CE1", PASS, "6 questions", partial=True), unlabelled, 0.3, "70/30").status == PASS


def test_ce5_answer_key(tmp_path):
    d = new_doc()
    add(d, "Question 1: what?")
    info = parse_docx(save(d, tmp_path))
    ctx = make_ctx()
    assert _answer_key(ctx, info, [q(1, "Question 1: what?", "Model answer: this")]).status == PASS
    assert _answer_key(ctx, info, [q(1, "Question 1: what?", "")]).status == NA


def test_asg1_all_or_none_of_the_above(tmp_path):
    d = new_doc()
    add(d, "D) All of the above")
    r = G.asg1(make_ctx(), ref(doc_type="chapter_exam"), parse_docx(save(d, tmp_path)))
    assert r.status == FAIL and r.marks


# ------------------------------------------------------------------ Storyboard
def storyboard(tmp_path, rows, top="Topic: pivotal concept. New video.", headers=None):
    d = new_doc()
    add(d, "Storyboard")
    add(d, top)
    headers = headers or ["Scene #", "Duration", "On-Screen Text", "Narration", "Animation / Visual Notes"]
    t = d.add_table(rows=1 + len(rows), cols=len(headers))
    for c, h in zip(t.rows[0].cells, headers):
        c.text = h
    for r, vals in zip(t.rows[1:], rows):
        for c, v in zip(r.cells, vals):
            c.text = v
    path = save(d, tmp_path, "Storyboard-Lesson-1.docx")
    sb = DocRef(rel="p/Storyboard-Lesson-1.docx", abs=path, ext="docx", doc_type="video_storyboard", chapter=1, lesson=1)
    return {f.code: f for f in G.storyboard_checks(make_ctx(), sb, parse_docx(path))}


GOOD = [["1", "0:00–0:15", "States of Matter", "Everything around us is matter.", "Title card"],
        ["2", "0:15–1:45", "Solid · Liquid · Gas", "In a solid, particles are packed tightly.", "Zoom into an ice cube"],
        ["3", "1:45–2:00", "Solid → Liquid → Gas", "Same particles, different energy.", "All three side by side"]]


def test_storyboard_from_the_template_passes(tmp_path):
    r = storyboard(tmp_path, GOOD)
    assert all(r[c].status == PASS for c in ("SB1", "SB2", "SB3", "SB4", "SB5")), {c: f.message for c, f in r.items()}


def test_storyboard_over_two_minutes_fails(tmp_path):
    rows = GOOD + [["4", "2:00–2:30", "Extra", "One more thing.", "Extra scene"]]
    assert storyboard(tmp_path, rows)["SB2"].status == FAIL


def test_storyboard_long_narration_and_sentences_on_screen_fail(tmp_path):
    long = " ".join(["word"] * 320)
    rows = [["1", "0:00–1:00", "This is a whole sentence that the narrator is already saying out loud to everyone", long, "x"]]
    r = storyboard(tmp_path, rows)
    assert r["SB3"].status == FAIL and r["SB5"].status == FAIL


def test_storyboard_without_new_or_edited_line_and_missing_columns(tmp_path):
    r = storyboard(tmp_path, GOOD, top="Topic: states of matter")
    assert r["SB4"].status == FAIL
    r = storyboard(tmp_path, [["1", "0:00-0:10", "x"]], headers=["Scene #", "Duration", "On-Screen Text"])
    assert r["SB1"].status == FAIL and "narration" in r["SB1"].message


def test_storyboard_urdu_template_columns(tmp_path):
    rows = [["1", "0:00–0:30", "صبر", "صبر کا مطلب ہے رکنا۔", "عنوان"]]
    r = storyboard(tmp_path, rows, top="نیا ویڈیو", headers=["منظر نمبر", "دورانیہ", "اسکرین پر متن", "بیانیہ", "متحرک تصویر / بصری نوٹس"])
    assert r["SB1"].status == PASS and r["SB2"].status == PASS and r["SB4"].status == PASS
