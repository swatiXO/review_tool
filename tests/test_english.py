"""English lessons built from the team's English templates: the section names sit in a small box with the
instructions in the title placeholder, the coverage table follows Key Takeaways, and the worked example sets
story sentences in a Heading style."""
import os

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from course_review import autofix
from course_review.checks import formatting as F
from course_review.checks import guidelines as G
from course_review.checks import slides as S
from course_review.docx_model import parse_docx
from course_review.models import FAIL, PASS
from course_review.pptx_model import parse_pptx

from helpers import add, make_ctx, new_doc, ref, save

GUIDE = ref(rel="p/Lesson-1-Guide.pptx", doc_type="facilitator_guide")


def template_deck(tmp_path, slides):
    """slides: [(instructions in the title placeholder, section name in a small box, notes)]"""
    prs = Presentation()
    for instructions, name, notes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = instructions
        s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(4), Inches(0.5)).text_frame.text = name
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    path = os.path.join(str(tmp_path), "guide.pptx")
    prs.save(path)
    return parse_pptx(path)


ENGLISH_GUIDE = [
    ("Grade:\n\nSubject:\n\nDuration:", "Facilitator's Guide", ""),
    ("Lesson SLOs (restated in simplified, learner-facing language):\n\nSLO-1\nSLO-2", "Session Overview", ""),
    ("Restate the introduction of the lesson with the scenario or story presented in the lesson.", "Introduction", ""),
    ("Start with the debrief, i.e. explain what learners need to do in the activity.\nDo not rewrite it.", "Warm-up Activity", ""),
    ("Write the main content in bullets here.\nBe careful: Only one idea per slide.", "Concept Building - Part 1", ""),
    ("Synthesise the Concept Building slides rather than re-listing them\nTie the recap to the SLOs", "Key Takeaways", ""),
    ("Close/Conclusion", "Close/Conclusion", ""),
]


def test_slide_name_comes_from_the_short_box_not_the_instructions(tmp_path):
    info = template_deck(tmp_path, ENGLISH_GUIDE)
    assert [s.title for s in info.slides_text][1:6] == ["Session Overview", "Introduction", "Warm-up Activity",
                                                         "Concept Building - Part 1", "Key Takeaways"]


def test_fg7_reads_the_english_guide_structure(tmp_path):
    assert S.fg7(make_ctx(), GUIDE, template_deck(tmp_path, ENGLISH_GUIDE)).status == PASS


def test_fg7_accepts_recap_for_key_takeaways(tmp_path):
    slides = [x if x[1] != "Key Takeaways" else (x[0], "Recap", "") for x in ENGLISH_GUIDE]
    assert S.fg7(make_ctx(), GUIDE, template_deck(tmp_path, slides)).status == PASS


def test_fg4_does_not_take_the_recap_for_a_concept_building_slide(tmp_path):
    r = G.fg4(make_ctx(), GUIDE, template_deck(tmp_path, ENGLISH_GUIDE))
    assert "slide 6" not in r.message


def english_plan(tmp_path):
    d = new_doc()
    add(d, "Lesson Plan")
    for name, body in [("Introduction", ["Today's lesson is about states of matter."]),
                       ("SLOs", ["• Identify solids, liquids, and gases"]),
                       ("Warm-up Activity", ["Can you name one solid?"]),
                       ("Concept Building", []),
                       ("Key Takeaways", ["• Matter exists in three states."])]:
        add(d, name, style="Heading 1")
        for line in body:
            add(d, line)
        if name == "Concept Building":
            add(d, "4.1. Solid", style="Heading 2")
            add(d, "Aliza wakes up. She pours water into a glass. She drops in an ice cube.", style="Heading 3")
    add(d, "Book and SLO Coverage Map", style="Heading 1")
    t = d.add_table(rows=2, cols=2)
    for c, v in zip(t.rows[0].cells + t.rows[1].cells, ["Book Heading", "Lesson Plan Location", "What Is Matter?", "1. Introduction"]):
        c.text = v
    return save(d, tmp_path, "Lesson-Plan-Lesson-2-Chapter-3.docx")


def test_coverage_table_cells_are_not_numbered_takeaways(tmp_path):
    assert G.lpg2(make_ctx(), ref(), parse_docx(english_plan(tmp_path))).status == PASS


def test_story_sentences_in_a_heading_style_are_body_text_not_headings(tmp_path):
    info = parse_docx(english_plan(tmp_path))
    heads = [p.text for p in F.doc_headings(make_ctx(), info)]
    assert not any(h.startswith("Aliza") for h in heads)
    r = F.we13(make_ctx(), ref(), info)
    assert r.status == FAIL and "running text" in r.message and r.marks[0]["text"].startswith("Aliza")


def test_coverage_map_heading_needs_no_number(tmp_path):
    info = parse_docx(english_plan(tmp_path))
    r = F.we12(make_ctx(), ref(), info)
    assert "Book and SLO" not in (r.message + " ".join(r.evidence or []))


def test_autofix_restyles_the_story_and_captions_the_table_from_its_heading(tmp_path):
    src = english_plan(tmp_path)
    dst = os.path.join(str(tmp_path), "fixed.docx")
    ctx = make_ctx()
    autofix.fix_docx(src, dst, ctx, ref(), parse_docx(src))
    d = Document(dst)
    by_text = {p.text: p for p in d.paragraphs}
    story = next(p for t, p in by_text.items() if t.startswith("Aliza"))
    assert not story.style.name.startswith("Heading")
    assert "Book and SLO Coverage Map" in by_text                      # left unnumbered
    assert "Table 1: Book and SLO Coverage Map" in by_text
    assert "5 Key Takeaways" in by_text and "4.1 Solid" in by_text
