"""Ingest safety and classification, workbook fill, and one small end-to-end review."""
import os
import zipfile

import openpyxl
import pytest
from docx.enum.text import WD_COLOR_INDEX
from pptx import Presentation
from pptx.util import Inches

from course_review import cli, ingest, workbook
from course_review.models import LessonKey

from helpers import PROFILE, add, checklist, new_doc, save, zip_dir


def touch(root, rel):
    path = os.path.join(root, *rel.split("/"))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    open(path, "wb").write(b"x")


# --------------------------------------------------------------- safety
def test_zip_slip_is_rejected(tmp_path):
    z = tmp_path / "evil.zip"
    with zipfile.ZipFile(z, "w") as f:
        f.writestr("../escape.txt", "x")
    with pytest.raises(ingest.UnsafeArchive):
        ingest.safe_extract(z, tmp_path / "out")


def test_entry_count_limit(tmp_path):
    z = tmp_path / "many.zip"
    with zipfile.ZipFile(z, "w") as f:
        for i in range(5):
            f.writestr(f"a/{i}.txt", "x")
    with pytest.raises(ingest.UnsafeArchive):
        ingest.safe_extract(z, tmp_path / "out", max_files=3)


# ------------------------------------------------------- classification
def test_classify_paths_variants_and_types(tmp_path):
    root = str(tmp_path / "Grade-6-Islamiat")
    for rel in [
        "Chapter-2-Faith/Lesson-1-Tauheed/Lesson-1-Chapter-2-Lesson-Plan.docx",
        "Chapter-2-Faith/Lesson-1-Tauheed/Lesson-1-Chapter-2-Facilitator-Guide.pptx",
        "Chapter-2-Faith/Lesson-1-Tauheed/Lesson-1-Chapter-2-Assessment.docx",
        "Chapter-2-Faith/Lesson-1-Tauheed/Lesson-1-Chapter-2-Video-Storyboard.docx",
        "Chapter-2-Faith/Chapter-Exam-Chapter-2/Chapter-Exam-Chapter-2.docx",
        "Chpater-7-Today/Lesson-2-Animals/Lesson-2-Chapter-7-Lesson Plan.docx",
        "Chapter-6-Heroes/Lesson-2-Chapter-6-Part-2-Fatima/Lesson-2-Chapter-6-Part-2-B-Assessment.docx",
        "Chapter-6-Heroes/Lesson-2-Chapter-6-Part-1-Family/Lesson-2-A-Chapter-6-Lesson-Plan.docx",
        "Pop Quiz/Pop Quiz.docx",
        "Data Bank/Grade-6-Data-Bank-v0.1.docx",
        "Data Bank/Grade-6-Data-Bank-v0.2.docx",
        "book.pdf",
        "notes.txt",
    ]:
        touch(root, rel)
    pkg = ingest.classify(root, PROFILE)
    by = {d.rel.split("/")[-1]: d for d in pkg.docs}
    assert by["Lesson-1-Chapter-2-Lesson-Plan.docx"].doc_type == "lesson_plan"
    assert by["Lesson-1-Chapter-2-Lesson-Plan.docx"].key == LessonKey(2, 1, "")
    assert by["Lesson-1-Chapter-2-Facilitator-Guide.pptx"].doc_type == "facilitator_guide"
    assert by["Lesson-1-Chapter-2-Assessment.docx"].doc_type == "chapter_exam"
    assert by["Lesson-1-Chapter-2-Video-Storyboard.docx"].doc_type == "video_storyboard"
    ws = by["Chapter-Exam-Chapter-2.docx"]
    assert ws.doc_type == "worksheet" and ws.scope == "chapter" and ws.chapter == 2 and ws.notes
    assert by["Lesson-2-Chapter-7-Lesson Plan.docx"].key == LessonKey(7, 2, "")   # typo'd folder, space in name
    assert by["Lesson-2-Chapter-6-Part-2-B-Assessment.docx"].key == LessonKey(6, 2, "B")
    assert by["Lesson-2-A-Chapter-6-Lesson-Plan.docx"].key == LessonKey(6, 2, "A")
    assert by["Grade-6-Data-Bank-v0.1.docx"].superseded and not by["Grade-6-Data-Bank-v0.2.docx"].superseded
    assert pkg.unclassified == ["notes.txt"]


# --------------------------------------------------------------- workbook
def test_workbook_rules_and_fill_extends_rows_and_validation(tmp_path):
    path = checklist(tmp_path)
    rules, layout = workbook.load_rules(path)
    assert rules["LP3"].scope == "Per lesson" and layout.matrix["Worksheet"]["scope"] == "chapter"
    rows = [{"label": f"Lesson {i}", "cells": {"LP3": "pass", "LP4": "fail", "LP5": None, "LP10": "na"}, "note": f"n{i}"} for i in range(1, 8)]
    out = tmp_path / "filled.xlsx"
    workbook.fill_workbook(path, out, {"Lesson Plan": rows}, {"WE1": ("fail", "bad name")}, [("LP3", "Yes", "how", "text")], {"Package": "t"})
    wb = openpyxl.load_workbook(out)
    ws = wb["Lesson Plan"]
    assert ws["B6"].value == "Lesson 1" and ws["C6"].value == "Pass" and ws["D6"].value == "Fail" and ws["E6"].value is None
    assert ws["B12"].value == "Lesson 7"                     # grew past the template's 3 rows
    assert ws["G12"].value.startswith("=COUNTIF(C12:F12")    # formula for the added row
    assert ws["C5"].value is None and ws["A5"].value == "Auto"  # example row replaced by the legend
    assert any("C5:F12" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert wb["Subject-Level"]["D5"].value == "Fail" and wb["Subject-Level"]["E5"].value == "bad name"
    assert "Review Summary" in wb.sheetnames


# ----------------------------------------------------------------- end to end
def build_package(tmp_path):
    root = tmp_path / "Grade-6-Test"
    lesson = root / "Chapter-1-Intro" / "Lesson-1-Roots"
    lesson.mkdir(parents=True)
    d = new_doc()
    for s in ("Introduction", "SLOs", "Warm-up", "Concept Building", "Key Takeaways"):
        add(d, s + ":", bold=True)
        if s == "SLOs":
            add(d, "Find square roots")      # plain paragraphs, not bullets, so LP5 fails
            add(d, "Explain perfect squares")
    d.save(str(lesson / "Lesson-1-Chapter-1-Lesson-Plan.docx"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.save(str(lesson / "Lesson-1-Chapter-1-Facilitator-Guide.pptx"))
    q = new_doc()
    add(q, "Lesson 1: Roots")
    for i in (1, 2):
        add(q, f"Question {i}: pick")
        add(q, "A) a")
        add(q, "B) b", highlight=WD_COLOR_INDEX.YELLOW)
        add(q, "C) c")
    (root / "Pop Quiz").mkdir()
    q.save(str(root / "Pop Quiz" / "Pop Quiz.docx"))
    return root


def test_end_to_end_review(tmp_path):
    root = build_package(tmp_path)
    zip_path = zip_dir(str(root), str(tmp_path / "pkg.zip"))
    out = tmp_path / "out"
    pkg, res, xlsx, _ = cli.review(zip_path, checklist(tmp_path), str(out))
    assert (out / "report.html").exists() and (out / "review.json").exists() and os.path.exists(xlsx)
    assert not res.errors
    wb = openpyxl.load_workbook(xlsx)
    lp = wb["Lesson Plan"]
    assert lp["B6"].value.startswith("Chapter 1 - Lesson 1")
    assert lp["C6"].value == "Pass"          # LP3: five sections present
    assert lp["E6"].value == "Fail"          # LP5: no SLO bullets
    pq = wb["Pop Quiz"]
    assert pq["C6"].value == "Pass" and pq["D6"].value == "Pass"   # PQ1 count, PQ4 yellow highlight
    assert wb["Subject-Level"]["D5"].value == "Fail"                # WE1: file names do not follow the pattern


def test_corrupt_file_is_reported_not_fatal(tmp_path):
    root = build_package(tmp_path)
    lesson = root / "Chapter-1-Intro" / "Lesson-1-Roots"
    (lesson / "Lesson-1-Chapter-1-Assessment.docx").write_bytes(b"this is not a docx")
    zip_path = zip_dir(str(root), str(tmp_path / "pkg.zip"))
    pkg, res, xlsx, _ = cli.review(zip_path, checklist(tmp_path), str(tmp_path / "out"))
    assert any(p.endswith("Assessment.docx") for p in res.errors)
    assert os.path.exists(xlsx)


def test_package_with_files_at_the_root_and_no_lesson_folders(tmp_path):
    root = tmp_path / "pkg"
    root.mkdir()
    new_doc().save(str(root / "Pop Quiz.docx"))
    zip_path = zip_dir(str(root), str(tmp_path / "pkg.zip"))
    pkg, res, xlsx, _ = cli.review(zip_path, checklist(tmp_path), str(tmp_path / "out"))
    assert os.path.exists(xlsx) and len(pkg.docs) == 1
