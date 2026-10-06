"""Each formatting rule in its pass and fail form, on tiny generated documents."""
import pytest
from docx.enum.text import WD_COLOR_INDEX
from PIL import Image

from course_review.checks import formatting as F
from course_review.docx_model import parse_docx
from course_review.models import FAIL, NA, PASS, REVIEW

from helpers import A4, add, make_ctx, new_doc, ref, save

JAMIL = "Jamil Noori Nastaliq"
URDU = "سبق کا منصوبہ"


def run(code, d, tmp_path, doc=None):
    info = parse_docx(save(d, tmp_path))
    return F.DOCX_CHECKS[code](make_ctx(), doc or ref(), info)


# WE4 ----------------------------------------------------------------------
def test_we4_passes_on_a4_one_inch_115(tmp_path):
    d = new_doc()
    add(d, "Hello world")
    assert run("WE4", d, tmp_path).status == PASS


def test_we4_fails_on_us_letter(tmp_path):
    d = new_doc(page=(8.5, 11))
    add(d, "Hello")
    r = run("WE4", d, tmp_path)
    assert r.status == FAIL and "A4" in r.message


def test_we4_fails_on_wrong_spacing_and_margins(tmp_path):
    d = new_doc(margin=0.5, spacing=1.5)
    add(d, "Hello")
    r = run("WE4", d, tmp_path)
    assert r.status == FAIL and "margins" in r.message


# WE6 / WE7 ----------------------------------------------------------------
def test_we6_poppins_passes_arial_fails_mixed_fails(tmp_path):
    d = new_doc()
    add(d, "Hello", font="Poppins")
    assert run("WE6", d, tmp_path).status == PASS
    d = new_doc()
    add(d, "Hello", font="Arial")
    assert run("WE6", d, tmp_path).status == FAIL
    d = new_doc()
    add(d, "Hello", font="Poppins")
    add(d, "World", font="Montserrat")
    r = run("WE6", d, tmp_path)
    assert r.status == FAIL and "mixed" in r.message


def test_we6_is_na_without_english(tmp_path):
    d = new_doc()
    add(d, URDU, urdu_font=JAMIL, rtl=True)
    assert run("WE6", d, tmp_path).status == NA


def test_we7_reads_complex_script_font(tmp_path):
    # team decision 2026-10-06: Urdu is set in Noto Nastaliq (the guideline's size table)
    d = new_doc()
    add(d, URDU, urdu_font="Noto Nastaliq Urdu", rtl=True)
    assert run("WE7", d, tmp_path).status == PASS
    d = new_doc()
    add(d, URDU, urdu_font=JAMIL, rtl=True)
    r = run("WE7", d, tmp_path)
    assert r.status == FAIL and "Jamil" in r.message


# WE8: sizes by grade from the Writing & Editing Guidelines (Grade 6+: title 20, heading 14, body 12)
def titled(size=12, **kw):
    d = new_doc(size=size)
    add(d, "Document title", size=20, bold=True)
    return d


def test_we8_body_size_follows_the_grade(tmp_path):
    d = titled()
    add(d, "Body text here, a normal sentence.", size=12)
    r = run("WE8", d, tmp_path)
    assert r.status == PASS and "Grade 6" in " ".join(r.evidence)
    d = titled()
    add(d, "Body text here, a normal sentence.", size=11)
    r = run("WE8", d, tmp_path)
    assert r.status == FAIL and "Body (12)" in r.message


def test_we8_grade_1_to_5_uses_the_larger_sizes(tmp_path):
    from course_review.checks.formatting import grade_sizes
    ctx = make_ctx()
    ctx.pkg.subject = "Grade-4-Science"
    assert grade_sizes(ctx)[1]["body"] == 14 and grade_sizes(ctx)[1]["section_heading"] == 16


def test_we8_uses_complex_script_size_for_urdu(tmp_path):
    d = titled()
    add(d, URDU, size=12, urdu_font=JAMIL, rtl=True)  # sz and szCs both 12
    assert run("WE8", d, tmp_path).status == PASS
    d = titled()
    p = add(d, URDU, size=12, urdu_font=JAMIL, rtl=True)
    for el in p._p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}szCs"):
        el.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "22")  # Word shows 11pt
    assert run("WE8", d, tmp_path).status == FAIL


def test_we8_section_heading_14_bold_including_bold_lines_used_as_headings(tmp_path):
    d = titled()
    h = d.add_heading("Title of section", level=1)
    for r in h.runs:
        r.font.size = __import__("docx").shared.Pt(14)
        r.font.bold = True
        r.font.color.rgb = None
    add(d, "Body text here, a normal sentence.", size=12)
    assert run("WE8", d, tmp_path).status == PASS
    d = titled()
    add(d, "A bold line used as a heading", size=11, bold=True)
    add(d, "Body text here, a normal sentence.", size=12)
    r = run("WE8", d, tmp_path)
    assert r.status == FAIL and "Section heading (14 bold)" in r.message


# WE10 / WE11 --------------------------------------------------------------
def test_we10_flags_eastern_digits_and_notes_verse_markers(tmp_path):
    d = new_doc()
    add(d, "Lesson 12 is fine")
    assert run("WE10", d, tmp_path).status == PASS
    d = new_doc()
    add(d, "﴿١﴾ ایک ۲۳")
    r = run("WE10", d, tmp_path)
    assert r.status == FAIL and "verse-number" in r.message


def test_we11_direction_matches_script(tmp_path):
    d = new_doc()
    add(d, URDU, rtl=True)
    add(d, "English paragraph")
    r = run("WE11", d, tmp_path)
    assert r.status == PASS and r.partial
    d = new_doc()
    add(d, URDU, rtl=False)
    assert run("WE11", d, tmp_path).status == FAIL


# WE12 / WE13 --------------------------------------------------------------
def test_we12_numbering(tmp_path):
    d = new_doc()
    d.add_heading("1 Intro", 1)
    d.add_heading("1.1 Sub", 2)
    d.add_heading("1.2 Sub two", 2)
    d.add_heading("2 Next", 1)
    assert run("WE12", d, tmp_path).status == PASS


@pytest.mark.parametrize("headings", [
    ["1 Intro", "1.1.1 Deep"],                 # skipped level
    ["1 Intro", "3 Jumped"],                   # out of order
    ["Intro", "Sub"],                          # not numbered
    ["1 A", "1.1 B", "1.1.1 C", "1.1.1.1 D"],  # more than three levels
])
def test_we12_failures(tmp_path, headings):
    d = new_doc()
    for h in headings:
        d.add_heading(h, 1)
    assert run("WE12", d, tmp_path).status == FAIL


def test_we13_requires_builtin_heading_styles(tmp_path):
    d = new_doc()
    d.add_heading("1 Real heading", 1)
    assert run("WE13", d, tmp_path).status == PASS
    d = new_doc()
    add(d, "Document title", bold=True)
    add(d, "Looks like a heading", bold=True)
    add(d, "Body text that is a normal sentence and goes on for a while. It ends with a full stop.")
    r = run("WE13", d, tmp_path)
    assert r.status == FAIL and "Heading style" in r.message


def test_title_question_labels_and_sublabels_are_not_headings(tmp_path):
    d = new_doc()
    add(d, "Assessment: Patience", bold=True)
    add(d, "سوال ١ (کثیر الانتخابی)", bold=True)
    add(d, "Some question text that is a sentence.")
    add(d, "Knowledge:", bold=True)
    add(d, "Question 2: pick one", bold=True)
    assert run("WE12", d, tmp_path).status == NA
    assert run("WE13", d, tmp_path).status == NA


# captions, figures, tables --------------------------------------------------
def test_we16_table_caption_above_and_numbered(tmp_path):
    d = new_doc()
    add(d, "Table 1: Names")
    d.add_table(rows=1, cols=2)
    assert run("WE16", d, tmp_path).status == PASS
    d = new_doc()
    d.add_table(rows=1, cols=2)
    assert run("WE16", d, tmp_path).status == FAIL


def _png(path, colour):
    Image.new("RGB", (400, 300), colour).save(path)


def test_figures_alt_text_colour_and_caption(tmp_path):
    grey = str(tmp_path / "g.png")
    red = str(tmp_path / "r.png")
    _png(grey, (128, 128, 128))
    _png(red, (220, 30, 30))
    d = new_doc()
    d.add_picture(grey, width=__import__("docx").shared.Inches(2))
    cap = add(d, "Figure 1: A grey square", size=11, italic=True)
    assert run("WE17", d, tmp_path).status == PASS
    assert run("WE18", d, tmp_path).status == PASS
    assert run("WE20", d, tmp_path).status == FAIL  # no alt text set
    d = new_doc()
    d.add_picture(red, width=__import__("docx").shared.Inches(2))
    assert run("WE18", d, tmp_path).status == FAIL
    assert run("WE17", d, tmp_path).status == FAIL  # no caption


# WE23 colour ---------------------------------------------------------------
def test_we23_colour_text_fails_grey_passes(tmp_path):
    d = new_doc()
    add(d, "Grey text", color="555555")
    assert run("WE23", d, tmp_path).status == PASS
    d = new_doc()
    add(d, "Red text", color="FF0000")
    assert run("WE23", d, tmp_path).status == FAIL


def test_we23_judges_visible_colour_not_unused_style_definitions(tmp_path):
    d = new_doc()                                   # the default template defines blue Heading styles
    add(d, "Plain black text")
    assert run("WE23", d, tmp_path).status == PASS
    d = new_doc()
    d.add_heading("1 A heading in the template's blue Heading 1", 1)
    assert run("WE23", d, tmp_path).status == FAIL
    d = new_doc()
    h = d.add_heading("1 Heading 1 overridden to black", 1)
    for r in h.runs:
        r.font.color.rgb = __import__("docx").shared.RGBColor(0, 0, 0)
    assert run("WE23", d, tmp_path).status == PASS


def test_we23_yellow_highlight_allowed_only_in_pop_quiz(tmp_path):
    d = new_doc()
    add(d, "answer", highlight=WD_COLOR_INDEX.YELLOW)
    assert run("WE23", d, tmp_path, doc=ref(doc_type="pop_quiz")).status == PASS
    assert run("WE23", d, tmp_path, doc=ref(doc_type="lesson_plan")).status == FAIL


# WE25 footer ---------------------------------------------------------------
def test_we25_footer(tmp_path):
    d = new_doc()
    add(d, "x")
    assert run("WE25", d, tmp_path).status == FAIL
    d = new_doc()
    add(d, "x")
    d.sections[0].footer.paragraphs[0].text = "Grade 6 Islamiat Lesson Plan v1.0 05/10/2026 Page 1"
    r = run("WE25", d, tmp_path)
    assert r.status == PASS and r.partial


# naming ---------------------------------------------------------------------
@pytest.mark.parametrize("name,w1,w2,w3", [
    ("LessonPlan-G6-Tauheed-v1.0.docx", PASS, PASS, REVIEW),
    ("Lesson-4-Chapter-3-Assessment.docx", FAIL, PASS, FAIL),
    ("LessonPlan-G6-Tauheed-Final-v1.0.docx", PASS, FAIL, REVIEW),
])
def test_naming_rules(name, w1, w2, w3):
    doc = ref(rel="Pkg/" + name)
    assert F.we1(make_ctx(), doc, None).status == w1
    assert F.we2(make_ctx(), doc, None).status == w2
    assert F.we3(make_ctx(), doc, None).status == w3
