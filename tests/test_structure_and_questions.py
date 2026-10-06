"""Lesson Plan structure, question parsing, Pop Quiz / exam / Data Bank rules."""
import docx
from docx.enum.text import WD_COLOR_INDEX

from course_review.checks import assessments as A
from course_review.checks import lessonplan as LP
from course_review.docx_model import parse_docx
from course_review.models import FAIL, NA, PASS, REVIEW
from course_review.questions import numbering_is_regular, parse_questions, segment_lessons
from course_review.textutil import normalize, starts_with_label, to_western_digits

from helpers import PROFILE, add, make_ctx, new_doc, ref, save

SECTIONS = ["Introduction", "SLOs", "Warm-up", "Concept Building", "Key Takeaways"]


def plan(sections, tmp_path, bullets=True):
    d = new_doc()
    for s in sections:
        add(d, s + ":", bold=True)
        if s == "SLOs":
            for item in ("Define a square", "Find a root"):
                d.add_paragraph(item, style="List Bullet" if bullets else None)
    return parse_docx(save(d, tmp_path))


def test_lp3_all_sections_in_order(tmp_path):
    r = LP.lp3(make_ctx(), ref(), plan(SECTIONS, tmp_path))
    assert r.status == PASS


def test_lp3_missing_and_out_of_order(tmp_path):
    r = LP.lp3(make_ctx(), ref(), plan([s for s in SECTIONS if s != "Warm-up"], tmp_path))
    assert r.status == FAIL and "warm_up" in r.message
    swapped = SECTIONS[:]
    swapped[2], swapped[3] = swapped[3], swapped[2]
    r = LP.lp3(make_ctx(), ref(), plan(swapped, tmp_path))
    assert r.status == FAIL and "out of order" in r.message


def test_lp3_matches_urdu_headings_with_arabic_letter_variants(tmp_path):
    d = new_doc()
    # Arabic yeh/kaf code points and diacritics, as found in real files
    for label in ("تعارف", "حاصلاتِ تعليم",
                  "تحريکی سرگرمی", "تصوراتی تعمیر",
                  "حاصل کلام"):
        add(d, label)
    r = LP.lp3(make_ctx(), ref(), parse_docx(save(d, tmp_path)))
    assert r.status == PASS, r.message


def test_lp5_bullets_pass_plain_paragraphs_fail(tmp_path):
    assert LP.lp5(make_ctx(), ref(), plan(SECTIONS, tmp_path, bullets=True)).status == PASS
    r = LP.lp5(make_ctx(), ref(), plan(SECTIONS, tmp_path, bullets=False))
    assert r.status == FAIL and "not a bullet" in r.message


def test_lp5_numbered_list_fails(tmp_path):
    d = new_doc()
    add(d, "SLOs:", bold=True)
    d.add_paragraph("Define a square", style="List Number")
    add(d, "Key Takeaways:", bold=True)
    r = LP.lp5(make_ctx(), ref(), parse_docx(save(d, tmp_path)))
    assert r.status == FAIL and "numbered" in r.message


# --------------------------------------------------------------- text utils
def test_normalize_unifies_variants_and_strips_marks():
    assert normalize("تعليم") == normalize("تعلیم")
    assert to_western_digits("١٢۳") == "123"
    assert starts_with_label("تعارف : abc", ["تعارف"])
    assert not starts_with_label("Introduction to a very long sentence " * 4, ["Introduction"])


# ------------------------------------------------------------ question parse
def questions_doc(tmp_path, lines, tables=()):
    d = new_doc()
    for l in lines:
        add(d, l)
    for cells in tables:
        t = d.add_table(rows=len(cells), cols=1)
        for r, text in zip(t.rows, cells):
            r.cells[0].text = text
    return parse_docx(save(d, tmp_path))


def test_questions_numbered_variants(tmp_path):
    plain = parse_questions(questions_doc(tmp_path, ["1. First", "2) Second", "3: Third"]), PROFILE)
    assert [q.num for q in plain] == [1, 2, 3] and numbering_is_regular(plain)
    labelled = parse_questions(questions_doc(tmp_path, ["Question 1: First", "سوال ٢ - second", "سوال ٣۔ third"]), PROFILE)
    assert [q.num for q in labelled] == [1, 2, 3] and numbering_is_regular(labelled)


def test_numbers_before_the_first_labelled_question_are_header_lines(tmp_path):
    info = questions_doc(tmp_path, ["31: lesson title", "Question 1: a", "Question 2: b"])
    assert [q.num for q in parse_questions(info, PROFILE)] == [1, 2]


def test_questions_with_eastern_digits_and_paren_tag(tmp_path):
    info = questions_doc(tmp_path, ["سوال ١ (تصوراتی)", "text", "سوال ٢ (تجزیہ)"])
    assert [q.num for q in parse_questions(info, PROFILE)] == [1, 2]


def test_questions_inside_tables(tmp_path):
    info = questions_doc(tmp_path, [], tables=[["سوال 1  قسم: open"], ["سوال 2  قسم: mcq"]])
    assert [q.num for q in parse_questions(info, PROFILE)] == [1, 2]


def test_questions_stop_at_teacher_guidance(tmp_path):
    info = questions_doc(tmp_path, ["سوال 1: a", "سوال 2: b", "اساتذہ کے لیے رہنمائی",
                                        "سوال 5(ب) کے جوابات"])
    assert [q.num for q in parse_questions(info, PROFILE)] == [1, 2]


def test_irregular_numbering_is_detected(tmp_path):
    info = questions_doc(tmp_path, ["سوال 1: a", "1. nested", "2. nested", "سوال 2: b"])
    assert not numbering_is_regular(parse_questions(info, PROFILE))


def test_mcq_detection_and_highlight(tmp_path):
    d = new_doc()
    add(d, "Question 1: pick one")
    add(d, "A) one")
    add(d, "B) two", highlight=WD_COLOR_INDEX.YELLOW)
    add(d, "C) three")
    add(d, "Question 2: pick again")
    add(d, "A) x")
    add(d, "✅ B) y")
    add(d, "C) z")
    qs = parse_questions(parse_docx(save(d, tmp_path)), PROFILE)
    assert [q.qtype for q in qs] == ["mcq", "mcq"]
    assert qs[0].highlighted and not qs[1].highlighted and qs[1].check_marked


# ------------------------------------------------------------------ segmenting
def test_lesson_segmentation_restarts_numbers_per_chapter_and_keeps_parts(tmp_path):
    d = new_doc()
    for h in ("سبق 1: a", "سبق 2: b", "سبق 1: c", "سبق 2(الف): d", "سبق 2(ب): e", "سبق 3: f"):
        add(d, h)
        add(d, "Question 1: x")
    groups = segment_lessons(parse_docx(save(d, tmp_path)), PROFILE)
    assert [len(g) for g in groups] == [2, 4]
    assert [h["variant"] for h in groups[1]] == ["", "A", "B", ""]


# ------------------------------------------------------------- worksheet WS6
def test_ws6_answer_key_must_be_a_section_at_the_end(tmp_path):
    ctx = make_ctx()
    doc = ref(doc_type="worksheet", scope="chapter")
    inline = questions_doc(tmp_path, ["1. Q one\nAnswer Key: a", "2. Q two"])
    qs = parse_questions(inline, PROFILE)
    assert A._ws6(ctx, doc, inline, qs).status == FAIL
    ok = questions_doc(tmp_path, ["1. Q one", "2. Q two", "Answer Key", "1. a", "2. b"])
    assert A._ws6(ctx, doc, ok, parse_questions(ok, PROFILE)).status == PASS


def test_exam_count_and_tags(tmp_path):
    ctx = make_ctx()
    doc = ref(doc_type="chapter_exam")
    lines = [f"Question {i}: text [SLO: M-06-A-01]" for i in range(1, 7)]
    info = questions_doc(tmp_path, lines)
    res = {f.code: f for f in A.exam_checks(ctx, doc, info, "chapter_exam")}
    assert res["CE1"].status == PASS and res["CE1"].partial and res["CE4"].status == PASS
    info = questions_doc(tmp_path, lines[:3])
    res = {f.code: f for f in A.exam_checks(ctx, doc, info, "chapter_exam")}
    assert res["CE1"].status == FAIL
    info = questions_doc(tmp_path, [f"Question {i}: text" for i in range(1, 7)])
    res = {f.code: f for f in A.exam_checks(ctx, doc, info, "chapter_exam")}
    assert res["CE4"].status == FAIL and "No SLO tags" in res["CE4"].message


import pytest


@pytest.mark.parametrize("fmt", ["{}.", "{})", "Q{}.", "Q.{}", "Q {}:", "Question {})", "({})", "{} -", "Q{}"])
def test_common_english_question_formats_are_recognised(tmp_path, fmt):
    info = questions_doc(tmp_path, [fmt.format(i) + " What is the capital of Pakistan?" for i in range(1, 7)])
    qs = parse_questions(info, PROFILE)
    assert [q.num for q in qs] == [1, 2, 3, 4, 5, 6]


def test_automatic_list_numbering_counts_as_questions(tmp_path):
    d = new_doc()
    for _ in range(6):
        d.add_paragraph("What is the capital of Pakistan?", style="List Number")
    assert len(parse_questions(parse_docx(save(d, tmp_path)), PROFILE)) == 6


@pytest.mark.parametrize("opts", [["(a) x", "(b) y", "(c) z"], ["a) x", "b) y", "c) z"], ["A. x", "B. y", "C. z"]])
def test_mcq_option_styles(tmp_path, opts):
    info = questions_doc(tmp_path, ["Question 1: pick one"] + opts)
    assert parse_questions(info, PROFILE)[0].qtype == "mcq"


def test_unrecognised_question_format_is_reviewed_not_failed(tmp_path):
    # Roman numerals are not a recognised format: the answer must be 'needs reviewer', never a false Fail
    info = questions_doc(tmp_path, [f"{r}. What is it?" for r in ("i", "ii", "iii", "iv", "v", "vi")])
    res = {f.code: f for f in A.exam_checks(make_ctx(), ref(doc_type="chapter_exam"), info, "chapter_exam")}
    assert res["CE1"].status == REVIEW and "not be supported" in res["CE1"].message


def test_lp5_slo_list_ends_at_the_next_heading_or_note_not_the_end_of_the_document(tmp_path):
    # Chapter 1 of the sample is not built on the five sections: an SLO block is followed by notes and content.
    d = new_doc()
    add(d, "Chapter 1: Quran", bold=True)
    add(d, "SLOs", bold=True)
    add(d, "Knowledge", bold=True)
    add(d, "• Understand the translation of the surahs")
    add(d, "Explain the background of the surahs")          # plain paragraph: not a bullet
    add(d, "Note: read these three surahs again")
    add(d, "Surah Al-Fatiha")
    for i in range(30):
        add(d, f"Content line {i} that is not an SLO.")
    add(d, "B - Memorisation", bold=True)
    add(d, "SLOs", bold=True)
    add(d, "•Recite the surahs in prayer")                     # bullet without a space after it
    r = LP.lp5(make_ctx(), ref(), parse_docx(save(d, tmp_path)))
    assert r.status == FAIL and "1 of 3 SLO items in 2 SLO sections" in r.message
    assert [m["text"] for m in r.marks] == ["Explain the background of the surahs"]
