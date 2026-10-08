"""The model reads the structure; code reconciles it with the parser and checks the rules on it.
A stand-in model answers by paragraph text, the way a good model would."""
import re

import docx
from docx.shared import Pt

from course_review import autofix, structure
from course_review.checks import formatting as F
from course_review.checks import guidelines as G
from course_review.checks import lessonplan as L
from course_review.checks import slides as S
from course_review.docx_model import parse_docx
from course_review.fallback import ModelConfig
from course_review.llm import LLMError
from course_review.models import FAIL, PASS, REVIEW

from fakes import _Base
from helpers import make_ctx, ref, save

LINE = re.compile(r"^\[(\d+)\] (?:\{[^}]*\})*\s?(.*)$")


class Reader(_Base):
    """Finds sections, headings and items by the text the test document was built with."""

    def __init__(self, sections, headings, slo=(), warm=(), takeaways=(), not_headings=(), bad=False, roles=None):
        super().__init__()
        self.sections, self.headings, self.not_headings = sections, headings, not_headings
        self.slo, self.warm, self.takeaways, self.bad, self.roles = slo, warm, takeaways, bad, roles or {}

    def _ids(self, user, texts):
        lines = {m.group(2).strip(): int(m.group(1)) for m in map(LINE.match, user.splitlines()) if m}
        return [lines[t] for t in texts if t in lines]

    def chat_json(self, system, user):
        self.calls += 1
        if self.bad:
            return {"something": "else"}
        if system.startswith("You map the structure"):
            ids = {k: (self._ids(user, [v]) or [None])[0] for k, v in self.sections.items()}
            return {"sections": ids, "headings": self._ids(user, self.headings), "slo_items": self._ids(user, self.slo),
                    "warm_up_questions": self._ids(user, self.warm), "takeaway_items": self._ids(user, self.takeaways) + [9999],
                    "glossary": None}
        if system.startswith("You decide which lines"):
            return {"headings": self._ids(user, self.headings), "not_headings": self._ids(user, self.not_headings)}
        if system.startswith("You map the slides"):
            return {"slides": {str(k): v for k, v in self.roles.items()}}
        raise LLMError("unexpected prompt")


def plan(tmp_path, lines):
    """lines: [(text, bold)]"""
    d = docx.Document()
    for t, bold in lines:
        r = d.add_paragraph().add_run(t)
        r.bold, r.font.size = bold, Pt(12)
    return parse_docx(save(d, tmp_path, "plan.docx"))


OFF_TEMPLATE = [
    ("Plants and their Roots", True),
    ("Story Time", True),
    ("Sara helps her grandmother in the garden and finds long white threads under a weed.", False),
    ("What you will learn today", True),
    ("1. Name the main parts of a root system", False),
    ("2. Explain how roots take in water", False),
    ("Let's Begin!", True),
    ("Have you ever pulled a plant out of the ground?", False),
    ("Why does a tree not fall over in a storm?", False),
    ("Main Lesson", True),
    ("Parts of a root", True),
    ("The primary root grows first. Root hairs cover the tips.", False),
    ("Roots hold the plant in the soil.", False),
    ("Root hairs take in water.", False),
]

GOOD_READER = dict(
    # the model points at each section's first line of content, as a real one often does
    sections={"introduction": OFF_TEMPLATE[2][0], "slos": OFF_TEMPLATE[4][0], "warm_up": OFF_TEMPLATE[7][0],
              "concept_building": "Main Lesson", "key_takeaways": "Roots hold the plant in the soil."},
    headings=["Story Time", "What you will learn today", "Let's Begin!", "Main Lesson", "Parts of a root"],
    slo=[OFF_TEMPLATE[4][0], OFF_TEMPLATE[5][0]], warm=[OFF_TEMPLATE[7][0], OFF_TEMPLATE[8][0]],
    takeaways=["Roots hold the plant in the soil.", "Root hairs take in water."])


def read(tmp_path, lines, reader):
    info = plan(tmp_path, lines)
    ctx = make_ctx()
    ctx.model = ModelConfig(client=reader)
    d = ref()
    d.abs, d.rel = str(tmp_path / "plan.docx"), "plan.docx"
    structure.attach(ctx, d, info)
    return ctx, info


def test_parser_alone_finds_nothing_in_an_off_template_plan(tmp_path):
    info = plan(tmp_path, OFF_TEMPLATE)
    r = L.lp3(make_ctx(), ref(), info)
    assert r.status == FAIL and "introduction, slos, warm_up, concept_building, key_takeaways" in r.message


def test_model_reading_finds_renamed_sections_as_suggestions(tmp_path):
    ctx, info = read(tmp_path, OFF_TEMPLATE, Reader(**GOOD_READER))
    st = structure.of(info)
    by = {k: info.paras[v].text for k, v in st.sections.items()}
    assert by["introduction"] == "Story Time" and by["slos"] == "What you will learn today"   # snapped to the heading above
    assert by["warm_up"] == "Let's Begin!" and by["key_takeaways"] == "Roots hold the plant in the soil."
    r = L.lp3(ctx, ref(), info)
    assert r.status == PASS and len(r.evidence) == 5          # the model decides; the evidence says how each was found
    assert "recognised by its content: 'Story Time'" in r.evidence[0]


def test_items_are_read_from_the_structure(tmp_path):
    ctx, info = read(tmp_path, OFF_TEMPLATE, Reader(**GOOD_READER))
    r5 = L.lp5(ctx, ref(), info)
    assert r5.status == FAIL and "2 of 2 SLO items" in r5.message
    assert G.lpg1(ctx, ref(), info).message == "The Warm-up has 2 question(s)"
    r2 = G.lpg2(ctx, ref(), info)
    assert r2.status == FAIL and "not written as bullets" in r2.message      # the made-up paragraph 9999 was dropped


def test_headings_come_from_the_reading(tmp_path):
    ctx, info = read(tmp_path, OFF_TEMPLATE, Reader(**GOOD_READER))
    heads = [p.text for p in F.doc_headings(ctx, info)]
    assert heads == ["Story Time", "What you will learn today", "Let's Begin!", "Main Lesson", "Parts of a root"]
    assert "Plants and their Roots" not in heads                            # the title is not a numbered heading


def test_a_section_the_model_puts_at_a_label_or_list_item_is_not_used(tmp_path):
    lines = [("Lesson", True), ("Introduction", True), ("Story text here for the lesson.", False), ("Knowledge:", True),
             ("• know roots", False)]
    reader = Reader(sections={"introduction": "Introduction", "slos": "Knowledge:", "warm_up": "• know roots"},
                    headings=["Introduction"])
    ctx, info = read(tmp_path, lines, reader)
    st = structure.of(info)
    assert st.sections["slos"] is None and st.sections["warm_up"] is None and st.source["introduction"] == "both"


def test_the_model_decides_which_lines_are_headings(tmp_path):
    lines = [("Lesson", True), ("Introduction", True), ("A letter from Hamza", True), ("Dear Fahad, I learned to be patient.", False),
             ("Why do roots matter?", False)]
    reader = Reader(sections={"introduction": "Introduction"}, headings=["Introduction", "Why do roots matter?"],
                    not_headings=["A letter from Hamza"])
    ctx, info = read(tmp_path, lines, reader)
    heads = [p.text for p in F.doc_headings(ctx, info)]
    assert heads == ["Introduction", "Why do roots matter?"]           # bold is not enough; a question mark does not stop it
    assert structure.heading_findings(ctx, info) == []


def test_an_unusable_reply_falls_back_to_the_parser(tmp_path):
    ctx, info = read(tmp_path, OFF_TEMPLATE, Reader(**GOOD_READER, bad=True))
    assert structure.of(info) is None
    assert any(s.get("structure") and not s["usable"] for s in ctx.model.stats)
    assert L.lp3(ctx, ref(), info).status == FAIL                           # the parser's answer, as before


def test_the_reading_is_cached(tmp_path):
    from course_review.fallback import Cache
    reader = Reader(**GOOD_READER)
    for n in (1, 2):
        info = plan(tmp_path, OFF_TEMPLATE)
        ctx = make_ctx()
        ctx.model = ModelConfig(client=reader, cache=Cache(tmp_path / "cache"))
        d = ref()
        d.abs, d.rel = str(tmp_path / "plan.docx"), "plan.docx"
        structure.attach(ctx, d, info)
        assert structure.of(info) is not None
    assert reader.calls == 2                                                # two questions the first time, none the second


def test_slide_roles_without_labels_come_from_the_model(tmp_path):
    from test_slides_citations import deck
    info = deck(tmp_path, [("Welcome", ["Grade 6"], ""), ("Today you will", ["learn roots"], ""), ("Sara's garden", ["story"], ""),
                           ("Think about it", ["why?"], ""), ("Roots", ["parts"], ""), ("Recap", ["summary"], "")])
    ctx = make_ctx()
    roles = {1: "cover", 2: "slos", 3: "introduction", 4: "warm_up", 5: "concept_building", 6: "key_takeaways"}
    st = structure.reconcile_slides(ctx, info, roles)
    info.structure = st
    r = S.fg7(ctx, ref(rel="g.pptx", doc_type="facilitator_guide"), info)
    assert r.status == PASS                                                   # the model's roles decide


def test_autofix_numbers_only_agreed_headings(tmp_path):
    ctx, info = read(tmp_path, OFF_TEMPLATE, Reader(**GOOD_READER))
    dst = tmp_path / "fixed.docx"
    d = ref()
    d.abs, d.rel = str(tmp_path / "plan.docx"), "plan.docx"
    autofix.fix_docx(str(tmp_path / "plan.docx"), str(dst), ctx, d, info)
    texts = [p.text for p in docx.Document(str(dst)).paragraphs]
    assert "1 Story Time" in texts and "3 Let's Begin!" in texts and "4.1 Parts of a root" in texts


# ------------------------------------------------------------- the corrected copy
RUN_ON = [
    ("Lesson", True),
    ("Introduction", True),
    ("Sara finds roots in the garden and asks what they are for.", False),
    ("Warm-up: Before you read on, answer these questions honestly in your notebook please.", False),
    ("What did you see when you pulled a plant?", False),
    ("Concept Building", True),
    ("Roots hold plants.", False),
    ("Key Takeaways", True),
    ("• Roots hold plants.", False),
]


def test_a_heading_with_its_text_run_on_is_split_and_explained(tmp_path):
    from course_review import annotate
    reader = Reader(sections={"introduction": "Introduction", "warm_up": RUN_ON[3][0], "concept_building": "Concept Building",
                              "key_takeaways": "Key Takeaways"},
                    headings=["Introduction", "Concept Building", "Key Takeaways"])
    ctx, info = read(tmp_path, RUN_ON, reader)
    assert structure.of(info).run_in == {3}
    d = ref()
    d.abs, d.rel = str(tmp_path / "plan.docx"), "plan.docx"
    fixed = tmp_path / "fixed.docx"
    lines, extra = autofix.fix_docx(d.abs, str(fixed), ctx, d, info)
    paras = [(p.style.name, p.text) for p in docx.Document(str(fixed)).paragraphs]
    assert ("Heading 1", "2 Warm-up") in paras
    assert ("Normal", "Before you read on, answer these questions honestly in your notebook please.") in paras   # word for word
    out = tmp_path / "marked.docx"
    annotate.annotate_docx(str(fixed), str(out), [], lines, extra["changes"])
    notes = [c.text for c in docx.Document(str(out)).comments]
    assert any(n.startswith("Fixed for you:") and "were on one line" in n for n in notes)
    marked = docx.Document(str(out))
    pink = [p.text for p in marked.paragraphs if any(r.font.highlight_color == 5 for r in p.runs if r.text.strip())]
    assert "2 Warm-up" in pink and "1 Introduction" in pink


def test_the_structure_follows_the_text_into_the_corrected_copy(tmp_path):
    reader = Reader(sections={"introduction": "Introduction", "warm_up": RUN_ON[3][0], "concept_building": "Concept Building",
                              "key_takeaways": "Key Takeaways"},
                    headings=["Introduction", "Concept Building", "Key Takeaways"])
    ctx, info = read(tmp_path, RUN_ON, reader)
    d = ref()
    d.abs, d.rel = str(tmp_path / "plan.docx"), "plan.docx"
    fixed = tmp_path / "fixed.docx"
    autofix.fix_docx(d.abs, str(fixed), ctx, d, info)
    new = parse_docx(str(fixed))
    moved = structure.transfer(structure.of(info), info, new)
    by = {k: new.paras[v].text for k, v in moved.sections.items() if v is not None}
    assert by == {"introduction": "1 Introduction", "warm_up": "2 Warm-up", "concept_building": "3 Concept Building",
                  "key_takeaways": "4 Key Takeaways"}


def test_a_document_nothing_could_place_becomes_a_lesson_plan_when_the_model_reads_one(tmp_path):
    from course_review.ingest import GENERIC
    info = plan(tmp_path, OFF_TEMPLATE)
    ctx = make_ctx()
    ctx.model = ModelConfig(client=Reader(**GOOD_READER))
    d = ref()
    d.abs, d.rel, d.doc_type, d.chapter, d.lesson = str(tmp_path / "plan.docx"), "plan.docx", GENERIC, None, None
    structure.attach(ctx, d, info)
    assert d.doc_type == "lesson_plan" and (d.chapter, d.lesson) == (1, 1)
    assert any("Recognised as lesson plan by the model" in n for n in d.notes)


def test_a_document_without_lesson_plan_sections_stays_general(tmp_path):
    from course_review.ingest import GENERIC
    info = plan(tmp_path, [("Staff memo", True), ("Please submit your timesheets by Friday.", False)])
    ctx = make_ctx()
    ctx.model = ModelConfig(client=Reader(sections={}, headings=[]))
    d = ref()
    d.abs, d.rel, d.doc_type = str(tmp_path / "plan.docx"), "memo.docx", GENERIC
    assert structure.attach(ctx, d, info) is None and d.doc_type == GENERIC


def test_the_coverage_check_gets_the_models_slo_list_not_every_line_of_the_section(tmp_path):
    from course_review.checks.judgement import lesson_slo_texts, section_texts
    lines = [("Lesson", True), ("Introduction", True), ("Sara finds roots in the garden.", False),
             ("Student Learning Objectives", True), ("• Name the parts of a root", False), ("• Explain how roots take in water", False),
             ("Words to Recall", True), ("Organ: a body part with a job", False), ("Warm-up", True), ("What holds a tree up?", False)]
    reader = Reader(sections={"introduction": "Introduction", "slos": "Student Learning Objectives", "warm_up": "Warm-up"},
                    headings=["Introduction", "Student Learning Objectives", "Words to Recall", "Warm-up"],
                    slo=["• Name the parts of a root", "• Explain how roots take in water"])
    ctx, info = read(tmp_path, lines, reader)
    slos = lesson_slo_texts(ctx, None, section_texts(ctx, info), info)
    assert slos == ["Name the parts of a root", "Explain how roots take in water"]    # not 'Words to Recall' or 'Organ: ...'


def test_slide_checks_find_slides_by_the_models_reading(tmp_path):
    from course_review.checks.judgement import slides_for
    from test_slides_citations import deck
    info = deck(tmp_path, [("Welcome", ["Grade 6"], ""), ("Let's think", ["why?"], ""), ("Roots", ["parts"], "")])
    ctx = make_ctx()
    info.structure = structure.reconcile_slides(ctx, info, {1: "cover", 2: "warm_up", 3: "concept_building"})
    assert [s.index for s in slides_for(ctx, info, "warm_up")] == [2]          # no 'Warm-up' in its title
