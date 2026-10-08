"""The model reads the layout of question and table documents (positions only); code reads the text."""
import re

import docx
from docx.enum.text import WD_COLOR_INDEX

from course_review import layout
from course_review.fallback import Cache, ModelConfig
from course_review.questions import load_bank_items
from course_review.slo import load_slos

from fakes import _Base
from helpers import add, make_ctx, new_doc, save

LINE = re.compile(r"^\[(\d+)\] (.*)$")


class LayoutModel(_Base):
    """Answers the layout questions the way a good model would; `wrong` makes it point at the wrong things."""

    def __init__(self, wrong=False):
        super().__init__()
        self.wrong = wrong

    def chat_json(self, system, user):
        self.calls += 1
        if system.startswith("You read a school quiz document"):
            rows = []
            for m in map(LINE.match, user.splitlines()):
                if m and m.group(2).lower().startswith(("unit", "section")):
                    n = int(re.search(r"\d+", m.group(2)).group())
                    rows.append({"para": int(m.group(1)), "lesson": n + (5 if self.wrong else 0), "part": ""})
            return {"lessons": rows}
        if system.startswith("You read the layout of a school Data Bank"):
            names = {"Item No": "item", "Topic SLO": "slo", "Stem": "question", "Choices": "options", "If right": "feedback_correct",
                     "If wrong": "feedback_incorrect", "Ch": "chapter", "Les": "lesson", "Course": "subject"}
            return {"fields": {k: ("level" if self.wrong else names.get(k)) for k in user.splitlines()}}
        if system.startswith("You read the layout of a school \"Document of Specifications\""):
            tables = [int(t) for t in re.findall(r"^Table (\d+)", user, re.M)]
            return {"tables": [{"table": t, "chapter": t + 1, "lesson_no": 1, "lesson_name": 2, "slo": 0 if not self.wrong else 9,
                                "bloom": None} for t in tables]}
        return {}


def ctx_with(model, tmp_path=None):
    ctx = make_ctx()
    ctx.model = ModelConfig(client=model, cache=Cache(tmp_path / "cache") if tmp_path else None)
    return ctx


class Doc:
    def __init__(self, path):
        self.abs, self.rel = path, path


# ------------------------------------------------------------------ Pop Quiz lessons
def quiz(tmp_path):
    d = new_doc()
    for unit in (1, 2):
        add(d, f"Unit {unit}: Plants")                 # not 'Lesson N': the rules do not know this heading
        add(d, "1. What do roots do?")
        add(d, "A) Make food")
        add(d, "B) Take in water", highlight=WD_COLOR_INDEX.YELLOW)
        add(d, "C) Hold seeds")
    path = save(d, tmp_path, "quiz.docx")
    from course_review.docx_model import parse_docx
    return path, parse_docx(path)


def test_lesson_sections_the_rules_miss_are_read_by_the_model(tmp_path):
    from course_review.questions import segment_lessons
    path, info = quiz(tmp_path)
    assert segment_lessons(info, make_ctx().profile) == []
    heads = layout.read_lesson_heads(ctx_with(LayoutModel()), Doc(path), info)
    assert [(h["num"], info.paras[h["start"]].text) for h in heads] == [(1, "Unit 1: Plants"), (2, "Unit 2: Plants")]
    assert [len(g) for g in layout.group_lessons(heads, len(info.paras))] == [2]


def test_a_lesson_number_that_is_not_in_the_heading_is_rejected(tmp_path):
    path, info = quiz(tmp_path)
    assert layout.read_lesson_heads(ctx_with(LayoutModel(wrong=True)), Doc(path), info) is None


# ------------------------------------------------------------------ Data Bank
def wide_bank(tmp_path, headers=("Item No", "Course", "Ch", "Les", "Topic SLO", "Stem", "Choices", "If right", "If wrong")):
    d = docx.Document()
    t = d.add_table(rows=3, cols=len(headers))
    for c, h in zip(t.rows[0].cells, headers):
        c.text = h
    for r, n in zip(t.rows[1:], (1, 2)):
        vals = [f"DB-{n}", "Science", "1", "3", "L3-SLO1", f"Which is a solid? ({n})", "A) Steam B) Wood", "Right: wood keeps its shape.",
                "Not quite: think about shape."]
        for c, v in zip(r.cells, vals):
            c.text = v
    return save(d, tmp_path, "bank.docx")


def test_a_bank_with_its_own_field_names_is_read_through_the_models_mapping(tmp_path):
    path = wide_bank(tmp_path)
    assert load_bank_items(path, make_ctx().profile) == []           # the profile does not know these names
    reading = layout.read_bank_fields(ctx_with(LayoutModel()), path)
    assert reading[0] == "header"
    items = load_bank_items(path, make_ctx().profile, reading)
    assert len(items) == 2 and items[0].fields["question"] == "Which is a solid? (1)"
    assert items[0].fields["feedback_incorrect"] == "Not quite: think about shape." and (items[0].chapter, items[0].lesson_no) == (1, 3)


def test_the_templates_wide_table_is_read_with_the_profiles_names_too(tmp_path):
    path = wide_bank(tmp_path, headers=("Item #", "Subject", "Chapter", "Lesson", "SLO", "Question", "Options", "Correct Feedback",
                                        "Incorrect Feedback"))
    items = load_bank_items(path, make_ctx().profile)
    assert len(items) == 2 and items[1].fields["slo"] == "L3-SLO1"


def test_a_field_mapping_without_a_question_field_is_rejected(tmp_path):
    assert layout.read_bank_fields(ctx_with(LayoutModel(wrong=True)), wide_bank(tmp_path)) is None


# ------------------------------------------------------------------ specifications
def spec(tmp_path):
    d = docx.Document()
    for ch in (1, 2):
        d.add_paragraph(f"باب {ch}")                     # no 'Chapter N' line: the rules cannot place the tables
        t = d.add_table(rows=3, cols=3)
        for c, h in zip(t.rows[0].cells, ("Outcomes", "No.", "Title")):
            c.text = h
        for r, (slo, no) in zip(t.rows[1:], ((f"Explain idea {ch}a", "1"), (f"Explain idea {ch}b\nApply idea {ch}c", "2"))):
            for c, v in zip(r.cells, (slo, no, "Lesson")):
                c.text = v
    return save(d, tmp_path, "spec.docx")


def test_specification_tables_the_rules_cannot_read_are_read_through_the_models_layout(tmp_path):
    path = spec(tmp_path)
    assert load_slos(path) == {}
    plan = layout.read_spec_tables(ctx_with(LayoutModel()), path)
    slos = load_slos(path, plan)
    assert [s.text for s in slos[(1, 2)]] == ["Explain idea 1b", "Apply idea 1c"]     # one SLO per line of the cell
    assert [s.text for s in slos[(2, 1)]] == ["Explain idea 2a"]


def test_a_column_that_does_not_exist_is_rejected(tmp_path):
    assert layout.read_spec_tables(ctx_with(LayoutModel(wrong=True)), spec(tmp_path)) is None


def test_layout_readings_are_cached(tmp_path):
    path = spec(tmp_path)
    m = LayoutModel()
    layout.read_spec_tables(ctx_with(m, tmp_path), path)
    layout.read_spec_tables(ctx_with(m, tmp_path), path)
    assert m.calls == 1
